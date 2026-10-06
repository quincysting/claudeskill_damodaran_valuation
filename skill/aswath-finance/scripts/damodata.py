"""damodata: Damodaran's datasets as queryable data (industry averages by region, historical returns, implied ERP).

Every number comes from Damodaran's own spreadsheets (pages.stern.nyu.edu/~adamodar, New_Home_Page/data.html), transcribed
column-for-column with the source column names and units (decimals, US $ millions where the source says so). Every answer carries the
source file, its URL and its as-of date. Period statistics on historical returns are the spreadsheet's 'Returns by year' formulas
(histretSP.xlsx rows 121-129) re-run over any start/end year.

Sources (URL = https://pages.stern.nyu.edu/~adamodar/pc/<path>):
  datasets/<file>  sheet 'Industry Averages' (+ 'Variables & FAQ' definitions), Jan 2026 update, 7 regions
                   (US, Global, Europe, Japan, Emerging, China, India):
                   betas/beta*.xls, wacc*.xls, margin*.xls, EVA*.xls, dbtfund*.xls, debtdetails*.xls, divfcfe*.xls,
                   divfund*.xls;  US only: taxrate.xls, mgnroc.xls
  blog/<file>      US & Global medians/quartiles: IndustryMargins{US,Global}.xlsx ('Margins'),
                   IndustryReturns{US,Global}.xlsx (both sheets), DebtIndustry{US,Global}2025.xlsx (both sheets),
                   DividendsIndustry{US,Global}.xlsx, BuybacksIndustry{US,Global}.xlsx, WACCdistn.xlsx (Global)
  datasets/histretSP.xlsx  'Returns by year' A21:H118 + U (annual returns 1928-2025, inflation)
  datasets/histimpl.xlsx   'Historical Impl Premiums' A7:S73 (implied ERP by year, 1960-2025)
  implprem/ERPbymonth.xlsx 'Historical ERP' A1:Q218 (implied ERP at the start of each month, Sep 2008-)
  datasets/indname.xlsx    'By industry' (48,156 companies -> industry group, country, region group)

Usage (add --json to any command for machine-readable output):
  python damodata.py datasets                                  # every dataset, region, file, as-of date
  python damodata.py columns wacc                              # Damodaran's column names + definitions
  python damodata.py industry "Semiconductor" --region Global  # every metric the datasets have for that industry
  python damodata.py industry Software --metric margin         # only columns matching 'margin'
  python damodata.py industry --region Europe                  # list Damodaran's industry names
  python damodata.py rank "Cost of Capital" --region US --top 10 [--asc]
  python damodata.py rank "IndustryMargins:Operating Margin: Median" --region Global
  python damodata.py history --start 1928 --end 2025           # arith/geo returns, real returns, premiums, std errors
  python damodata.py history --year 2008                       # one year's returns
  python damodata.py erp [--year 2025 | --month 2026-09 | --start 2016 --end 2025 | --monthly --start 2025-01]
  python damodata.py company NVDA                              # Damodaran's industry classification for a company
  python damodata.py --selftest                                # spot-check cells + the histretSP/histimpl summary rows
  python damodata.py --build-cases <spreadsheets/*.xls*>       # rebuild data/damodata_data.json + cases (openpyxl, xlrd)

Region names: US, Global, Europe, Japan, Emerging, China, India. Industry and metric names match case-insensitively,
exact first, else a unique substring; 'dataset:column' pins a column (e.g. 'wacc:Cost of Capital').
"""

from __future__ import annotations

import argparse
import json
import math
import re
import statistics
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
DATA = HERE.parent / "data" / "damodata_data.json"
CASES = HERE.parent / "data" / "damodata_cases.json"
SITE = "https://pages.stern.nyu.edu/~adamodar/pc/"
REGIONS = {"US": "", "Global": "Global", "Europe": "Europe", "Japan": "Japan", "Emerging": "emerg", "China": "China",
           "India": "India"}
TOTALS = ("Total Market", "Total Market (without financials)", "All firms")
STOCKS, SMALL, BILLS, BONDS, BAA = ("S&P 500 (includes dividends)", "US Small cap (bottom decile)", "3-month T.Bill",
                                    "US T. Bond (10-year)", "Baa Corporate Bond")
PREMIUMS = {"Stocks - T.Bills": (STOCKS, BILLS), "Stocks - T.Bonds": (STOCKS, BONDS),
            "Stocks - Baa Corp Bond": (STOCKS, BAA), "Small Cap Premium (small cap - S&P 500)": (SMALL, STOCKS)}


# ------------------------------------------------------------------ lookup
def load(path: Path = DATA) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def pick(query: str, options: list[str], what: str) -> str:
    """Exact (case-insensitive) match, else a unique substring match, else a helpful error."""
    q = query.strip().lower()
    hits = [o for o in options if o.lower() == q] or [o for o in options if q in o.lower()]
    if len(hits) == 1:
        return hits[0]
    if not hits:
        raise KeyError(f"no {what} matches '{query}'. Options: {', '.join(options)}")
    raise KeyError(f"'{query}' matches several ({what}): {' | '.join(hits)}")


def region_name(region: str) -> str:
    alias = {"emerg": "Emerging", "em": "Emerging", "emerging markets": "Emerging", "europe": "Europe"}
    return alias.get(region.lower()) or pick(region, list(REGIONS), "region")


def tables(data: dict, region: str, dataset: str | None = None) -> list[tuple[str, dict]]:
    names = [pick(dataset, list(data["datasets"]), "dataset")] if dataset else list(data["datasets"])
    return [(d, data["datasets"][d]["regions"][region]) for d in names if region in data["datasets"][d]["regions"]]


def meta(name: str, t: dict) -> dict:
    return {"dataset": name, "file": t["file"], "url": t["url"], "as_of": t["as_of"], "sheet": t["sheet"]}


def industry(data: dict, name: str, region: str = "US", dataset: str | None = None, metric: str | None = None) -> dict:
    region = region_name(region)
    tabs = tables(data, region, dataset)
    names = sorted({i for _, t in tabs for i in t["rows"]})
    ind = pick(name, names, "industry")
    out = []
    for d, t in tabs:
        row = t["rows"].get(ind)
        if row is None:
            continue
        vals = {c: v for c, v in zip(t["columns"], row) if not metric or metric.lower() in c.lower()}
        if vals:
            out.append({**meta(d, t), "values": vals})
    return {"industry": ind, "region": region, "results": out}


def resolve_metric(data: dict, metric: str, region: str, dataset: str | None = None) -> tuple[str, str]:
    if ":" in metric and metric.split(":")[0].strip().lower() in [d.lower() for d in data["datasets"]]:
        dataset, metric = metric.split(":", 1)
    cols = [(d, c) for d, t in tables(data, region, dataset) for c in t["columns"]]
    exact = [dc for dc in cols if dc[1].lower() == metric.strip().lower()]
    if exact:  # same column name in several of the source files (e.g. Beta in betas/wacc/EVA): first in dataset order
        return exact[0]
    hits = [dc for dc in cols if metric.strip().lower() in dc[1].lower()]
    if len(hits) == 1:
        return hits[0]
    listing = " | ".join(f"{d}:{c}" for d, c in (hits or cols))
    raise KeyError(f"metric '{metric}' {'is ambiguous' if hits else 'not found'} in {region}; use dataset:column. "
                   f"Candidates: {listing}")


def rank(data: dict, metric: str, region: str = "US", dataset: str | None = None, top: int | None = None,
         ascending: bool = False) -> dict:
    region = region_name(region)
    d, col = resolve_metric(data, metric, region, dataset)
    t = data["datasets"][d]["regions"][region]
    k, nf = t["columns"].index(col), next((i for i, c in enumerate(t["columns"]) if "firms" in c.lower()), None)
    rows = [(i, r[k], r[nf] if nf is not None else None) for i, r in t["rows"].items() if r[k] is not None]
    body = sorted([x for x in rows if x[0] not in TOTALS], key=lambda x: x[1], reverse=not ascending)
    return {**meta(d, t), "region": region, "metric": col, "order": "ascending" if ascending else "descending",
            "ranking": [{"industry": i, "value": v, "firms": n} for i, v, n in body[:top]],
            "count": len(body), "totals": [{"industry": i, "value": v, "firms": n} for i, v, n in rows if i in TOTALS]}


def company(data: dict, query: str, limit: int = 25) -> dict:
    c = data["companies"]
    q = query.strip().upper()
    tiers: list[list] = [[], [], []]
    for name, ii, ci in c["rows"]:
        u = name.upper()
        if u.endswith(f":{q})") or u.endswith(f"({q})"):
            tiers[0].append((name, ii, ci))
        elif u.startswith(q):
            tiers[1].append((name, ii, ci))
        elif q in u:
            tiers[2].append((name, ii, ci))
    hits = [h for tier in tiers for h in sorted(tier)]
    rows = [{"company": n, "industry": c["industries"][i], "country": c["countries"][k],
             "broad_group": c["country_groups"][c["countries"][k]][0],
             "sub_group": c["country_groups"][c["countries"][k]][1]} for n, i, k in hits[:limit]]
    return {"file": c["file"], "url": c["url"], "as_of": c["as_of"], "query": query, "matches": len(hits), "rows": rows}


# ------------------------------------------------------------------ history (histretSP / histimpl / ERPbymonth)
def _geo(xs: list[float]) -> float:
    return math.prod(1 + x for x in xs) ** (1 / len(xs)) - 1


def _series(data: dict, start: int, end: int) -> tuple[list[int], dict, dict]:
    h = data["history"]
    rows = [r for r in h["rows"] if start <= r[0] <= end]
    if not rows:
        raise ValueError(f"no annual returns in {start}-{end}; Damodaran's data covers {h['rows'][0][0]}-{h['rows'][-1][0]}")
    nominal = {c: [r[i] for r in rows] for i, c in enumerate(h["columns"]) if i}
    infl = nominal.pop("Inflation Rate")
    real = {c: [(1 + x) / (1 + p) - 1 for x, p in zip(v, infl)] for c, v in nominal.items()}  # the sheet's columns V:AB
    return [r[0] for r in rows], nominal, real


def period_stats(data: dict, start: int, end: int) -> dict:
    """The spreadsheet's 'Returns by year' summary formulas for any window: AVERAGE (arithmetic), compounded value^(1/n)-1
    (geometric), premium = difference of the two averages, std error = STDEV(annual premium)/sqrt(n)."""
    years, nominal, real = _series(data, start, end)
    n = len(years)

    def block(ser: dict) -> tuple[dict, dict]:
        avg = {c: {"arithmetic": statistics.fmean(v), "geometric": _geo(v)} for c, v in ser.items()}
        prem = {}
        for name, (a, b) in PREMIUMS.items():
            diff = [x - y for x, y in zip(ser[a], ser[b])]
            prem[name] = {"arithmetic": statistics.fmean(diff), "geometric": _geo(ser[a]) - _geo(ser[b]),
                          "std_error": statistics.stdev(diff) / math.sqrt(n) if n > 1 else None}
        return avg, prem

    nom, nprem = block(nominal)
    rl, rprem = block(real)
    h = data["history"]
    # the sheet's figure, the sheet's way: its column J is =100*(1+C) every year (not cumulative), so C126:C128 = (J_end/J_start-1)^(1/n)-1
    c = {r[0]: r[h["columns"].index(SMALL)] for r in h["rows"]}
    his_small = ((1 + c[years[-1]]) / (1 + c.get(years[0] - 1, 0.0))) ** (1 / n) - 1
    his_sheet = {"nominal": {SMALL: {"geometric": his_small}}, "premiums": {
        "Small Cap Premium (small cap - S&P 500)": {"geometric": his_small - nom[STOCKS]["geometric"]}}}
    return {"file": h["file"], "url": h["url"], "as_of": h["as_of"], "start": years[0], "end": years[-1], "n": n,
            "nominal": nom, "premiums": nprem, "real": rl, "real_premiums": rprem, "his_sheet": his_sheet}


def year_row(data: dict, year: int) -> dict:
    years, nominal, real = _series(data, year, year)
    h = data["history"]
    infl = next(r for r in h["rows"] if r[0] == year)[-1]
    since = period_stats(data, h["rows"][0][0], year)["premiums"]["Stocks - T.Bonds"]["geometric"] if year > h["rows"][0][0] else None
    return {"file": h["file"], "url": h["url"], "as_of": h["as_of"], "year": year,
            "nominal": {c: v[0] for c, v in nominal.items()}, "Inflation Rate": infl,
            "real": {c: v[0] for c, v in real.items()},
            "Historical ERP (geometric stocks - T.Bonds since 1928)": since}


def implied_erp(data: dict, year: int | None = None, start: int | None = None, end: int | None = None) -> dict:
    """histimpl.xlsx. Period average = the sheet's C78:D80 formulas: AVERAGE of 'Implied ERP (FCFE)' over the years that have
    one, and ERP + AVERAGE('T.Bond Rate') over the same rows."""
    t = data["implied_erp"]
    cols = t["columns"]
    out = {"file": t["file"], "url": t["url"], "as_of": t["as_of"], "columns": cols}
    if year is not None:
        row = next((r for r in t["rows"] if r[0] == year), None)
        if row is None:
            raise KeyError(f"no implied ERP for {year}; Damodaran's data covers {t['rows'][0][0]}-{t['rows'][-1][0]}")
        return {**out, "rows": [row]}
    start = start if start is not None else t["rows"][-1][0] - 9
    end = end if end is not None else t["rows"][-1][0]
    rows = [r for r in t["rows"] if start <= r[0] <= end]
    e, b = cols.index("Implied ERP (FCFE)"), cols.index("T.Bond Rate")
    used = [r for r in rows if r[e] is not None]
    if not used:
        raise KeyError(f"no implied ERP in {start}-{end}")
    avg, rf = statistics.fmean(r[e] for r in used), statistics.fmean(r[b] for r in used)
    return {**out, "rows": rows, "average": {"years": f"{used[0][0]}-{used[-1][0]}", "n": len(used),
                                            "Implied ERP (FCFE)": avg, "T.Bond Rate": rf, "ERP + Riskfree Rate": avg + rf}}


def monthly_erp(data: dict, month: str | None = None, start: str | None = None, end: str | None = None) -> dict:
    t = data["erp_monthly"]
    out = {"file": t["file"], "url": t["url"], "as_of": t["as_of"], "columns": t["columns"]}
    if month:
        rows = [r for r in t["rows"] if r[0] == month]
        if not rows:
            raise KeyError(f"no month {month}; Damodaran's data covers {t['rows'][0][0]} to {t['rows'][-1][0]} (YYYY-MM)")
        return {**out, "rows": rows}
    start, end = start or t["rows"][-12][0], end or t["rows"][-1][0]
    return {**out, "rows": [r for r in t["rows"] if start <= r[0] <= end]}


# ------------------------------------------------------------------ reports
def fmt(v) -> str:
    if v is None:
        return "NA"
    if isinstance(v, float):
        return f"{v:.6g}"
    return str(v)


def src(m: dict) -> str:
    return f"[{m['file']} | as of {m['as_of']} | {m['url']}]"


def report_industry(r: dict) -> str:
    lines = [f"{r['industry']} ({r['region']})"]
    for t in r["results"]:
        lines.append(f"\n  {t['dataset']}  {src(t)}")
        lines += [f"    {c:60s} {fmt(v)}" for c, v in t["values"].items()]
    return "\n".join(lines)


def report_rank(r: dict) -> str:
    lines = [f"{r['metric']} by industry, {r['region']} ({r['order']}, {r['count']} industries)  {src(r)}"]
    lines += [f"  {k:3d}. {x['industry']:45s} {fmt(x['value']):>12s}   firms {fmt(x['firms'])}"
              for k, x in enumerate(r["ranking"], 1)]
    lines += [f"       {x['industry']:45s} {fmt(x['value']):>12s}   firms {fmt(x['firms'])}" for x in r["totals"]]
    return "\n".join(lines)


def report_history(s: dict) -> str:
    lines = [f"US annual returns {s['start']}-{s['end']} (n={s['n']})  {src(s)}",
             f"  {'':32s} {'arith':>9s} {'geom':>9s} {'real arith':>11s} {'real geom':>10s}"]
    for c, v in s["nominal"].items():
        rv = s["real"][c]
        lines.append(f"  {c:32s} {v['arithmetic']:9.2%} {v['geometric']:9.2%} {rv['arithmetic']:11.2%} {rv['geometric']:10.2%}")
    lines.append(f"\n  {'Risk premium':40s} {'arith':>9s} {'geom':>9s} {'std err':>9s} {'real ar':>9s} {'real geo':>9s}")
    for c, v in s["premiums"].items():
        rv = s["real_premiums"][c]
        se = "NA" if v["std_error"] is None else f"{v['std_error']:.2%}"
        lines.append(f"  {c:40s} {v['arithmetic']:9.2%} {v['geometric']:9.2%} {se:>9s} {rv['arithmetic']:9.2%} {rv['geometric']:9.2%}")
    hs = s["his_sheet"]
    lines.append(f"\n  NOTE SHEET ERROR: the histretSP small-cap 'Cum Value' column J is =100*(1+C) every year (not compounded), "
                 f"so the source sheet's formula gives small-cap geometric {hs['nominal'][SMALL]['geometric']:.2%} and small-cap "
                 f"premium {hs['premiums']['Small Cap Premium (small cap - S&P 500)']['geometric']:.2%} for this window "
                 f"(JSON 'his_sheet'); the figures above compound the annual returns.")
    return "\n".join(lines)


def report_rows(r: dict, title: str) -> str:
    lines = [f"{title}  {src(r)}"]
    cols = r["columns"]
    for row in r["rows"]:
        lines.append("  " + str(row[0]) + ": " + ", ".join(f"{c}={fmt(v)}" for c, v in zip(cols[1:], row[1:]) if v is not None))
    if "average" in r:
        a = r["average"]
        lines.append(f"  average {a['years']} (n={a['n']}): Implied ERP (FCFE) {a['Implied ERP (FCFE)']:.4%}, "
                     f"T.Bond Rate {a['T.Bond Rate']:.4%}, ERP + Riskfree Rate {a['ERP + Riskfree Rate']:.4%}")
    return "\n".join(lines)


# ------------------------------------------------------------------ selftest
def _dig(obj, path):
    for p in path:
        obj = obj[p]
    return obj


def case_value(data: dict, c: dict):
    k = c["kind"]
    if k == "cell":
        t = data["datasets"][c["dataset"]]["regions"][c["region"]]
        return t["rows"][c["industry"]][t["columns"].index(c["column"])]
    if k == "company":
        r = [x for x in company(data, c["company"], limit=10**6)["rows"] if x["company"] == c["company"]]
        return r[0][c["field"]] if r else None
    if k == "table":
        t = data[c["table"]]
        return next(r for r in t["rows"] if r[0] == c["key"])[t["columns"].index(c["column"])]
    if k == "period":
        return _dig(period_stats(data, c["start"], c["end"]), c["path"])
    if k == "implied_avg":
        return implied_erp(data, start=c["start"], end=c["end"])["average"][c["field"]]
    raise ValueError(k)


def close(got, his) -> bool:
    if isinstance(his, str) or isinstance(got, str) or his is None or got is None:
        return got == his
    return abs(got - his) <= 1e-6 or abs(got - his) <= 5e-4 * abs(his)


def selftest() -> int:
    data, cases = load(), json.loads(CASES.read_text(encoding="utf-8"))
    groups: dict[str, list[int]] = {}
    bad = known = 0
    for c in cases:
        try:
            got = case_value(data, c)
        except Exception as e:  # a lookup that fails is a failed case
            got = f"ERROR {e}"
        if c.get("his_error") and c["kind"] == "period":  # the sheet's erroneous figure, reproduced its way where we can
            try:
                got = _dig(period_stats(data, c["start"], c["end"])["his_sheet"], c["path"])
                c = {k: v for k, v in c.items() if k != "his_error"}
            except KeyError:
                pass
        ok = close(got, c["value"])
        g = groups.setdefault(c["file"], [0, 0])
        if c.get("his_error"):
            known += 1
            print(f"src-err {c['file']} {c['cell']}: Damodaran {fmt(c['value'])} ours {fmt(got)} -- {c['his_error']}")
            continue
        g[0] += ok
        g[1] += 1
        if not ok:
            bad += 1
            print(f"BAD {c['file']} {c['cell']} {c.get('industry', c.get('key', ''))} {c.get('column', c.get('path', ''))}: "
                  f"Damodaran {fmt(c['value'])} ours {fmt(got)}")
    for f, (ok, n) in groups.items():
        if n:
            print(f"{'ok ' if ok == n else 'BAD'} {f:32s} {ok}/{n}")
    total = sum(n for _, n in groups.values())
    print(f"{total - bad}/{total} cases match Damodaran's spreadsheets; {known} cells documented as source-sheet errors")
    return 1 if bad else 0


# ------------------------------------------------------------------ build (maintenance only: openpyxl + xlrd)
XLS_RE = re.compile(r"^(betas|beta|wacc|margin|EVA|dbtfund|debtdetails|divfcfe|divfund|taxrate|mgnroc)"
                    r"(|Global|Europe|Japan|emerg|China|India)\.xls$")
BLOG_RE = re.compile(r"^(IndustryMargins|IndustryReturns|DividendsIndustry|BuybacksIndustry|DebtIndustry)(US|Global)"
                     r"(2025)?\.xlsx$")
ORDER = ["betas", "wacc", "margin", "EVA", "dbtfund", "debtdetails", "divfund", "divfcfe", "taxrate", "mgnroc",
         "IndustryMargins", "IndustryReturns", "DebtIndustry", "DividendsIndustry", "BuybacksIndustry", "WACCdistn"]
HEADER_KEYS = ("industry name", "industry group", "primary sector", "sub region")
SKIP_SETTINGS = ("date updated", "created by", "home page", "data website", "companies in each", "variable definitions")


def _sig(v):
    """Numbers kept to 10 significant digits (files stay small, far inside the 1e-6 / 0.05% calibration tolerance)."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        if not math.isfinite(v):
            return None
        v = float(f"{v:.10g}")
        return int(v) if v.is_integer() and abs(v) < 1e15 else v
    return None


def _num(v, notes: list, where: str):
    if isinstance(v, str):
        s = orig = v.strip()
        if s in ("", "NA", "N/A", "-") or s.startswith("#"):
            return None
        if re.fullmatch(r"-?\d+,\d+%?\*?", s):  # decimal comma typed into the source sheet ('6,36%', '84,88')
            s = s.replace(",", ".")
        try:
            x = float(s[:-1]) / 100 if s.endswith("%") else float(s.rstrip("*"))
        except ValueError:
            notes.append(f"{where}: text '{orig}' stored as null")
            return None
        notes.append(f"{where}: text '{orig}' in the source sheet stored as {x:g}")
        return _sig(x)
    return _sig(v)


def _rows(path: Path):
    """[(sheet name, rows as lists of python values)] with Excel error cells as None."""
    if path.suffix == ".xls":
        import xlrd
        wb = xlrd.open_workbook(str(path))
        out = []
        for sh in wb.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, wb.datemode))
                    elif cell.ctype in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_TEXT):
                        row.append(cell.value if cell.value != "" else None)
                    else:
                        row.append(None)
                rows.append(row)
            out.append((sh.name, rows))
        return out
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]


def _text(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def _names(header: list, above: list) -> list[str]:
    """The source header text; a column whose name repeats in the row or is a bare year/'Average..' gets the group header
    above it as prefix ('Return on Equity: Median', 'Unlevered beta corrected for cash - Over time: 2024')."""
    names = [_text(h) for h in header]
    group, groups = "", []
    for a in list(above) + [None] * (len(names) - len(above)):
        if isinstance(a, str) and a.strip() and not a.strip().startswith("http"):
            group = _text(a)
        groups.append(group)
    out = []
    for i, n in enumerate(names):
        needs = names.count(n) > 1 or re.match(r"^(\d{4}|Average)", n)
        out.append(f"{groups[i]}: {n}" if needs and groups[i] else n)
    return out


def _table(rows: list, notes: list, where: str) -> dict | None:
    h = next((i for i, r in enumerate(rows) if r and isinstance(r[0], str) and r[0].strip().lower() in HEADER_KEYS), None)
    if h is None:
        return None
    last = max(i for i, x in enumerate(rows[h]) if _text(x))
    cols = _names(rows[h][1:last + 1], (rows[h - 1] if h else [])[1:last + 1])
    body, coords = {}, {}
    for j in range(h + 1, len(rows)):
        r = rows[j]
        if not r or not _text(r[0]):
            break
        name = _text(r[0])
        body[name] = [_num(r[k] if k < len(r) else None, notes, f"{where} row {j + 1}") for k in range(1, last + 1)]
        coords[name] = j
    return {"h": h, "columns": cols, "rows": body, "coords": coords}


def _settings(rows: list, h: int) -> dict:
    out = {}
    for r in rows[:h]:
        if r and isinstance(r[0], str) and not r[0].strip().lower().startswith(SKIP_SETTINGS):
            val = next((x for x in r[1:6] if _text(x)), None)
            if val is not None:
                out[_text(r[0])] = _sig(val) if isinstance(val, (int, float)) else _text(val)
    return out


def _definitions(sheets: list) -> tuple[str, dict]:
    purpose, defs = "", {}
    for name, rows in sheets:
        if "faq" not in name.lower():
            continue  # 'Variables & FAQ' / 'Explanations & FAQ'
        start = next((i for i, r in enumerate(rows) if r and _text(r[0]) == "Variable"), None)
        purpose = next((_text(r[1]) for r in rows if r and _text(r[0]) == "End Game" and len(r) > 1), "")
        for r in rows[start + 1:] if start is not None else []:
            if r and _text(r[0]) and len(r) > 1 and _text(r[1]):
                defs[_text(r[0])] = _text(r[1])
    return purpose, defs


def _posts(md_dir: Path) -> dict:
    """spreadsheet file name -> latest post (published date, title, url) that links it."""
    out = {}
    for p in sorted(md_dir.glob("*.md")) if md_dir.exists() else []:
        text = p.read_text(encoding="utf-8", errors="ignore")
        head = {k: v.strip().strip("'\"") for k, v in re.findall(r"^(title|source_url|published): (.*)$", text[:1500], re.M)}
        for f in set(re.findall(r"adamodar/pc/(?:[\w]+/)*([\w&]+\.xlsx?)", text)):
            if f not in out or head.get("published", "") > out[f]["published"]:
                out[f] = {"published": head.get("published", "")[:10], "title": head.get("title", ""),
                          "url": head.get("source_url", "")}
    return out


def _col(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def build(files: list[str]) -> int:
    import random
    paths = [Path(f) for f in files]
    urls = json.loads((paths[0].parent / "_sources.json").read_text(encoding="utf-8")) \
        if (paths[0].parent / "_sources.json").exists() else {}
    posts = _posts(paths[0].resolve().parent.parent / "markdown")
    notes: list[str] = []
    data = {"source": "Aswath Damodaran, https://pages.stern.nyu.edu/~adamodar/New_Home_Page/data.html",
            "notes": notes, "datasets": {}}
    cases: list[dict] = []

    def spot(fname, sheet, tab, dataset, region, rows, k=5):
        rnd = random.Random(f"{fname}|{sheet}")
        cells = [(name, j, c) for name, j in tab["coords"].items() for c in range(len(tab["columns"]))
                 if tab["rows"][name][c] is not None]
        for name, j, c in rnd.sample(cells, min(k, len(cells))):
            raw = rows[j][c + 1]
            cases.append({"kind": "cell", "file": fname, "sheet": sheet, "cell": f"{_col(c + 1)}{j + 1}",
                          "dataset": dataset, "region": region, "industry": name, "column": tab["columns"][c],
                          "his_header": _text(rows[tab["h"]][c + 1]), "value": float(raw) if not isinstance(raw, str) else raw})

    for p in sorted(paths, key=lambda x: x.name.lower()):
        m, b = XLS_RE.match(p.name), BLOG_RE.match(p.name)
        if not (m or b or p.name == "WACCdistn.xlsx"):
            continue
        key = ("betas" if m.group(1) == "beta" else m.group(1)) if m else (b.group(1) if b else "WACCdistn")
        region = next(r for r, s in REGIONS.items() if s == m.group(2)) if m else (b.group(2) if b else "Global")
        sheets = _rows(p)
        url = urls.get(p.name, SITE + ("datasets/" if m else "blog/") + p.name)
        ds = data["datasets"].setdefault(key, {"regions": {}})
        entry = None
        for sheet, rows in sheets:
            if "raw data" in sheet.lower() or "faq" in sheet.lower():  # raw sums repeat the 'Margins' quartiles
                continue
            tab = _table(rows, notes, f"{p.name}!{sheet}")
            if not tab:
                continue
            spot(p.name, sheet, tab, key, region, rows)
            if entry is None:
                post = posts.get(p.name, {})
                if m:
                    d = rows[0][1]
                    if isinstance(d, (int, float)):  # Excel serial date (1900 system)
                        import datetime
                        d = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=d)
                    as_of = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else _text(d)
                else:
                    as_of = f"{post.get('published', '?')} (date of the post that links it; the file has no date cell)"
                entry = {"file": p.name, "url": url, "as_of": as_of, "sheet": sheet,
                         "settings": _settings(rows, tab["h"]) if m else {}, "columns": [], "rows": {}}
                if post:
                    entry["post"] = post
                ds["regions"][region] = entry
            else:
                entry["sheet"] += f" + {sheet}"
            new = [i for i, c in enumerate(tab["columns"]) if c not in entry["columns"]]
            entry["columns"] += [tab["columns"][i] for i in new]
            for name, vals in tab["rows"].items():
                entry["rows"].setdefault(name, []).extend(vals[i] for i in new)
        if m and region == "US":
            ds["purpose"], ds["definitions"] = _definitions(sheets)
    # the source header cell is stale in EVAGlobal.xls (reads 2025-01-05; firm counts and betas equal the Jan 2026 files)
    eg = data["datasets"].get("EVA", {}).get("regions", {}).get("Global")
    if eg and eg["as_of"].startswith("2025"):
        eg["as_of"] = "2026-01-05"
        notes.append("EVAGlobal.xls: 'Date updated' cell reads 2025-01-05 but its firm counts, betas and settings "
                     "equal the 2026-01-05 waccGlobal.xls; as_of set to 2026-01-05")

    byname = {p.name: p for p in paths}
    if "histretSP.xlsx" in byname:
        cases += _build_history(byname["histretSP.xlsx"], data, urls, notes)
    if "histimpl.xlsx" in byname:
        cases += _build_implied(byname["histimpl.xlsx"], data, urls, notes)
    if "ERPbymonth.xlsx" in byname:
        cases += _build_monthly(byname["ERPbymonth.xlsx"], data, urls, notes)
    if "indname.xlsx" in byname:
        cases += _build_companies(byname["indname.xlsx"], data, urls)

    order = [k for k in ORDER if k in data["datasets"]] + [k for k in data["datasets"] if k not in ORDER]
    data["datasets"] = {k: data["datasets"][k] for k in order}
    DATA.write_text(json.dumps(data, separators=(",", ":"), ensure_ascii=False), encoding="utf-8")
    CASES.write_text(json.dumps(cases, indent=0, ensure_ascii=False), encoding="utf-8")
    names = {k: sorted(set(r["rows"])) for k, d in data["datasets"].items() for r in d["regions"].values()}
    ref = names.get("wacc", [])
    for k, n in names.items():
        if set(n) - set(TOTALS) != set(ref) - set(TOTALS):
            print(f"note: {k} industry names differ from wacc: {sorted(set(n) ^ set(ref))[:6]}")
    print(f"wrote {DATA.name} ({DATA.stat().st_size / 1e6:.2f} MB), {CASES.name} ({len(cases)} cases); {len(notes)} notes")
    return 0


def _build_history(p: Path, data: dict, urls: dict, notes: list) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(str(p), data_only=True)
    ws = wb["Returns by year"]
    fx = openpyxl.load_workbook(str(p), data_only=False)["Returns by year"]
    cols = {"B": STOCKS, "C": SMALL, "D": BILLS, "E": BONDS, "F": BAA, "G": "Real Estate", "H": "Gold*",
            "U": "Inflation Rate"}
    for c, name in cols.items():  # guard: the source header text must still be what we label it
        assert _text(ws[f"{c}20"].value) == name, (c, ws[f"{c}20"].value)
    year_row = {ws[f"A{r}"].value: r for r in range(21, 200) if isinstance(ws[f"A{r}"].value, int)}
    rows = [[y] + [_sig(ws[f"{c}{r}"].value) for c in cols] for y, r in year_row.items()]
    data["history"] = {"file": p.name, "url": urls.get(p.name, SITE + "datasets/" + p.name),
                       "as_of": ws["B1"].value.strftime("%Y-%m-%d"), "sheet": "Returns by year",
                       "columns": ["Year", *cols.values()], "rows": rows,
                       "notes": ["Gold*: prior to 1971 gold prices were fixed (gold standard).",
                                 "Small cap = bottom decile; the sheet's 'Small Cap Premium' column R is small cap minus S&P 500 "
                                 "(its header text says the reverse)."]}
    f = p.name
    rnd = random_cells(f"{f}|annual", [(y, c) for y in year_row for c in cols.values()], 8)
    cases = [{"kind": "table", "file": f, "cell": f"{k}{year_row[y]}", "table": "history", "key": y, "column": c,
              "value": ws[f"{k}{year_row[y]}"].value}
             for y, c in rnd for k in [next(kk for kk, vv in cols.items() if vv == c)]]
    per = lambda label: tuple(int(x) for x in _text(label).split("-"))  # '1976-2025' -> (1976, 2025)
    series = {c: cols[c] for c in "BCDEFGH"}
    real = dict(zip("VWXYZ", [STOCKS, SMALL, BILLS, BONDS, BAA]), AA="Real Estate", AB="Gold*")

    def add(cell, start, end, path, err=None):
        c = {"kind": "period", "file": f, "cell": cell, "start": start, "end": end, "path": path,
             "value": ws[cell].value}
        if err:
            c["his_error"] = err
        cases.append(c)

    for r_ar, r_geo, r_prem_ar, r_prem_geo in [(121, 126, 122, 127), (122, 127, 123, 128), (123, 128, 124, 129)]:
        s, e = per(ws[f"A{r_ar}"].value)
        for c, name in series.items():
            add(f"{c}{r_ar}", s, e, ["nominal", name, "arithmetic"])
            bug = ("the sheet's column J (value of $100 in small caps) is =100*(1+C) every year, not cumulative, so the sheet's "
                   "geometric small-cap return is wrong" if c == "C" else None)
            add(f"{c}{r_geo}", s, e, ["nominal", name, "geometric"], bug)
        add(f"L{r_prem_ar}", s, e, ["premiums", "Stocks - T.Bills", "arithmetic"])
        add(f"M{r_prem_ar}", s, e, ["premiums", "Stocks - T.Bonds", "arithmetic"])
        add(f"Q{r_prem_ar}", s, e, ["premiums", "Small Cap Premium (small cap - S&P 500)", "arithmetic"])
        add(f"L{r_prem_geo}", s, e, ["premiums", "Stocks - T.Bills", "geometric"])
        add(f"M{r_prem_geo}", s, e, ["premiums", "Stocks - T.Bonds", "geometric"])
        add(f"Q{r_prem_geo}", s, e, ["premiums", "Small Cap Premium (small cap - S&P 500)", "geometric"],
            "built on the sheet's non-cumulative small-cap column J (see cell C126)")
        rs, re_ = per(ws[f"U{r_ar + 1}"].value)
        for c, name in real.items():
            add(f"{c}{r_ar + 1}", rs, re_, ["real", name, "arithmetic"])
            add(f"{c}{r_geo}", *per(ws[f"U{r_geo}"].value), ["real", name, "geometric"])
    # standard errors: read the sheet's STDEV range and sqrt(n) from the formula text
    for cell, path in [(f"{c}{r}", [grp, name, "std_error"]) for r in (122, 123, 124)
                       for c, grp, name in [("N", "premiums", "Stocks - T.Bills"), ("O", "premiums", "Stocks - T.Bonds"),
                                            ("R", "premiums", "Small Cap Premium (small cap - S&P 500)")]] + \
                      [(f"{c}{r}", ["real_premiums", name, "std_error"]) for r in (123, 124, 125)
                       for c, name in [("AH", "Stocks - T.Bills"), ("AI", "Stocks - T.Bonds")]]:
        m = re.match(r"=STDEV\([A-Z]+(\d+):[A-Z]+(\d+)\)/\(\(\$A\$(\d+)-\$A\$(\d+)\+1\)\^0\.5\)", fx[cell].value)
        r0, r1, d1, d0 = map(int, m.groups())
        y0, y1 = fx[f"A{r0}"].value, fx[f"A{r1}"].value
        n_obs, n_den = r1 - r0 + 1, ws[f"A{d1}"].value - ws[f"A{d0}"].value + 1
        add(cell, y0, y1, path, None if n_obs == n_den else
            f"the sheet's formula is STDEV of {n_obs} years ({y0}-{y1}) divided by sqrt({n_den})")
    # 'Historical ERP' column T: geometric stocks minus T.Bonds from 1928 through each year; E17 = Damodaran's estimator
    for y in (1960, 1980, 2000, 2008, 2020, 2025):  # the sheet's T column is filled from 1960 on
        add(f"T{year_row[y]}", 1928, y, ["premiums", "Stocks - T.Bonds", "geometric"])
    if ws["D9"].value == "LT":
        add("E17", ws["D10"].value, max(year_row), ["premiums", "Stocks - T.Bonds", "geometric"])
    # stale labels noted, values are live formulas
    notes.append("histretSP 'Summary for ppt' G4:G8 labels (1928-2023, 1974-2023, 2014-2023) are stale; the cells "
                 "link to the 1928-2025 / 1976-2025 / 2016-2025 real averages.")
    return cases


def random_cells(seed: str, cells: list, k: int) -> list:
    import random
    return random.Random(seed).sample(cells, min(k, len(cells)))


def _build_implied(p: Path, data: dict, urls: dict, notes: list) -> list[dict]:
    import openpyxl
    ws = openpyxl.load_workbook(str(p), data_only=True)["Historical Impl Premiums"]
    fx = openpyxl.load_workbook(str(p), data_only=False)["Historical Impl Premiums"]
    ncol = max(c for c in range(1, 40) if ws.cell(7, c).value)
    cols = [_text(ws.cell(7, c).value) for c in range(1, ncol + 1)]
    year_row = {ws.cell(r, 1).value: r for r in range(8, 200) if isinstance(ws.cell(r, 1).value, int)}
    rows = [[y] + [_num(ws.cell(r, c).value, notes, f"{p.name} row {r}") for c in range(2, ncol + 1)]
            for y, r in year_row.items()]
    last = max(year_row)
    data["implied_erp"] = {"file": p.name, "url": urls.get(p.name, SITE + "datasets/" + p.name),
                           "as_of": f"{last + 1}-01 (data through end of {last}; the sheet's 'Date updated' cell still reads "
                                    f"{ws['B1'].value:%Y-%m-%d})", "sheet": "Historical Impl Premiums",
                           "columns": cols, "rows": rows}
    f = p.name
    cases = [{"kind": "table", "file": f, "cell": f"{_col(cols.index(c))}{year_row[y]}", "table": "implied_erp",
              "key": y, "column": c, "value": ws.cell(year_row[y], cols.index(c) + 1).value}
             for y, c in random_cells(f, [(y, c) for y in year_row for c in cols[1:]
                                          if ws.cell(year_row[y], cols.index(c) + 1).value is not None], 8)]
    for r in range(70, 100):  # the sheet's period table: 'Period | ERP | ERP + Riskfree Rate'
        label = _text(ws.cell(r, 2).value)
        if re.match(r"^\d{4}-\d{4}$", label):
            s, e = map(int, label.split("-"))
            for c, field in ((3, "Implied ERP (FCFE)"), (4, "ERP + Riskfree Rate")):
                cases.append({"kind": "implied_avg", "file": f, "cell": f"{_col(c - 1)}{r}", "start": s, "end": e,
                              "field": field, "value": ws.cell(r, c).value})
        elif label.startswith("End of"):
            form = fx.cell(r, 4).value  # =C82+K73: ERP of the year in C, plus that year's T.Bond rate
            erp_row = int(re.search(r"P(\d+)", fx.cell(r, 3).value).group(1))
            y = ws.cell(erp_row, 1).value
            cases.append({"kind": "table", "file": f, "cell": f"C{r}", "table": "implied_erp", "key": y,
                          "column": "Implied ERP (FCFE)", "value": ws.cell(r, 3).value})
            k_row, c_row = int(re.search(r"K(\d+)", form).group(1)), int(re.search(r"C(\d+)", form).group(1))
            case = {"kind": "table", "file": f, "cell": f"D{r}", "table": "implied_erp", "key": y,
                    "column": "Implied ERP (FCFE)", "value": ws.cell(r, 4).value - ws.cell(k_row, 11).value}
            if c_row != r or k_row != erp_row:
                case["his_error"] = f"the sheet's {form} adds the ERP in C{c_row} to the T.Bond rate of row {k_row} (should be C{r} + K{erp_row})"
            cases.append(case)
            if label != f"End of {y}":
                notes.append(f"histimpl {_col(1)}{r} label '{label}' but the cell links to {y}")
    return cases


def _build_monthly(p: Path, data: dict, urls: dict, notes: list) -> list[dict]:
    import openpyxl
    ws = openpyxl.load_workbook(str(p), data_only=True)["Historical ERP"]
    ncol = max(c for c in range(1, 40) if ws.cell(1, c).value)
    cols = ["Month"] + [_text(ws.cell(1, c).value) for c in range(2, ncol + 1)]
    month_row = {ws.cell(r, 1).value.strftime("%Y-%m"): r for r in range(2, ws.max_row + 1)
                 if hasattr(ws.cell(r, 1).value, "strftime")}
    rows = [[mth] + [(_text(ws.cell(r, c).value) or None) if cols[c - 1] == "Notes"
                     else _num(ws.cell(r, c).value, notes, f"{p.name}!{_col(c - 1)}{r}") for c in range(2, ncol + 1)]
            for mth, r in month_row.items()]
    data["erp_monthly"] = {"file": p.name, "url": urls.get(p.name, SITE + "implprem/" + p.name),
                           "as_of": max(month_row) + "-01", "sheet": "Historical ERP", "columns": cols, "rows": rows}
    return [{"kind": "table", "file": p.name, "cell": f"{_col(c - 1)}{month_row[mth]}", "table": "erp_monthly",
             "key": mth, "column": cols[c - 1], "value": ws.cell(month_row[mth], c).value}
            for mth, c in random_cells(p.name, [(mth, c) for mth in month_row for c in range(2, ncol + 1)
                                                if isinstance(ws.cell(month_row[mth], c).value, (int, float))], 8)]


def _build_companies(p: Path, data: dict, urls: dict) -> list[dict]:
    rows = dict(_rows(p))["By industry"]
    hdr = [_text(x) for x in rows[0]]
    ci = {h: i for i, h in enumerate(hdr)}
    inds = sorted({r[ci["Industry Group"]] for r in rows[1:]})
    ctry = sorted({r[ci["Country"]] for r in rows[1:]})
    groups = {r[ci["Country"]]: [r[ci["Broad Group"]], r[ci["Sub Group"]]] for r in rows[1:]}
    body = []
    for r in rows[1:]:
        name, tick = _text(r[ci["Company Name"]]), _text(r[ci["Exchange:Ticker"]])
        if tick not in ("", "-") and f"({tick})" not in name:
            name += f" ({tick})"
        body.append([name, inds.index(r[ci["Industry Group"]]), ctry.index(r[ci["Country"]])])
    data["companies"] = {"file": p.name, "url": urls.get(p.name, SITE + "datasets/" + p.name),
                         "as_of": "2026-01 (no date cell; 48,156 firms = the Jan 2026 Global dataset count)",
                         "sheet": "By industry", "columns": ["Company Name (Exchange:Ticker)", "Industry Group", "Country"],
                         "note": "Primary Sector and SIC Code columns not stored.", "industries": inds,
                         "countries": ctry, "country_groups": groups, "rows": body}
    cases = []
    for j in random_cells(p.name, list(range(1, len(rows))), 6):
        r = rows[j]
        name = body[j - 1][0]
        for field, col in (("industry", "Industry Group"), ("country", "Country"), ("sub_group", "Sub Group")):
            cases.append({"kind": "company", "file": p.name, "cell": f"row {j + 1}", "company": name, "field": field,
                          "value": r[ci[col]]})
    return cases


# ------------------------------------------------------------------ CLI
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--build-cases", nargs="+", metavar="SPREADSHEET")
    sub = ap.add_subparsers(dest="cmd")
    sub.add_parser("datasets")
    p = sub.add_parser("columns")
    p.add_argument("dataset")
    p = sub.add_parser("industry")
    p.add_argument("name", nargs="?")
    p.add_argument("--region", default="US")
    p.add_argument("--dataset")
    p.add_argument("--metric")
    p = sub.add_parser("rank")
    p.add_argument("metric")
    p.add_argument("--region", default="US")
    p.add_argument("--dataset")
    p.add_argument("--top", type=int)
    p.add_argument("--asc", action="store_true")
    p = sub.add_parser("history")
    p.add_argument("--start", type=int, default=1928)
    p.add_argument("--end", type=int)
    p.add_argument("--year", type=int)
    p = sub.add_parser("erp")
    p.add_argument("--year", type=int)
    p.add_argument("--month")
    p.add_argument("--start")
    p.add_argument("--end")
    p.add_argument("--monthly", action="store_true")
    p = sub.add_parser("company")
    p.add_argument("query")
    for sp in sub.choices.values():  # allow --json after the subcommand too
        sp.add_argument("--json", action="store_true", default=argparse.SUPPRESS)
    a = ap.parse_args()
    if a.build_cases:
        return build(a.build_cases)
    if a.selftest:
        return selftest()
    if not a.cmd:
        ap.print_help()
        return 1
    data = load()
    try:
        out, text = run(a, data)
    except (KeyError, ValueError) as e:
        print(f"error: {e.args[0] if e.args else e}", file=sys.stderr)
        return 2
    print(json.dumps(out, indent=1, ensure_ascii=False) if a.json else text)
    return 0


def run(a, data: dict) -> tuple[dict, str]:
    if a.cmd == "datasets":
        out = {d: {r: {k: t[k] for k in ("file", "url", "as_of")} for r, t in v["regions"].items()}
               for d, v in data["datasets"].items()}
        extra = {k: {x: data[k][x] for x in ("file", "url", "as_of")} for k in ("history", "implied_erp", "erp_monthly", "companies")}
        lines = [f"{d}: {v.get('purpose', '')}\n" + "\n".join(f"   {r:9s} {m['file']:28s} as of {m['as_of']}" for r, m in rs.items())
                 for (d, rs), v in zip(out.items(), data["datasets"].values())]
        lines += [f"{k}: {m['file']} as of {m['as_of']}" for k, m in extra.items()]
        return {**out, **extra}, "\n".join(lines)
    if a.cmd == "columns":
        d = pick(a.dataset, list(data["datasets"]), "dataset")
        v = data["datasets"][d]
        cols = list(dict.fromkeys(c for t in v["regions"].values() for c in t["columns"]))
        defs = v.get("definitions", {})
        out = {"dataset": d, "purpose": v.get("purpose"), "columns": cols, "definitions": defs,
               "settings": {r: t.get("settings") for r, t in v["regions"].items()}}
        text = f"{d}: {v.get('purpose', '')}\n" + "\n".join(f"  {c}" for c in cols)
        text += "\n\nhis definitions:\n" + "\n".join(f"  {k}: {x}" for k, x in defs.items()) if defs else ""
        text += "\n\nhis settings (cells above the table):\n" + "\n".join(
            f"  {r}: {s}" for r, s in out["settings"].items() if s)
        return out, text
    if a.cmd == "industry":
        if not a.name:
            region = region_name(a.region)
            names = sorted({i for _, t in tables(data, region, a.dataset) for i in t["rows"]})
            return {"region": region, "industries": names}, "\n".join(names)
        r = industry(data, a.name, a.region, a.dataset, a.metric)
        return r, report_industry(r)
    if a.cmd == "rank":
        r = rank(data, a.metric, a.region, a.dataset, a.top, a.asc)
        return r, report_rank(r)
    if a.cmd == "history":
        if a.year:
            r = year_row(data, a.year)
            erp = r["Historical ERP (geometric stocks - T.Bonds since 1928)"]
            text = "\n".join([f"{a.year}  {src(r)}"] + [
                f"  {c:32s} {v:9.2%}   real {r['real'][c]:9.2%}" for c, v in r["nominal"].items()] +
                [f"  Inflation Rate {r['Inflation Rate']:.2%}"] +
                ([f"  Historical ERP (geometric stocks - T.Bonds, 1928-{a.year}) {erp:.2%}"] if erp is not None else []))
            return r, text
        end = a.end or data["history"]["rows"][-1][0]
        r = period_stats(data, a.start, end)
        return r, report_history(r)
    if a.cmd == "erp":
        if a.month or a.monthly:
            r = monthly_erp(data, a.month, a.start, a.end)
            return r, report_rows(r, "Implied ERP at the start of each month (S&P 500)")
        r = implied_erp(data, a.year, int(a.start) if a.start else None, int(a.end) if a.end else None)
        return r, report_rows(r, "Implied ERP by year (the sheet's value for year t is at the end of t = start of t+1)")
    if a.cmd == "company":
        r = company(data, a.query)
        text = f"{r['matches']} match(es) for '{a.query}'  {src(r)}\n" + "\n".join(
            f"  {x['company']:60s} {x['industry']:34s} {x['country']} ({x['sub_group']})" for x in r["rows"])
        return r, text
    raise ValueError(a.cmd)


if __name__ == "__main__":
    sys.exit(main())
