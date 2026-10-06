"""Build calibration cases from Damodaran's own Ginzu spreadsheets and check ginzu.py against them.

Reads Damodaran's cached inputs (Input sheet, R& D converter, mature ERP) and the cached value per share,
runs ginzu.value(), and writes data/calibration_cases.json for `ginzu.py --selftest`.
If the workbook customizes the revenue-growth or margin rows (e.g. segment models), the cached
paths are passed as revenue_growth_path / margin_path and the case is marked "custom".

Usage: python calibrate.py <xlsx> [<xlsx> ...]      (needs openpyxl)
"""

import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

from ginzu import CASES, value

warnings.filterwarnings("ignore")

FIELDS = {  # Input sheet label prefix -> input key (first match wins, in sheet order)
    "revenues": "revenues", "operating income or ebit": "ebit", "book value of equity": "book_equity",
    "book value of debt": "book_debt", "cash and marketable": "cash",
    "cross holdings": "non_operating_assets", "minority interests": "minority_interests",
    "number of shares outstanding": "shares_outstanding", "current stock price": "price",
    "effective tax rate": "effective_tax_rate", "marginal tax rate": "marginal_tax_rate",
    "revenue growth rate for next year": "growth_next_year", "operating margin for next year": "margin_next_year",
    "compounded annual revenue growth": "growth_years_2_5", "target pre-tax operating margin": "target_margin",
    "year of convergence": "margin_convergence_year", "riskfree rate": "riskfree_rate",
    "initial cost of capital": "initial_cost_of_capital",
}
# default-assumption statements (in sheet order) -> handler key
STATEMENTS = [("in stable growth", "coc"), ("i will assume that your firm will earn", "roic"),
              ("i will assume that your firm has no chance", "fail"), ("i assume that reinvestment", "lag"),
              ("i will assume that your effective tax", "tax"), ("i will assume that you have no losses", "nol"),
              ("i will asssume that today's risk free", "rf10"), ("i will assume that today's risk free", "rf10"),
              ("i will assume that the growth rate in perpetuity", "growth"), ("i have assumed that none of the cash", "trap")]


def label_rows(ws):
    for r in range(1, 130):
        a = ws.cell(r, 1).value
        if isinstance(a, str):
            yield r, a.strip().lower(), ws.cell(r, 2).value


def row_by_label(ws, prefix):
    return next((r for r, a, _ in label_rows(ws) if a.startswith(prefix)), None)


def read_case(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ins, vo = wb["Input sheet"], wb["Valuation output"]
    inp, flags, section, s2c, override = {}, {}, None, [], {}
    for r, a, b in label_rows(ins):
        key = next((k for p, k in FIELDS.items() if a.startswith(p)), None)
        if key and key not in inp and a != "revenue growth rate for next year" or key == "growth_next_year" and key not in inp:
            if key == "revenues" and a != "revenues":
                continue
            inp[key] = b
        if a.startswith("sales to capital ratio"):
            s2c.append(b)
        for p, k in [("do you have r&d", "rd"), ("do you have employee options", "options")]:
            if a.startswith(p):
                flags[k] = str(b).strip().lower() == "yes"
        for p, k in [("number of options outstanding", "number"), ("average strike", "strike"),
                     ("average maturity", "maturity"), ("standard deviation on stock", "volatility")]:
            if a.startswith(p):
                flags.setdefault("opt", {})[k] = b
        sec = next((k for p, k in STATEMENTS if a.startswith(p)), None)
        if sec:
            section = sec
        elif section and a.startswith("do you want"):
            override[section] = str(b).strip().lower() == "yes"
        elif section and override.get(section):
            if section == "coc" and a.startswith("if yes"):
                inp["terminal_cost_of_capital"] = b
            elif section == "roic" and a.startswith("if yes"):
                inp["terminal_roic"] = b
            elif section == "fail":
                if a.startswith("if yes"):
                    inp["failure_probability"] = b
                elif a.startswith("what do you want to tie"):
                    inp["failure_proceeds_basis"] = str(b).strip().upper()
                elif a.startswith("enter the distress"):
                    inp["failure_proceeds_pct"] = b
            elif section == "lag" and a.startswith("if yes"):
                inp["reinvestment_lag"] = int(b)
            elif section == "tax":
                inp["keep_effective_tax"] = True
            elif section == "nol" and a.startswith("if yes"):
                inp["nol"] = b
            elif section == "rf10" and a.startswith("if yes"):
                inp["riskfree_after_10"] = b
            elif section == "growth" and a.startswith("if yes"):
                inp["terminal_growth"] = b
            elif section == "trap":
                if a.startswith("if yes"):
                    inp["trapped_cash"] = b
                elif a.startswith("&"):
                    inp["trapped_cash_tax_rate"] = b
    std_s2c = len(s2c) == 2  # standard Ginzu; DIY/segment workbooks have 1 or one per segment
    inp["sales_to_capital_1_5"], inp["sales_to_capital_6_10"] = (s2c[0], s2c[1]) if std_s2c else (s2c[0], s2c[0])
    for r, a, b in label_rows(ins):  # IPO workbooks: proceeds retained by the firm
        if any(a.startswith(p) for p in ("offer proceeds", "expected ipo proceeds", "expected proceeds on ipo")) and b:
            inp["ipo_proceeds"] = b
            shares_f = openpyxl.load_workbook(path, data_only=False)["Valuation output"].cell(
                row_by_label(vo, "number of shares"), 2).value
            inp["ipo_shares_at_value"] = "/" in str(shares_f)
    for k in ("cash", "non_operating_assets", "minority_interests", "book_debt"):  # blank cells = 0 in Excel
        inp[k] = inp.get(k) or 0
    if "Country equity risk premiums" in wb.sheetnames:
        inp["mature_erp"] = wb["Country equity risk premiums"]["B1"].value
    if flags.get("rd") and "R& D converter" in wb.sheetnames:
        rd = wb["R& D converter"]
        life = int(rd["F6"].value)
        inp["rd"] = {"life": life, "current": rd["F7"].value,
                     "past": [rd.cell(11 + k, 2).value or 0 for k in range(life)]}
    if flags.get("options"):
        inp["options"] = flags["opt"]
        # which stock price does the option sheet use: market price or the sheet's estimated value?
        wf = openpyxl.load_workbook(path, data_only=False)["Option value"]
        refs = " ".join(str(wf.cell(r, c).value) for r in range(2, 6) for c in (2, 4))
        inp["options"]["price_basis"] = "value" if "Valuation output" in refs else "price"

    # The workbook's totals (segment workbooks sum several revenue / operating-income rows). If they differ
    # from the standard Ginzu paths, pass them as explicit paths: with a shared tax rate and
    # sales-to-capital ratio, FCFF depends only on total revenue and total operating income.
    rows = [(r, a) for r, a, _ in label_rows(vo) if r < 40]
    total = [r for r, a in rows if a.startswith("total revenues")]
    rev_rows = total or [r for r, a in rows if "revenue" in a and "growth" not in a]
    ebit_rows = [r for r, a in rows if a.startswith("ebit (operating income)") or a == "operating income"]
    col = lambda rs, c: sum(vo.cell(r, c).value or 0 for r in rs)
    rev = [col(rev_rows, c) for c in range(2, 13)]  # base year + years 1-10
    his_growth = [rev[t] / rev[t - 1] - 1 for t in range(1, 11)]
    his_margin = [col(ebit_rows, c) / col(rev_rows, c) for c in range(3, 13)]
    notes = []
    inp.setdefault("revenues", rev[0])
    if "mature_erp" not in inp or "terminal_cost_of_capital" not in inp and not std_s2c:
        inp["terminal_cost_of_capital"] = vo.cell(row_by_label(vo, "cost of capital"), 13).value
        notes.append("terminal_cost_of_capital(cached)")
    if not std_s2c:  # segment-specific sales-to-capital: take the sheet's summed reinvestment
        re_rows = [r for r, a in rows if "reinvestment" in a and "rate" not in a]
        re_rows = [r for r in re_rows if "total" in vo.cell(r, 1).value.lower()] or re_rows
        inp["reinvestment_path"] = [col(re_rows, c) for c in range(3, 13)]
        notes.append("reinvestment_path(cached)")
    if any(k not in inp for k in ("growth_next_year", "growth_years_2_5")):
        inp["revenue_growth_path"], _ = his_growth, notes.append("revenue_growth_path(no std inputs)")
    if any(k not in inp for k in ("margin_next_year", "target_margin", "margin_convergence_year")):
        inp["margin_path"], _ = his_margin, notes.append("margin_path(no std inputs)")
    ours = value(inp)["years"][:10]
    if abs(rev[0] - inp["revenues"]) > 1e-6 * abs(inp["revenues"]):
        notes.append(f"base revenue {rev[0]:.0f} vs input {inp['revenues']:.0f}")
    if any(abs(h - y["growth"]) > 1e-9 for h, y in zip(his_growth, ours)):
        inp["revenue_growth_path"], _ = his_growth, notes.append("revenue_growth_path")
    if any(abs(h - y["margin"]) > 1e-9 for h, y in zip(his_margin, ours)):
        inp["margin_path"], _ = his_margin, notes.append("margin_path")
    expected = vo.cell(row_by_label(vo, "estimated value /share"), 2).value
    case = {"name": path.stem, "custom": notes, "inputs": inp, "expected_value_per_share": expected}
    mkt_in = [r for r, a, _ in label_rows(ins) if a.startswith("total market")]
    mkt_out = [r for r, a, _ in label_rows(vo) if a.startswith("total market")]
    if mkt_in and len(mkt_in) == len(mkt_out):  # market x share segments, as in Damodaran's Nvidia 2025 model
        segs = [{"name": ins.cell(r - 1, 1).value, "market": [ins.cell(r, 2).value, vo.cell(ro, 7).value, ins.cell(r, 3).value],
                 "share": [ins.cell(r + 1, 2).value, ins.cell(r + 1, 3).value],
                 "margin": [ins.cell(r + 2, 2).value, ins.cell(r + 2, 3).value]} for r, ro in zip(mkt_in, mkt_out)]
        seg_inp = {k: v for k, v in inp.items() if k not in ("revenue_growth_path", "margin_path")}
        case["segment_case"] = {"name": path.stem + " (segments helper)", "custom": ["segments"],
                                "inputs": {**seg_inp, "segments": segs}, "expected_value_per_share": expected}
    return case


def num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def extract_data(path: Path) -> dict:
    """Damodaran's data tables embedded in a Ginzu workbook -> dict used by cost_of_capital.py."""
    wb = openpyxl.load_workbook(path, data_only=True)
    wf = openpyxl.load_workbook(path, data_only=False)
    cr = wb["Country equity risk premiums"]
    countries, regions = {}, {}
    split = next((r for r in range(5, cr.max_row + 1) if cr.cell(r, 1).value == "Region"), cr.max_row + 1)
    for r in range(5, cr.max_row + 1):
        a, b, c, d, e, f = (cr.cell(r, j).value for j in range(1, 7))
        if r < split and isinstance(a, str) and num(c) and num(d):
            countries[a.strip()] = {"rating": b, "default_spread": c, "erp": d, "crp": e, "tax_rate": f}
        elif r > split and isinstance(a, str) and num(b):  # regional ERPs (first block feeds the calculator)
            regions.setdefault(a.strip(), b)

    def industries(name):
        ws = wb[name]
        hdr = [str(ws.cell(1, j).value).strip() for j in range(1, ws.max_column + 1)]
        out = {}
        for r in range(2, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if isinstance(a, str) and num(ws.cell(r, 7).value):
                out[a.strip()] = {h: ws.cell(r, j + 1).value for j, h in enumerate(hdr[1:], start=1) if h != "None"}
        return out

    glob_name = next(s for s in wb.sheetnames if s.startswith("Industry Average") and "Global" in s)
    us_name = next(s for s in wb.sheetnames if s.startswith("Industry Average") and "US" in s)
    coc, cocf, syn = wb["Cost of capital worksheet"], wf["Cost of capital worksheet"], wb["Synthetic rating"]
    base = re.search(r"-\s*([\d.]+)%", str(cocf["B67"].value))
    hdr = next((r for r in range(1, 150) if coc.cell(r, 1).value == "Region" and coc.cell(r, 2).value == "First Decile"), None)
    if hdr:  # 2024+ layout: regions down, percentiles across
        dist = {coc.cell(r, 1).value: [coc.cell(r, j).value for j in range(2, 7)]
                for r in range(hdr + 1, hdr + 8) if coc.cell(r, 1).value and num(coc.cell(r, 2).value)}
    else:  # older layout: percentiles down ("1st Decile", ...), regions across
        hdr = next(r for r in range(1, 150) if str(coc.cell(r, 1).value).startswith("Decile/Quartile"))
        order = [("1st decile", "first decile"), ("first quartile",), ("median",), ("third quartile",),
                 ("9th decile", "ninth decile")]
        rows = {str(coc.cell(r, 1).value).lower(): r for r in range(hdr + 1, hdr + 7)}
        pick = lambda keys: next(r for k, r in rows.items() if k.startswith(keys))
        dist = {str(coc.cell(hdr, j).value).replace("Eruope", "Europe"): [coc.cell(pick(k), j).value for k in order]
                for j in range(2, 7) if coc.cell(hdr, j).value}
    table = lambda r0, r1: [[syn.cell(r, j).value for j in range(1, 5)] for r in range(r0, r1) if num(syn.cell(r, 4).value)]
    dist_cell = next(r for r in range(1, 150) if str(cocf.cell(r, 1).value).lower().startswith("cost of capital based upon decile"))
    dist_rf = re.search(r"-\s*([\d.]+)%", str(cocf.cell(dist_cell, 2).value))  # pre-2026 vintages adjust for riskfree
    return {
        "source": f"Damodaran {path.name} (pages.stern.nyu.edu/~adamodar)", "updated": str(cr["C1"].value),
        "distribution_riskfree": float(dist_rf.group(1)) / 100 if dist_rf else None,
        "mature_erp": cr["B1"].value, "industry_riskfree": float(base.group(1)) / 100 if base else None,
        "countries": countries, "regions": regions,
        "industries_us": industries(us_name), "industries_global": industries(glob_name),
        "coc_distribution": dist, "synthetic_large": table(22, 37), "synthetic_small": table(41, 56),
        "rating_spreads": {syn.cell(r, 7).value: syn.cell(r, 8).value for r in range(42, 57) if num(syn.cell(r, 8).value)},
    }


def read_coc_inputs(path: Path, data: dict) -> tuple[dict, float] | None:
    """Damodaran's cost-of-capital worksheet entries -> cost_of_capital.compute() inputs, plus the sheet's result."""
    wb = openpyxl.load_workbook(path, data_only=True)
    c, ins = wb["Cost of capital worksheet"], wb["Input sheet"]
    b = lambda p: c.cell(row_by_label(c, p), 2).value if row_by_label(c, p) else None  # layouts differ by vintage
    lab = lambda p: ins.cell(row_by_label(ins, p), 2).value
    approach = {"Detailed": "detailed", "Industry Average": "industry_average", "Distribution": "distribution",
                "I will input": "direct"}.get(b("which approach"))
    if not approach:
        return None
    ci = {"approach": approach, "riskfree_rate": b("riskfree rate"), "marginal_tax_rate": b("tax rate"),
          "cost_of_capital": b("if direct input, enter cost of capital"),
          "region": b("which grouping"), "group": b("which risk grouping") or b("what risk")}
    bt = b("industry") or ""
    ci["industry"], ci["scope"] = lab("industry (us)"), "Global" if "Global" in bt else "US"

    def block(title):  # G/H table under a title in column G, until the "Total"/"Company" row
        r0 = next(r for r in range(1, 120) if str(c.cell(r, 7).value).lower().startswith(title))
        out = []
        for r in range(r0 + 2, r0 + 16):
            name = c.cell(r, 7).value
            if str(name).strip().lower() in ("total", "company"):
                break
            out.append((name, c.cell(r, 8).value, c.cell(r, 9).value))
        return out

    beta_map = {"Single Business(US)": ("single_us", "industry (us)"),
                "Single Business(Global)": ("single_global", "industry (global)")}
    ba = b("approach for estimating beta")
    if ba in beta_map:
        beta = {"approach": beta_map[ba][0], "industry": lab(beta_map[ba][1])}
    elif ba == "Direct Input":
        beta = {"approach": "direct", "levered_beta": b("if direct input, enter levered beta")}
    else:
        us = ba == "Multibusiness(US)"
        beta = {"approach": "multi_us" if us else "multi_global",
                "businesses": [{"industry": n, "revenues": rev} for n, rev, _ in
                               block("multi business (us" if us else "multi business (global") if n and num(rev) and rev]}
    if approach == "industry_average" and bt.startswith("Multibusiness"):
        us = "US" in bt
        ci["businesses"] = [{"industry": n, "revenues": rev} for n, rev, _ in
                            block("multi business (us" if us else "multi business (global") if n and num(rev) and rev]
    ea = b("what approach do you want")
    erp = {"approach": {"Will Input": "direct", "Country of Incorporation": "country",
                        "Operating regions": "regions"}.get(ea, "countries"),
           "value": b("direct input for erp"), "country": lab("country of incorporation")}
    if erp["approach"] in ("regions", "countries"):
        known = data["regions" if erp["approach"] == "regions" else "countries"]
        rows = block("operating regions" if erp["approach"] == "regions" else "operating countries")
        lookup = lambda n: known[n]["erp"] if erp["approach"] == "countries" else known[n]
        erp["weights"] = {(str(n).strip() if n else f"unnamed {i}"):
                          rev if str(n).strip() in known and abs(lookup(str(n).strip()) - p) < 1e-9
                          else {"revenues": rev, "erp": p} for i, (n, rev, p) in enumerate(rows) if num(rev) and rev > 0}
    debt = {"book_debt": b("book value of straight debt"), "interest_expense": b("interest expense on debt"),
            "maturity": b("average maturity"),
            "approach": {"Direct Input": "direct", "Synthetic Rating": "synthetic"}.get(
                b("approach for estimating pre-tax cost of debt"), "rating"),
            "pre_tax_cost": b("if direct input, input the pre-tax"), "rating": b("if actual rating"),
            "firm_type": b("if synethetic rating") or b("if synthetic rating"),
            "ebit": wb["Synthetic rating"]["F8"].value, "country": lab("country of incorporation")}
    ci.update(shares_outstanding=b("number of shares outstanding"), price=b("current market price per share ="),
              beta=beta, erp=erp, debt=debt,
              preferred={"shares": b("number of preferred shares") or 0,
                         "price": b("current market price per share=") or 0,
                         "dividend": b("annual dividend per share") or 0})
    return ci, b("cost of capital based upon approach")


def main():
    from cost_of_capital import DATA, compute
    if len(sys.argv) < 2:
        sys.exit(__doc__)  # never overwrite the cases file with an empty set
    cases = []
    for f in sys.argv[1:]:
        wbro = openpyxl.load_workbook(f, data_only=True)
        std = "Cost of capital worksheet" in wbro.sheetnames and any(
            str(wbro["Cost of capital worksheet"].cell(r, 2).value) in ("Detailed", "Industry Average", "Distribution")
            for r in range(5, 20))
        got = None
        if std:  # DIY/IPO workbooks use a custom cost-of-capital sheet feeding the input directly
            try:
                data = extract_data(Path(f))
                got = read_coc_inputs(Path(f), data)
            except (KeyError, StopIteration, TypeError) as e:
                print(f"SKIP COC {Path(f).stem}: layout not recognized ({type(e).__name__}: {e})")
            if got:
                ci, his = got
                ours = compute(ci, data)["cost_of_capital"]
                print(f"{'ok ' if abs(ours - his) < 5e-5 else 'BAD'} COC {Path(f).stem:24s} {ci['approach']:16s} "
                      f"Damodaran {his:.4%} ours {ours:.4%}")
        c = read_case(Path(f))
        got = value(c["inputs"])["value_per_share"]
        err = abs(got / c["expected_value_per_share"] - 1)
        print(f"{'ok ' if err <= 0.005 else 'BAD'} {c['name']:28s} Damodaran {c['expected_value_per_share']:10.2f} "
              f"ours {got:10.2f} diff {err:.3%} {'custom:' + ','.join(c['custom']) if c['custom'] else ''}")
        cases.append(c)
        if "segment_case" in c:
            sc = c.pop("segment_case")
            got = value(sc["inputs"])["value_per_share"]
            err = abs(got / sc["expected_value_per_share"] - 1)
            print(f"{'ok ' if err <= 0.005 else 'BAD'} {sc['name']:28s} Damodaran {sc['expected_value_per_share']:10.2f} "
                  f"ours {got:10.2f} diff {err:.3%} custom:segments")
            cases.append(sc)
    CASES.parent.mkdir(exist_ok=True)
    CASES.write_text(json.dumps(cases, indent=1, default=str), encoding="utf-8")
    print(f"wrote {len(cases)} cases -> {CASES}")
    latest = [f for f in sys.argv[1:] if Path(f).stem.startswith("fcffsimpleginzu")]
    if latest:  # the current template carries Damodaran's latest data tables
        DATA.write_text(json.dumps(extract_data(Path(latest[0])), indent=1, default=str), encoding="utf-8")
        print(f"wrote data tables from {Path(latest[0]).name} -> {DATA}")


if __name__ == "__main__":
    main()
