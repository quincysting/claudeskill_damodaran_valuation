"""Cost of capital, transcribed from the 'Cost of capital worksheet' / 'Synthetic rating' sheets of
Damodaran's fcffsimpleginzu.xlsx. Data (country ERPs, industry betas, rating spreads, cost-of-capital
distribution) comes from data/ginzu_data.json, extracted from Damodaran's spreadsheet by calibrate.py.

Usage: python cost_of_capital.py coc_inputs.json [--data data/ginzu_data.json]
       python cost_of_capital.py --selftest        # replay data/coc_cases.json (the Ginzu template, every branch)
       python cost_of_capital.py --build-cases fcffsimpleginzu.xlsx   # pycel recalculation of the Ginzu template

Inputs (decimals; market values in the same units as shares x price):
  approach: "detailed" | "industry_average" | "distribution" | "direct"
  riskfree_rate, marginal_tax_rate
  direct:            cost_of_capital
  industry_average:  industry | businesses: [{industry, revenues}], scope "US"|"Global"
                     (Damodaran's industry cost of capital, value-weighted if multi-business, + (rf - data riskfree))
  distribution:      region "US"|"Global"|"Europe"|"Emerging"|"Japan",
                     group "First Decile"|"First Quartile"|"Median"|"Third Quartile"|"Ninth Decile"
  detailed:
    shares_outstanding, price
    beta: {approach: "single_us"|"single_global"|"multi_us"|"multi_global"|"direct",
           industry | businesses: [{industry, revenues}] | levered_beta}
    erp:  {approach: "direct"|"country"|"regions"|"countries", value | country |
           weights: {name: revenues | {revenues, erp}}}          (name = country or region)
    debt: {book_debt, interest_expense, maturity, approach: "direct"|"rating"|"synthetic",
           pre_tax_cost | rating ("A2/A") | {firm_type 1 (large) | 2 (small/risky), ebit, country}}
    preferred: {shares, price, dividend}   optional
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

DATA = Path(__file__).resolve().parents[1] / "data" / "ginzu_data.json"
CASES = DATA.parent / "coc_cases.json"
GROUPS = ["First Decile", "First Quartile", "Median", "Third Quartile", "Ninth Decile"]


def load(path: Path = DATA) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def industry(data: dict, name: str, scope: str) -> dict:
    table = data["industries_us" if scope.upper() == "US" else "industries_global"]
    if name not in table:
        raise KeyError(f"industry '{name}' not in Damodaran's {scope} table; options: {sorted(table)[:8]}...")
    return table[name]


def unlevered_beta(beta: dict, data: dict) -> float:
    a = beta["approach"]
    if a in ("single_us", "single_global"):
        return industry(data, beta["industry"], "US" if a == "single_us" else "Global")["Unlevered Beta"]
    # multi-business: weight each industry beta by estimated value = revenues x industry EV/Sales
    scope = "US" if a == "multi_us" else "Global"
    rows = [(industry(data, b["industry"], scope), b["revenues"]) for b in beta["businesses"]]
    values = [r["EV/Sales"] * rev for r, rev in rows]
    return sum(r["Unlevered Beta"] * v for (r, _), v in zip(rows, values)) / sum(values)


def erp(e: dict, data: dict) -> float:
    a = e["approach"]
    if a == "direct":
        return e["value"]
    if a == "country":
        return data["countries"][e["country"]]["erp"]
    table = data["countries"] if a == "countries" else data["regions"]
    total = sum(v["revenues"] if isinstance(v, dict) else v for v in e["weights"].values())
    out = 0.0
    for name, v in e["weights"].items():
        rev = v["revenues"] if isinstance(v, dict) else v
        p = v["erp"] if isinstance(v, dict) else (table[name]["erp"] if a == "countries" else table[name])
        out += p * rev / total
    return out


def pre_tax_cost_of_debt(d: dict, rf: float, data: dict) -> tuple[float, str]:
    a = d["approach"]
    if a == "direct":
        return d["pre_tax_cost"], "direct"
    if a == "rating":
        return rf + data["rating_spreads"][d["rating"]], d["rating"]
    # synthetic rating from interest coverage (+ country default spread, as in the source sheet)
    interest, ebit = d["interest_expense"], d["ebit"]
    cov = 1e6 if interest == 0 else (-1e5 if ebit < 0 else ebit / interest)
    table = data["synthetic_large" if d.get("firm_type", 2) == 1 else "synthetic_small"]
    lo, hi, rating, spread = next(row for row in reversed(table) if cov >= row[0])
    country = data["countries"].get(d.get("country", ""), {}).get("default_spread", 0.0)
    return rf + spread + country, f"{rating} (coverage {cov:.2f})"


def compute(ci: dict, data: dict) -> dict:
    rf, a = ci["riskfree_rate"], ci["approach"]
    if a == "direct":
        return {"cost_of_capital": ci["cost_of_capital"], "approach": a}
    if a == "industry_average":
        scope = ci.get("scope", "US")
        if ci.get("businesses"):  # multi-business: industry costs of capital weighted by revenues x EV/Sales
            rows = [(industry(data, x["industry"], scope), x["revenues"]) for x in ci["businesses"]]
            vals = [r["EV/Sales"] * rev for r, rev in rows]
            base = sum(r["Cost of capital"] * v for (r, _), v in zip(rows, vals)) / sum(vals)
        else:
            base = industry(data, ci["industry"], scope)["Cost of capital"]
        return {"cost_of_capital": base + (rf - data["industry_riskfree"]), "approach": a,
                "industry_cost_of_capital": base, "riskfree_adjustment": rf - data["industry_riskfree"]}
    if a == "distribution":
        base = data["coc_distribution"][ci["region"]][GROUPS.index(ci["group"])]
        adj = rf - data["distribution_riskfree"] if data.get("distribution_riskfree") else 0.0
        return {"cost_of_capital": base + adj, "approach": a, "distribution_value": base, "riskfree_adjustment": adj}

    t = ci["marginal_tax_rate"]
    d = ci["debt"]
    kd, rating = pre_tax_cost_of_debt(d, rf, data)
    n = d.get("maturity", 0)
    mv_debt = (d["interest_expense"] * (1 - (1 + kd) ** -n) / kd + d["book_debt"] / (1 + kd) ** n) if n else d["book_debt"]
    mv_equity = ci["shares_outstanding"] * ci["price"]
    pref = ci.get("preferred") or {}
    mv_pref = pref.get("shares", 0) * pref.get("price", 0)
    b = ci["beta"]
    bu = None if b["approach"] == "direct" else unlevered_beta(b, data)
    bl = b["levered_beta"] if bu is None else bu * (1 + (1 - t) * mv_debt / mv_equity)
    p = erp(ci["erp"], data)
    ke = rf + bl * p
    cap = mv_equity + mv_debt + mv_pref
    kp = pref["dividend"] / pref["price"] if mv_pref else 0.0
    wacc = (mv_equity * ke + mv_debt * kd * (1 - t) + mv_pref * kp) / cap
    return {"cost_of_capital": wacc, "approach": a, "unlevered_beta": bu, "levered_beta": bl, "erp": p,
            "cost_of_equity": ke, "pre_tax_cost_of_debt": kd, "rating": rating,
            "after_tax_cost_of_debt": kd * (1 - t), "market_value_debt": mv_debt,
            "market_value_equity": mv_equity, "debt_to_capital": mv_debt / cap}


def selftest() -> int:
    """Replay the Ginzu template's 'Cost of capital based upon approach' (B13) under every approach/branch."""
    data, doc = load(), json.loads(CASES.read_text(encoding="utf-8"))
    if doc["data_updated"] != data["updated"]:
        print(f"BAD cases built on data '{doc['data_updated']}', data file is '{data['updated']}': rerun --build-cases")
        return 1
    bad = 0
    for c in doc["cases"]:
        got = compute(c["inputs"], data)["cost_of_capital"]
        ok = abs(got - c["expected"]) <= 1e-6
        bad += not ok
        print(f"{'ok ' if ok else 'BAD'} {c['name']:36s} Damodaran {c['expected']:.6%}  ours {got:.6%}")
    print(f"{len(doc['cases']) - bad}/{len(doc['cases'])} cases match the Ginzu template's cost of capital (1e-6 absolute)")
    return 1 if bad else 0


def build_cases(path: str) -> None:
    """Recalculate Damodaran's fcffsimpleginzu.xlsx with pycel under each worksheet choice; only plain-value
    cells are overridden. Needs openpyxl + pycel (maintenance venv); data must be that template's tables."""
    import copy
    import warnings
    warnings.filterwarnings("ignore")
    import openpyxl
    from pycel import ExcelCompiler
    from calibrate import extract_data, read_coc_inputs
    x = Path(path)
    data = extract_data(x)
    if json.loads(json.dumps(data, default=str)) != load():
        sys.exit("data/ginzu_data.json is not this template's data: run calibrate.py on it first")
    base, _ = read_coc_inputs(x, data)
    wb = openpyxl.load_workbook(x, data_only=True)
    ws = wb["Cost of capital worksheet"]

    def block(title: str, cols: int = 2) -> list:  # the template's G/H(/I) tables, as read_coc_inputs reads them
        r0 = next(r for r in range(1, 120) if str(ws.cell(r, 7).value).lower().startswith(title))
        rows = [[ws.cell(r, 7 + k).value for k in range(cols)] for r in range(r0 + 2, r0 + 16)]
        return [row for row in rows[:next(i for i, row in enumerate(rows) if str(row[0]).lower() in ("total", "company"))]
                if row[0] and row[1]]
    us = [{"industry": n, "revenues": r} for n, r in block("multi business (us")]
    gl = [{"industry": n, "revenues": r} for n, r in block("multi business (global")]
    ctry = {n: (r if n in data["countries"] else {"revenues": r, "erp": p}) for n, r, p in block("operating countries", 3)}
    d = lambda **k: lambda ci: ci["debt"].update(k)
    scen = [("template as published", {}, lambda ci: None),
            ("debt: synthetic, firm type 2", {"B33": "Synthetic Rating", "B36": 2}, d(approach="synthetic", firm_type=2)),
            ("debt: synthetic, firm type 1", {"B33": "Synthetic Rating", "B36": 1}, d(approach="synthetic", firm_type=1)),
            ("debt: synthetic t1, EBIT 300", {"B33": "Synthetic Rating", "B36": 1, "Input sheet!B12": 300},
             d(approach="synthetic", firm_type=1, ebit=300)),
            ("debt: direct 6%", {"B33": "Direct Input", "B34": 0.06}, d(approach="direct", pre_tax_cost=0.06)),
            ("debt: rating B1/B+", {"B35": "B1/B+"}, d(rating="B1/B+")),
            ("erp: direct 6%", {"B25": "Will Input", "B26": 0.06}, lambda ci: ci["erp"].update(approach="direct", value=0.06)),
            ("erp: country of incorporation", {"B25": "Country of Incorporation"},
             lambda ci: ci["erp"].update(approach="country")),
            ("erp: operating countries", {"B25": "Operating countries"},
             lambda ci: ci["erp"].update(approach="countries", weights=ctry)),
            ("beta: direct 1.2", {"B21": "Direct Input", "B22": 1.2},
             lambda ci: ci.update(beta={"approach": "direct", "levered_beta": 1.2})),
            ("beta: single global", {"B21": "Single Business(Global)"},
             lambda ci: ci.update(beta={"approach": "single_global", "industry": wb["Input sheet"]["B9"].value})),
            ("beta: multibusiness US", {"B21": "Multibusiness(US)"},
             lambda ci: ci.update(beta={"approach": "multi_us", "businesses": us})),
            ("beta: multibusiness global", {"B21": "Multibusiness(Global)"},
             lambda ci: ci.update(beta={"approach": "multi_global", "businesses": gl})),
            ("preferred 100 sh @70, div 5", {"B48": 100}, lambda ci: ci["preferred"].update(shares=100, price=70, dividend=5)),
            ("approach: direct 9%", {"B11": "I will input"}, lambda ci: ci.update(approach="direct", cost_of_capital=0.09)),
            ("industry avg: single US", {"B11": "Industry Average", "B66": "Single Business(US)"},
             lambda ci: ci.update(approach="industry_average", scope="US")),
            ("industry avg: single global", {"B11": "Industry Average", "B66": "Single Business(Global)"},
             lambda ci: ci.update(approach="industry_average", scope="Global")),
            ("industry avg: multibusiness US", {"B11": "Industry Average", "B66": "Multibusiness(US)"},
             lambda ci: ci.update(approach="industry_average", scope="US", businesses=us)),
            ("industry avg: multibusiness global", {"B11": "Industry Average", "B66": "Multibusiness(Global)"},
             lambda ci: ci.update(approach="industry_average", scope="Global", businesses=gl)),
            ("distribution: Global third quartile", {"B11": "Distribution", "B70": "Global", "B71": "Third Quartile"},
             lambda ci: ci.update(approach="distribution", region="Global", group="Third Quartile")),
            ("distribution: US ninth decile", {"B11": "Distribution", "B70": "US", "B71": "Ninth Decile"},
             lambda ci: ci.update(approach="distribution", region="US", group="Ninth Decile"))]
    cases, sheet = [], "Cost of capital worksheet!"
    for name, cells, mod in scen:
        xl = ExcelCompiler(filename=str(x), cycles={"iterations": 200, "tolerance": 1e-12})
        xl.evaluate(sheet + "B13")
        for c, v in cells.items():
            xl.set_value(c if "!" in c else sheet + c, v)
        xl.recalculate()
        ci = copy.deepcopy(base)
        mod(ci)
        sheet_val, ours = xl.evaluate(sheet + "B13"), compute(ci, data)["cost_of_capital"]
        print(f"{'ok ' if abs(ours - sheet_val) <= 1e-6 else 'BAD'} {name:36s} Damodaran {sheet_val:.6%}  ours {ours:.6%}")
        cases.append({"name": name, "cells": cells, "inputs": ci, "expected": sheet_val})
    CASES.write_text(json.dumps({"source": f"{x.name}: 'Cost of capital worksheet'!B13 recalculated (pycel) per "
                                 "override; 'template as published' equals the template's cached cell",
                                 "url": "https://pages.stern.nyu.edu/~adamodar/pc/fcffsimpleginzu.xlsx",
                                 "data_updated": data["updated"], "cases": cases}, indent=1, default=str), encoding="utf-8")
    print(f"wrote {len(cases)} cases -> {CASES}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("inputs", nargs="?")
    ap.add_argument("--data", default=str(DATA))
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--build-cases", metavar="XLSX")
    a = ap.parse_args()
    if a.selftest:
        sys.exit(selftest())
    if a.build_cases:
        return build_cases(a.build_cases)
    if not a.inputs:
        ap.error("inputs JSON required")
    data = load(Path(a.data))
    ci = {k: v for k, v in json.loads(Path(a.inputs).read_text(encoding="utf-8")).items() if not k.startswith("_")}
    out = compute(ci, data)
    print(f"data: {data['source']} (updated {data['updated']})")
    for k, v in out.items():
        print(f"  {k:24s} {v:.4%}" if isinstance(v, float) and abs(v) < 5 and k != "levered_beta" and k != "unlevered_beta"
              else f"  {k:24s} {v:,.4f}" if isinstance(v, float) else f"  {k:24s} {v}")


if __name__ == "__main__":
    main()
