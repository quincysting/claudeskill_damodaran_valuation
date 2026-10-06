"""Shared styling and sheet helpers for the Excel export (investment-banking conventions:
blue = hard-coded input, black = formula, green = link to another sheet)."""
import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- style (blue input, black formula, green link)
NAVY, LIGHT = "0B2545", "DCE6F1"
F = "Arial"


def font(size=10, **k):
    return Font(name=F, size=size, **k)


FONT, BLUE, GREEN, BOLD = font(), font(color="0000FF"), font(color="008000"), font(bold=True)
WHITE_B, TITLE, SUB = font(bold=True, color="FFFFFF"), font(16, bold=True, color=NAVY), font(italic=True, color="595959")
NOTE, LINK = font(8, italic=True, color="595959"), font(color="1F4E79", underline="single")
FILL_NAVY, FILL_LIGHT, FILL_FLEX = (PatternFill("solid", fgColor=c) for c in (NAVY, LIGHT, "FFF2CC"))
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
THIN = Side(style="thin", color="808080")
TOP, BOX = Border(top=THIN), Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
M, PCT, PCT2 = '#,##0_);(#,##0);"–"_)', '0.0%_);(0.0%);"–"_)', '0.00%_);(0.00%);"–"_)'
X, NUM2, NUM4 = '0.00"x"', "0.00_);(0.00)", "0.0000"
LINK_RE = re.compile(r"=('?[\w ]+'?!)?\$?[A-Z]{1,3}\$?\d+")
ITER = 60  # unrolled fixed-point steps on the Dilution tab (converges to < 1e-12 well before this)


def col(t: int) -> str:
    """DCF column for year t: 0 = base (C), 1..10 = D..M, 11 = terminal (N)."""
    return get_column_letter(3 + t)


def title(ws, text, sub):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.column_dimensions["A"].width = 2
    ws["B2"], ws["B3"] = text, sub
    ws["B2"].font, ws["B3"].font = TITLE, SUB


def section(ws, row, text, first="B", last="N"):
    for c in range(ord(first), ord(last) + 1):
        ws[f"{chr(c)}{row}"].fill = FILL_NAVY
    ws[f"{first}{row}"] = text
    ws[f"{first}{row}"].font = WHITE_B


def hdr(ws, row, labels: dict):
    for ref, text in labels.items():
        c = ws[f"{ref}{row}"]
        c.value, c.font, c.fill = text, BOLD, FILL_LIGHT
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")


def put(ws, ref, v, fmt=None, kind=None, bold=False):
    """kind: None = auto (formula -> black, single cross-sheet reference -> green, constant -> blue), or 'label'."""
    c = ws[ref]
    c.value = v
    if kind is None:
        kind = ("link" if "!" in v and LINK_RE.fullmatch(v) else "calc") if isinstance(v, str) and v.startswith("=") else "input"
    c.font = {"input": BLUE, "link": GREEN, "calc": FONT, "label": FONT}[kind]
    if bold:
        c.font = font(bold=True, color=c.font.color)
    if fmt:
        c.number_format = fmt
    return c


def label(ws, ref, text, bold=False):
    return put(ws, ref, text, kind="label", bold=bold)


def widths(ws, spec):
    for k, w in spec.items():
        ws.column_dimensions[k].width = w


def printable(ws, landscape=True, one_page=False):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1 if one_page else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


# flex cells on the Sensitivity tab (the data tables vary them; neutral = base case)
FX = {k: f"Sensitivity!$C${r}" for r, k in enumerate(("d_init", "d_term", "d_growth", "d_margin", "d_rf"), start=5)}
