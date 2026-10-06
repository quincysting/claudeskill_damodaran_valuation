"""Country risk premiums (CRP) and equity risk premiums (ERP) by country, transcribed from Damodaran's
ctryprem spreadsheets (https://pages.stern.nyu.edu/~adamodar/pc/datasets/ctryprem.xlsx and the dated
ctryprem<Mon><YY>.xlsx vintages) and checked against every country in six of them (--selftest).

Damodaran's pipeline (sheet -> what this file transcribes):
  'Sovereign Ratings'       rating = Moody's local-currency rating; if Moody's does not rate the country,
                            the S&P rating mapped notch-for-notch to Moody's          -> moodys_equivalent()
  'Default Spreads for Ratings'  spread(rating) = previous update's spread x (1 + median % change in sovereign
                            CDS across countries quoted at both dates)                  -> update_spreads()
  'Relative Equity Volatility'   scalar = avg annual std dev (EM equity index, daily) / avg annual std dev
                            (EM sovereign bond ETF, daily), 5 annual windows            -> relative_volatility()
  '10-year CDS Spreads'     CDS net of a benchmark (US until 2025; Switzerland after the May-2025 US
                            downgrade), floored at 0                                   -> net_cds()
  'ERPs by country'         CRP = spread x scalar (cell "adjust for equity volatility" = Yes), ERP = mature
                            ERP + CRP; same with net CDS (CDS columns are always scaled). From Jan 2026 the US
                            row is the US implied ERP and its CRP is the unscaled Aa1 spread -> country_row()
  'Summary' / FAQ (post-May 2025)  mature ERP = implied S&P 500 ERP (over T.Bond - US default spread)
                            - US default spread (unscaled; source files, e.g. 4.42% - 0.22% = 4.20%) -> mature_erp()
  'PRS Worksheet'           unrated countries: PRS composite score -> rating band -> ERP of that band
                                                                                        -> prs_premium()
  'Regional Weighted Averages'  GDP-weighted spread/CRP/tax; ERP = mature + weighted CRP -> weighted()
Company use (the Ginzu 'countries' ERP option and Damodaran's posts): revenue-weighted average of country ERPs
(production weights for natural-resource firms) -> company_erp(); local riskfree = govt bond - default spread.

Usage:
  python country_risk.py inputs.json [--json]     inputs: see below
  python country_risk.py --country India [--as-of 2025-01-15]
  python country_risk.py --selftest               replay data/country_risk_cases.json (Damodaran's files)
  python country_risk.py --build-cases [files]    re-extract data/country_risk_cases.json + _data.json from the
                                                  xlsx (maintenance venv: openpyxl; pycel for the E5=No branch)

inputs.json (any combination; rates are decimals):
  {"country": "India", "as_of": "2025-07-01"}                           lookup (latest if as_of omitted)
  {"premium": {"rating": "Baa3" | "default_spread": 0.0175 | "cds": 0.0085, "cds_benchmark": 0.0013,
               "mature_erp": 0.042, "scalar": 1.5545, "scale": true}}  compute; omitted -> latest vintage
  {"weights": {"United States": 600, "China": 150, "Western Europe": 250}, "basis": "revenues"}
                                     revenue-weighted ERP; names may be countries or Damodaran's regions/groupings
  {"riskfree": {"govt_bond": 0.0632, "country": "India"}}              local-currency riskfree rate
  {"tax": "Brazil"}                                                    marginal corporate tax rate
"""

from __future__ import annotations

import argparse
import difflib
import json
import math
import re
import statistics
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
DATA = HERE.parent / "data" / "country_risk_data.json"
CASES = HERE.parent / "data" / "country_risk_cases.json"
URL = "https://pages.stern.nyu.edu/~adamodar/pc/datasets/"
NOT_RATED = {None, "", "NR", "NA", "#N/A"}
ALIASES = {"us": "United States", "usa": "United States", "u.s.": "United States", "uk": "United Kingdom"}


def num(x) -> bool:
    return isinstance(x, (int, float)) and not isinstance(x, bool)


# ---------------------------------------------------------------- spreadsheet formulas
def moodys_equivalent(moodys: str | None, sp: str | None, sp_map: dict, fitch: str | None = None) -> str:
    """Moody's local-currency rating; if Moody's does not rate, the S&P rating mapped to Moody's; if S&P does
    not rate either, the Fitch rating (the sheet's Nepal row, 2026 files)."""
    if moodys not in NOT_RATED:
        return moodys
    return sp_map.get(sp if sp not in NOT_RATED or fitch in NOT_RATED else fitch, "NA")


def update_spreads(prior_bp: dict, cds_pairs: list) -> tuple[dict, float]:
    """Spread table (basis points) = prior table x (1 + median CDS % change); pairs = [cds_then, cds_now]."""
    med = statistics.median(now / then - 1 for then, now in cds_pairs if num(then) and num(now))
    return {r: (s * (1 + med) if num(s) else None) for r, s in prior_bp.items()}, med


def net_cds(cds, benchmark: float):
    """Sovereign CDS net of the benchmark's CDS, floored at 0 (IF(C>bench, C-bench, 0))."""
    return None if not num(cds) else (cds - benchmark if cds > benchmark else 0.0)


def annual_std(prices: list, lo: int, hi: int, factor: float, first_row: int) -> float:
    """STDEV of daily returns P[k]/P[k-1]-1 over sheet rows lo..hi, x sqrt(factor)."""
    rets = [prices[k - first_row] / prices[k - first_row - 1] - 1
            for k in range(max(lo, first_row + 1), hi + 1) if k - first_row < len(prices)]
    return statistics.stdev(rets) * math.sqrt(factor)


def relative_volatility(vol: dict) -> tuple[float, list]:
    """Ratio of the average annualized equity std dev to the average annualized bond std dev."""
    rows = [(annual_std(vol["equity"], *w[0:3], vol["first_row"]), annual_std(vol["bond"], *w[3:6], vol["first_row"]))
            for w in vol["windows"]]
    return statistics.fmean(e for e, _ in rows) / statistics.fmean(b for _, b in rows), rows


def mature_erp(us_implied_erp: float, us_default_spread: float) -> float:
    """Post-May-2025 mature market premium = US implied ERP - US (Aa1) default spread."""
    return us_implied_erp - us_default_spread


def premiums(spread: float, scalar: float, mature: float, scale: bool = True) -> tuple[float, float]:
    """(CRP, ERP): CRP = default spread x relative equity volatility (or the spread, if not scaled)."""
    crp = spread * scalar if scale else spread
    return crp, mature + crp


def country_row(name: str, rating: str, cds_net, p: dict) -> dict:
    """One row of 'ERPs by country'. p: spreads (bp), scalar, mature_erp, scale, us_erp, us_country."""
    s = p["spreads"].get(rating)
    spread = s / 10000 if num(s) else None
    if spread is None:
        crp = erp = None
    elif name == p.get("us_country"):  # Jan 2026 on: US ERP typed in, CRP = unscaled spread
        crp, erp = spread, p["us_erp"]
    else:
        crp, erp = premiums(spread, p["scalar"], p["mature_erp"], p.get("scale", True))
    # CDS columns ignore the Yes/No switch in the source sheet (I = G x E6); to leave CDS unscaled, set scalar = 1
    crp_cds, erp_cds = premiums(cds_net, p["scalar"], p["mature_erp"]) if num(cds_net) else (None, None)
    return {"rating": rating, "spread": spread, "crp": crp, "erp": erp,
            "cds_net": cds_net, "crp_cds": crp_cds, "erp_cds": erp_cds}


def prs_premium(score: float, bands: list, spreads_bp: dict, scalar: float, mature: float, override=None) -> dict:
    """Unrated country: ERP of the rating band [lower bound, rating] its PRS score falls in (approximate
    VLOOKUP on lower bounds); CRP = ERP - mature; implied default spread = CRP / scalar."""
    if num(override) and override > 0:
        erp, rating = override, None
    else:
        lower, rating = max((b for b in bands if b[0] <= score), key=lambda b: b[0])
        erp = mature + spreads_bp[rating] / 10000 * scalar
    return {"band_rating": rating, "erp": erp, "crp": erp - mature, "spread": (erp - mature) / scalar}


def weighted(rows: list, fields=("spread", "crp", "tax")) -> dict:
    """rows = [(weight, {field: value})] -> weight-averaged fields (the sheet's SUM(weight/total x value))."""
    total = sum(w for w, _ in rows)
    return {f: sum(w / total * r[f] for w, r in rows) for f in fields}


def region_premiums(members: list, mature: float, erp_rule: str = "mature_plus_crp") -> dict:
    """'Regional Weighted Averages': GDP weights; ERP = mature + weighted CRP (the July 2024 file: weighted ERP)."""
    out = weighted([(m["gdp"], m) for m in members], ("spread", "crp", "tax", "erp"))
    return {**out, "erp": out["crp"] + mature if erp_rule == "mature_plus_crp" else out["erp"],
            "gdp": sum(m["gdp"] for m in members)}


def local_riskfree(govt_bond: float, default_spread: float) -> float:
    """Riskfree rate in a currency = local-currency government bond rate - sovereign default spread."""
    return govt_bond - default_spread


# ---------------------------------------------------------------- data lookup
def load() -> dict:
    return json.loads(DATA.read_text(encoding="utf-8"))


def pick_vintage(data: dict, as_of: str | None = None) -> tuple[str, dict]:
    keys = sorted(data["vintages"])
    ok = [k for k in keys if as_of is None or k <= as_of]
    if not ok:
        raise SystemExit(f"no vintage on or before {as_of}; earliest is {keys[0]} "
                         "(earlier numbers: see Damodaran's posts in references/country-risk.md)")
    return ok[-1], data["vintages"][ok[-1]]


def match(name: str, names) -> str:
    names = list(names)
    q = ALIASES.get(name.strip().lower(), name).strip().lower()
    low = {n.lower(): n for n in names}
    if q in low:
        return low[q]
    hits = [n for n in names if q in n.lower() or n.lower() in q]
    if len(hits) == 1:
        return hits[0]
    close = difflib.get_close_matches(q, low, n=1, cutoff=0.75)
    if close:
        return low[close[0]]
    raise KeyError(f"'{name}' not found; similar: {difflib.get_close_matches(q, low, n=5, cutoff=0.4) or hits[:5]}")


def lookup(data: dict, country: str, as_of: str | None = None) -> dict:
    key, v = pick_vintage(data, as_of)
    name = match(country, v["countries"])
    return {"vintage": key, "file": v["file"], "country": name, "mature_erp": v["mature_erp"],
            "scalar": v["scalar"], **v["countries"][name]}


def premium_from(data: dict, q: dict) -> dict:
    """Compute CRP/ERP Damodaran's way from a rating, a default spread or a CDS spread; defaults = chosen vintage."""
    key, v = pick_vintage(data, q.get("as_of"))
    mature, scalar = q.get("mature_erp", v["mature_erp"]), q.get("scalar", v["scalar"])
    if "rating" in q:
        spread, basis = v["spreads_bp"][q["rating"]] / 10000, f"rating {q['rating']} ({key} table)"
    elif "default_spread" in q:
        spread, basis = q["default_spread"], "default spread given"
    else:
        bench = q.get("cds_benchmark", v["cds_benchmark_value"])
        spread, basis = net_cds(q["cds"], bench), f"CDS net of {bench:.4%} ({v['cds_benchmark']} on {key})"
    crp, erp = premiums(spread, scalar, mature, q.get("scale", True))
    return {"vintage": key, "basis": basis, "default_spread": spread, "scalar": scalar, "mature_erp": mature,
            "crp": crp, "erp": erp}


def company_erp(data: dict, weights: dict, as_of: str | None = None) -> dict:
    """Revenue- (or production-) weighted ERP: sum(ERP_i x w_i)/sum(w_i), as in the Ginzu 'countries' option.
    Names may be countries, Damodaran's GDP-weighted regions or regional groupings."""
    key, v = pick_vintage(data, as_of)
    pool = {**v["groupings"], **v["regions"], **v["countries"]}
    rows = []
    for n, w in weights.items():
        m = match(n, pool)
        rows.append((m, w, pool[m]["erp"], pool[m].get("crp", pool[m]["erp"] - v["mature_erp"])))
    total = sum(w for _, w, _, _ in rows)
    return {"vintage": key, "rows": [{"name": m, "weight": w / total, "erp": e, "crp": c} for m, w, e, c in rows],
            "erp": sum(w * e for _, w, e, _ in rows) / total, "crp": sum(w * c for _, w, _, c in rows) / total}


def tax_rate(data: dict, country: str, as_of: str | None = None) -> dict:
    t = data["tax"]
    out = {}
    if as_of is None or as_of >= t["as_of"]:
        name = match(country, t["rates"])
        out = {"country": name, "marginal": t["rates"][name][0], "with_global_min_tax": t["rates"][name][1],
               "as_of": t["as_of"], "source": t["file"]}
    else:
        key, v = pick_vintage(data, as_of)
        name = match(country, v["countries"])
        out = {"country": name, "marginal": v["countries"][name].get("tax"), "as_of": key, "source": v["file"]}
    e = data.get("effective_tax", {})
    try:
        en = match(country, e.get("rates", {}))
        out["effective_2025"] = dict(zip(e["fields"], e["rates"][en]))
    except KeyError:
        pass
    return out


# ---------------------------------------------------------------- report / CLI
def pct(x):
    return "NA" if x is None else f"{x:.2%}"


# the source file's cell kept as-is; Damodaran's post gives a different number (see references/tool-country-risk.md, step 5)
POST_NOTES = {("2025-07-01", "United States"): "Damodaran's post gives a US ERP of 4.48% (implied ERP vs the default-free "
              "rate); this file's US row is mature ERP + scaled Aa1 CRP"}


def report(q: dict, data: dict) -> tuple[list, dict]:
    lines, out = [], {}
    if "country" in q:
        r = out["lookup"] = lookup(data, q["country"], q.get("as_of"))
        lines += [f"{r['country']} ({r['vintage']}, {r['file']}; mature ERP {pct(r['mature_erp'])}, "
                  f"scalar {r['scalar']:.4f})",
                  f"  region {r.get('region')} | basis {r.get('basis', 'rating')} | rating {r.get('rating')}"
                  + (f" | PRS {r['prs']}" if 'prs' in r else ""),
                  f"  rating-based: default spread {pct(r.get('spread'))}  CRP {pct(r.get('crp'))}  ERP {pct(r.get('erp'))}",
                  f"  CDS-based:    CDS {pct(r.get('cds'))} net {pct(r.get('cds_net'))}  CRP {pct(r.get('crp_cds'))}  "
                  f"ERP {pct(r.get('erp_cds'))}",
                  f"  marginal tax rate {pct(r.get('tax'))}"]
        if (r["vintage"], r["country"]) in POST_NOTES:
            lines.append(f"  caution: {POST_NOTES[r['vintage'], r['country']]}")
    if "premium" in q:
        r = out["premium"] = premium_from(data, q["premium"])
        lines += [f"Premium from {r['basis']}: spread {pct(r['default_spread'])} x scalar {r['scalar']:.4f} = CRP "
                  f"{pct(r['crp'])}; ERP = {pct(r['mature_erp'])} + CRP = {pct(r['erp'])}"]
    if "weights" in q:
        r = out["weights"] = company_erp(data, q["weights"], q.get("as_of"))
        lines += [f"Weighted ERP ({q.get('basis', 'revenues')}, {r['vintage']}): {pct(r['erp'])}  (CRP {pct(r['crp'])})"]
        lines += [f"  {x['name']:32s} w {x['weight']:6.1%}  ERP {pct(x['erp'])}" for x in r["rows"]]
        lines += [f"  caution: {POST_NOTES[r['vintage'], x['name']]}" for x in r["rows"]
                  if (r["vintage"], x["name"]) in POST_NOTES]
    if "riskfree" in q:
        f = q["riskfree"]
        s = f.get("default_spread")
        if s is None:
            s = lookup(data, f["country"], f.get("as_of"))["spread"]
        rf = out["riskfree"] = local_riskfree(f["govt_bond"], s)
        lines += [f"Local riskfree = govt bond {pct(f['govt_bond'])} - default spread {pct(s)} = {pct(rf)}"]
    if "tax" in q:
        r = out["tax"] = tax_rate(data, q["tax"], q.get("as_of"))
        lines += [f"Tax {r['country']}: marginal {pct(r['marginal'])} ({r['as_of']}, {r['source']})"
                  + (f", with global min tax {pct(r['with_global_min_tax'])}" if "with_global_min_tax" in r else "")
                  + (f"; 2025 effective (aggregate) {pct(r['effective_2025']['effective_aggregate'])}"
                     if "effective_2025" in r else "")]
    return lines, out


# ---------------------------------------------------------------- selftest
def close(a, b) -> bool:
    if not (num(a) and num(b)):
        return a == b or (a in NOT_RATED and b in NOT_RATED)
    return abs(a - b) <= 1e-6 or abs(a - b) <= 5e-4 * abs(b)


def replay(c: dict) -> tuple[list, dict]:
    """Recompute one vintage from the source inputs; return [(check, his, ours)] mismatches and counts."""
    i, e = c["inputs"], c["expected"]
    bad, n = [], {}

    def chk(group, label, his, ours):
        n[group] = n.get(group, 0) + 1
        if not close(ours, his):
            bad.append((group, label, his, ours))

    scalar, _ = relative_volatility(i["vol"])
    chk("scalar", "relative volatility", e["scalar"], scalar)
    table, med = update_spreads(i["prior_spreads_bp"], i["cds_pairs"])
    table.update(i.get("spread_overrides_bp", {}))
    chk("spreads", "median CDS change", e["median_cds_change"], med)
    for r, s in e["spreads_bp"].items():
        chk("spreads", r, s, table.get(r))
    p = {"spreads": table, "scalar": scalar, "mature_erp": i["mature_erp"], "scale": i["scale"],
         "us_erp": i.get("us_erp"), "us_country": i.get("us_country")}
    for name, moodys, sp, fitch, cds in i["countries"]:
        rating = moodys_equivalent(moodys, sp, i["sp_to_moodys"], fitch)
        row = country_row(name, rating, net_cds(cds, i["cds_benchmark"]), p)
        for f, his in e["countries"][name].items():
            chk("countries", f"{name}.{f}", his, row[f])
    for name, score, override in i["prs"]["countries"]:
        row = prs_premium(score, i["prs"]["bands"], table, scalar, i["mature_erp"], override)
        for f, his in e["prs"][name].items():
            chk("prs", f"{name}.{f}", his, row[f])
    regs = {}
    for reg, members in i["regions"].items():
        regs[reg] = region_premiums([dict(zip(("gdp", "spread", "erp", "crp", "tax"), m[1:])) for m in members],
                                    i["mature_erp"], i["region_erp_rule"])
        for f, his in e["regions"][reg].items():
            chk("regions", f"{reg}.{f}", his, regs[reg][f])
    if "Global" in e:
        glob = weighted([(r["gdp"], r) for r in regs.values()], ("spread", "crp", "tax", "erp"))
        for f, his in e["Global"].items():
            chk("regions", f"Global.{f}", his, glob[f])
    if i.get("us_erp") is not None:  # the sheet's mature ERP = US ERP - US default spread, typed to 0.01%
        us = e["countries"][i["us_country"]]["spread"]
        chk("mature", "US ERP - US spread (rounded 0.01%)", i["mature_erp"], round(mature_erp(i["us_erp"], us), 4))
    return bad, n


def replay_branch(c: dict) -> tuple[list, dict]:
    """Damodaran's template recalculated by pycel with the 'scale the default spread' switch set to No."""
    p = {**c["params"], "scale": False}
    bad = [("branch", f"{name}.{f}", his, country_row(name, rating, cds_net, p)[f])
           for name, rating, cds_net, exp in c["rows"] for f, his in exp.items()
           if not close(country_row(name, rating, cds_net, p)[f], his)]
    return bad, {"branch": sum(len(r[3]) for r in c["rows"])}


def selftest() -> int:
    cases = json.loads(CASES.read_text(encoding="utf-8"))
    total_bad = 0
    for c in cases:
        bad, n = replay_branch(c) if "rows" in c else replay(c)
        total_bad += bool(bad)
        counts = ", ".join(f"{g} {n[g] - sum(b[0] == g for b in bad)}/{n[g]}" for g in n)
        print(f"{'ok ' if not bad else 'BAD'} {c['name']:16s} {counts}")
        for g, label, his, ours in bad[:15]:
            print(f"      {g}: {label:40s} Damodaran {his!r:>24}  ours {ours!r}")
    print(f"{len(cases) - total_bad}/{len(cases)} cases match Damodaran's spreadsheets on every check")
    return 1 if total_bad else 0


# ---------------------------------------------------------------- build (only place openpyxl is used)
def build(files: list[Path]) -> None:
    import datetime as dt
    import warnings

    import openpyxl
    warnings.filterwarnings("ignore")
    r8 = lambda x: round(x, 10) if num(x) else (None if x in (None, "NA", "#N/A") else x)
    cases, vintages, notes = [], {}, []
    for f in sorted(files, key=lambda p: p.name):
        if not f.name.lower().startswith("ctryprem"):
            continue
        wv, wf = openpyxl.load_workbook(f, data_only=True), openpyxl.load_workbook(f)
        case, v = read_vintage(f, wv, wf, r8, notes)
        cases.append(case)
        vintages[v["date"]] = v
    tax = next((read_tax(p, r8) for p in files if p.name.lower() == "countrytaxrates.xlsx"), None)
    eff = next((read_effective_tax(p, r8) for p in files if p.name.lower() == "countrytaxrates2025.xlsx"), None)
    latest_file = next(p for p in files if p.name == vintages[max(vintages)]["file"])
    try:
        cases.append(branch_case(latest_file, vintages[max(vintages)]))
    except ImportError:
        notes.append("pycel not installed: branch case (scale switch = No) not rebuilt")
    rc = next((p for p in files if p.name.lower() == "ratings&cds2023.xlsx"), None)
    if rc and "2023-07-01" in vintages:  # the July 2023 companion file: same spread table?
        ws = openpyxl.load_workbook(rc, data_only=True)["Default Spreads"]
        t = vintages["2023-07-01"]["spreads_bp"]
        diff = [r[1].value for r in ws.iter_rows(min_row=2) if r[1].value in t and not close(r[2].value, t[r[1].value] / 1e4)]
        notes.append(f"Ratings&CDS2023 'Default Spreads' vs ctrypremJuly23 table: {'identical' if not diff else diff}")
    latest = max(vintages)
    data = {"source": "Damodaran ctryprem spreadsheets (country risk premiums), countrytaxrates.xlsx, "
                      "CountryTaxRates2025.xlsx", "url": URL + "ctryprem.xlsx", "as_of": latest,
            "note": "rates as decimals; spreads_bp in basis points; 'basis' PRS = unrated, ERP from PRS score band",
            "vintages": {k: {kk: vv for kk, vv in v.items() if kk != "date"} for k, v in sorted(vintages.items())},
            "tax": tax, "effective_tax": eff}
    DATA.write_text(json.dumps(data, indent=0, ensure_ascii=False), encoding="utf-8")
    CASES.write_text(json.dumps(cases, separators=(",", ":"), ensure_ascii=False), encoding="utf-8")
    print(f"wrote {DATA.name} ({len(vintages)} vintages) and {CASES.name} ({len(cases)} cases)")
    for n in notes:
        print("  note:", n)


def read_vintage(f, wv, wf, r8, notes):
    import datetime as dt
    ev, ef = wv["ERPs by country"], wf["ERPs by country"]
    hdr = next(r for r in range(1, 15) if ev.cell(r, 1).value == "Country")
    label = {str(ev.cell(r, 1).value).lower(): r for r in range(1, hdr)}
    row_of = lambda s: next(r for t, r in label.items() if s in t)
    date = ev["B2"].value.date().isoformat()
    mature = ev.cell(row_of("mature equity market"), 5).value
    us_r = next((r for t, r in label.items() if "premium for the us" in t), None)
    us_erp = ev.cell(us_r, 5).value if us_r else None
    scale = ev.cell(row_of("adjust the country default spread"), 5).value == "Yes"
    scalar = ev.cell(row_of("multiplier"), 5).value
    # rating -> spread table used by the ERP sheet (J:K), and the typed overrides in the helper table below it
    jk = {}
    for r in range(hdr + 2, hdr + 26):
        if ev.cell(r, 10).value:
            jk[str(ev.cell(r, 10).value).strip()] = ev.cell(r, 11).value
    overrides = {}
    for r in range(150, ef.max_row + 1):
        b, c = ef.cell(r, 2).value, ef.cell(r, 3).value
        if isinstance(b, str) and b in jk and num(c):
            overrides[b] = c
    # spread update sheet
    ds = wv["Default Spreads for Ratings"]
    prior = {ds.cell(r, 1).value: (ds.cell(r, 2).value if num(ds.cell(r, 2).value) else None) for r in range(2, 23)
             if ds.cell(r, 1).value}
    dsf = wf["Default Spreads for Ratings"]
    overrides.update({dsf.cell(r, 1).value: dsf.cell(r, 3).value for r in range(2, 23) if num(dsf.cell(r, 3).value)})
    pairs = [[ds.cell(r, 7).value, ds.cell(r, 8).value] for r in range(2, 83) if ds.cell(r, 5).value]
    med = ds["I84"].value
    # ratings
    rw = wv["Ratings worksheet"]
    raw = {rw.cell(r, 1).value: (rw.cell(r, 2).value, rw.cell(r, 3).value) for r in range(2, rw.max_row + 1) if rw.cell(r, 1).value}
    sr, srf = wv["Sovereign Ratings (Moody's,S&P)"], wf["Sovereign Ratings (Moody's,S&P)"]
    col = next(c for c in range(5, 10) if str(sr.cell(8, c).value or "").startswith("S&P"))
    ratings = {}
    unadj = str(sr["E1"].value).startswith("Moody's Rating (unadj)")  # 2025+ layout: B S&P, C Fitch, E Moody's
    for r in range(2, 160):
        n = sr.cell(r, 1).value
        if not n:
            continue
        if unadj:
            ratings[n] = (sr.cell(r, 5).value, sr.cell(r, 2).value, sr.cell(r, 3).value)
        else:  # 2023-24 layout: C = IF(Moody's="NA", map(S&P), Moody's); a typed C is a manual rating
            typed = not str(srf.cell(r, 3).value).startswith("=")
            ratings[n] = (srf.cell(r, 3).value if typed else raw.get(n, (None, None))[1], sr.cell(r, 2).value, None)
    sp_map = {sr.cell(r, col).value: sr.cell(r, col + 1).value for r in range(9, 36) if sr.cell(r, col).value}
    # CDS: benchmark cell from the net-CDS formula
    cs_f, cs_v = wf["10-year CDS Spreads"], wv["10-year CDS Spreads"]
    brow = int(re.search(r"\$C\$(\d+)", cs_f["D2"].value).group(1))
    bench, bench_name = cs_v.cell(brow, 3).value, cs_v.cell(brow, 1).value
    cds = {cs_v.cell(r, 1).value: cs_v.cell(r, 3).value for r in range(2, 160) if cs_v.cell(r, 1).value}
    # country rows
    countries, expected, data_c = [], {}, {}
    us_country = None
    region_of = {wv["Regional lookup table"].cell(r, 1).value: wv["Regional lookup table"].cell(r, 2).value
                 for r in range(2, 200)}
    taxes = {wv["Country Tax Rates"].cell(r, 1).value: wv["Country Tax Rates"].cell(r, 2).value for r in range(2, 230)}
    r = hdr + 1
    while isinstance(ev.cell(r, 1).value, str) and not ev.cell(r, 1).value.startswith("Frontier"):
        name = ev.cell(r, 1).value
        vals = [ev.cell(r, c).value for c in range(2, 10)]
        if ef.cell(r, 5).value == "=E4":
            us_country = name
        countries.append([name, *ratings[name], cds.get(name)])
        keys = ("rating", "spread", "erp", "crp", "cds_net", "erp_cds", "crp_cds")
        expected[name] = {k: r8(x) for k, x in zip(keys, vals[1:])}
        data_c[name] = {"region": vals[0], **expected[name], "cds": r8(cds.get(name)), "tax": r8(taxes.get(name)),
                        "basis": "rating"}
        r += 1
    # PRS frontier
    pv, pf = wv["PRS Worksheet"], wf["PRS Worksheet"]
    bands = []
    for rr in range(3, 19):
        m = re.fullmatch(r"=J(\d+)", str(pf.cell(rr, 12).value))
        rating = str(pv.cell(int(m.group(1)), 9).value).strip() if m else None
        bands.append([pv.cell(rr, 9).value, rating])
    prs_in, prs_exp = [], {}
    for rr in range(100, pv.max_row + 1):
        a, e_f = pf.cell(rr, 1).value, str(pf.cell(rr, 5).value)
        if isinstance(a, str) and not a.startswith("=") and "VLOOKUP(B" in e_f:
            c = pv.cell(rr, 3).value
            prs_in.append([a, pv.cell(rr, 2).value, c if num(c) else None])
            prs_exp[a] = {"erp": r8(pv.cell(rr, 5).value), "crp": r8(pv.cell(rr, 7).value), "spread": r8(pv.cell(rr, 4).value)}
            data_c[a] = {"region": region_of.get(a), "rating": None, "prs": pv.cell(rr, 2).value, **prs_exp[a],
                         "tax": r8(pv.cell(rr, 6).value), "basis": "PRS"}
    # regions
    rg_v, rg_f = wv["Regional Weighted Averages"], wf["Regional Weighted Averages"]
    regions_in, regions_exp, regions_data, glob, groupings = {}, {}, {}, None, {}
    for rr in range(2, rg_f.max_row + 1):
        bform = str(rg_f.cell(rr, 2).value)
        m = re.fullmatch(r"=SUM\(B(\d+):B(\d+)\)", bform)
        if m:
            reg = rg_v.cell(rr, 1).value
            erp_rule = "weighted_erp" if str(rg_f.cell(rr, 5).value).startswith("=SUM(K") else "mature_plus_crp"
            regions_in[reg] = [[rg_v.cell(k, 1).value] + [r8(rg_v.cell(k, c).value) for c in (2, 4, 5, 6, 7)]
                               for k in range(int(m.group(1)), int(m.group(2)) + 1)]
            regions_exp[reg] = {"erp": r8(rg_v.cell(rr, 5).value), "spread": r8(rg_v.cell(rr, 4).value),
                                "crp": r8(rg_v.cell(rr, 6).value), "tax": r8(rg_v.cell(rr, 7).value)}
            regions_data[reg] = {**regions_exp[reg], "gdp": r8(rg_v.cell(rr, 2).value)}
        if rg_v.cell(rr, 1).value == "Global" and str(rg_f.cell(rr, 2).value).startswith("=SUM(H"):
            glob = {"erp": r8(rg_v.cell(rr, 2).value), "crp": r8(rg_v.cell(rr, 3).value),
                    "spread": r8(rg_v.cell(rr, 4).value), "tax": r8(rg_v.cell(rr, 5).value)}
            regions_data["Global"] = glob
        a = rg_v.cell(rr, 1).value
        if rr > 180 and isinstance(a, str) and num(rg_v.cell(rr, 2).value) and a != "Global" and a not in regions_data:
            groupings[a] = {"erp": r8(rg_v.cell(rr, 2).value), "spread": r8(rg_v.cell(rr, 3).value),
                            "tax": r8(rg_v.cell(rr, 4).value)}
    stale = [n for reg in regions_in.values() for n, _, s, _, c, _ in reg
             if n in expected and not close(c, expected[n]["crp"])]
    if stale:
        notes.append(f"{f.name}: regional sheet CRPs differ from ERP sheet for {len(stale)} countries (e.g. {stale[:4]})")
    # relative volatility
    vf, vv = wf["Relative Equity Volatility"], wv["Relative Equity Volatility"]
    windows = []
    for rr in range(2, 7):
        e_m = re.search(r"STDEV\(I(\d+):I(\d+)\)\*\((\d+)\^0\.5\)", vf.cell(rr, 2).value)
        b_m = re.search(r"STDEV\(C(\d+):C(\d+)\)\*\((\d+)\^0\.5\)", vf.cell(rr, 3).value)
        windows.append([int(x) for x in e_m.groups() + b_m.groups()])
    first = 15
    series = {}
    for pc, rc in (("B", "C"), ("H", "I")):
        prices, k = [], first
        while num(vv[f"{pc}{k}"].value):
            if k > first and vf[f"{rc}{k}"].value != f"={pc}{k}/{pc}{k - 1}-1":
                notes.append(f"{f.name}: return cell {rc}{k} is {vf[f'{rc}{k}'].value!r}")
            prices.append(vv[f"{pc}{k}"].value)
            k += 1
        series[pc] = prices
    his_scalar = vv["D7"].value
    if not close(his_scalar, scalar):
        notes.append(f"{f.name}: ERP sheet scalar {scalar} != volatility sheet D7 {his_scalar}")
    case = {"name": f.stem, "file": f.name, "date": date,
            "inputs": {"mature_erp": mature, "us_erp": us_erp, "us_country": us_country, "scale": scale,
                       "sp_to_moodys": sp_map, "prior_spreads_bp": prior, "spread_overrides_bp": overrides,
                       "cds_pairs": pairs, "cds_benchmark": bench, "countries": countries,
                       "prs": {"bands": bands, "countries": prs_in}, "regions": regions_in, "region_erp_rule": erp_rule,
                       "vol": {"first_row": first, "windows": windows, "equity": series["H"], "bond": series["B"]}},
            "expected": {"scalar": scalar, "median_cds_change": med, "spreads_bp": {k: v for k, v in jk.items()
                                                                                    if k != "Rating" and num(v)},
                         "countries": expected, "prs": prs_exp, "regions": regions_exp,
                         **({"Global": glob} if glob else {})}}
    v = {"date": date, "file": f.name, "url": URL + f.name, "mature_erp": mature, "us_erp": us_erp,
         "scalar": r8(scalar), "cds_benchmark": bench_name, "cds_benchmark_value": bench,
         "spreads_bp": {k: r8(v) for k, v in jk.items() if k != "Rating" and num(v)},
         "median_cds_change": r8(med), "prs_bands": bands, "countries": data_c, "regions": regions_data,
         "groupings": groupings}
    return case, v


def branch_case(path: Path, v: dict) -> dict:
    """pycel: set 'ERPs by country'!E5 (scale default spread by relative volatility?) to No, recalculate."""
    import statistics as st
    import types
    import warnings
    warnings.filterwarnings("ignore")
    extra = types.ModuleType("pycel_extra")  # pycel lacks MEDIAN and STDEV
    flat = lambda a: [y for x in a for y in (flat(x) if isinstance(x, (tuple, list)) else [x]) if num(y)]
    extra.median = lambda *a: st.median(flat(a))
    extra.stdev = lambda *a: st.stdev(flat(a))
    sys.modules["pycel_extra"] = extra
    from pycel import ExcelCompiler
    xl, sh = ExcelCompiler(filename=str(path), plugins=["pycel_extra"]), "ERPs by country!"
    names = [xl.evaluate(f"{sh}A{r}") for r in range(9, 166)]
    rows = [9 + names.index(n) for n in ("Abu Dhabi", "Angola", "Ethiopia", "India", "Brazil", "United States")]
    cols = {"D": "spread", "E": "erp", "F": "crp", "H": "erp_cds", "I": "crp_cds"}
    for r in rows:
        for c in cols:
            xl.evaluate(f"{sh}{c}{r}")
    xl.set_value(sh + "E5", "No")
    xl.recalculate()
    out = []
    for r in rows:
        n = names[r - 9]
        exp = {k: (None if xl.evaluate(f"{sh}{c}{r}") == "NA" else xl.evaluate(f"{sh}{c}{r}")) for c, k in cols.items()}
        out.append([n, v["countries"][n]["rating"], v["countries"][n]["cds_net"], exp])
    return {"name": f"{path.stem} E5=No", "file": path.name,
            "params": {"spreads": v["spreads_bp"], "scalar": v["scalar"], "mature_erp": v["mature_erp"],
                       "us_erp": v["us_erp"], "us_country": "United States"}, "rows": out}


def read_tax(p: Path, r8) -> dict:
    import openpyxl
    ws = openpyxl.load_workbook(p, data_only=True).worksheets[0]
    rates = {ws.cell(r, 1).value: [r8(ws.cell(r, 2).value), r8(ws.cell(r, 3).value)]
             for r in range(7, ws.max_row + 1) if ws.cell(r, 1).value}
    return {"file": p.name, "url": URL + p.name, "as_of": ws["B1"].value.date().isoformat(),
            "fields": ["marginal", "with_global_min_tax"], "rates": rates}


def read_effective_tax(p: Path, r8) -> dict:
    import openpyxl
    ws = openpyxl.load_workbook(p, data_only=True).worksheets[0]
    fields = ["firms", "effective_aggregate", "q1", "median", "q3", "marginal"]
    rates = {ws.cell(r, 1).value: [r8(ws.cell(r, c).value) for c in (2, 5, 6, 7, 8, 9)]
             for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value}
    return {"file": p.name, "url": "https://pages.stern.nyu.edu/~adamodar/pc/blog/" + p.name,
            "note": "2025 taxes paid / taxable income across listed firms, by country", "fields": fields, "rates": rates}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("inputs", nargs="?")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--country")
    ap.add_argument("--as-of")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--build-cases", nargs="*")
    a = ap.parse_args()
    if a.selftest:
        return selftest()
    if a.build_cases is not None:
        default = HERE.parents[2] / "spreadsheets"
        names = ["ctryprem.xlsx", "ctrypremApr26.xlsx", "ctrypremJuly23.xlsx", "ctrypremJuly24.xlsx",
                 "ctrypremJuly25.xlsx", "ctrypremJuly26.xlsx", "countrytaxrates.xlsx", "CountryTaxRates2025.xlsx",
                 "Ratings&CDS2023.xlsx"]
        build([Path(x) for x in a.build_cases] or [default / n for n in names])
        return 0
    if a.inputs:
        q = {k: v for k, v in json.loads(Path(a.inputs).read_text(encoding="utf-8")).items() if not k.startswith("_")}
    elif a.country:
        q = {"country": a.country, "as_of": a.as_of}
    else:
        ap.print_help()
        return 1
    lines, out = report(q, load())
    print(json.dumps(out, indent=1) if a.json else "\n".join(lines))
    return 0


if __name__ == "__main__":
    sys.exit(main())
