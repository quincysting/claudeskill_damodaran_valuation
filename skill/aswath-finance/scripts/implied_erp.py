"""Implied equity risk premium (S&P 500), intrinsic riskfree rate and currency riskfree rates, Damodaran's way.

Transcribed from Damodaran's own spreadsheets (https://pages.stern.nyu.edu/~adamodar/pc/...):
  implprem/ERP{Jan23,Aug23,Jan24,Jan25,July25,Oct25,Jan26,March26,July26}.xlsx, sheet 'Impl premium calculator':
    Expected Earnings  E0 = estimated earnings for the most recent year ('Expected growth rate'!H4/H5), E_t = E_t-1 (1+g)
    Expected cash payout  p0 = dividends+buybacks / reported LTM earnings; constant, or (sustainable = "Yes")
                          linear to 1 - g_stable / ROE(avg last decade) by year 5
    Expected Dividends + Buybacks  CF_t = E_t p_t;  terminal CF = E_5 (1+g_stable) p_T
    Expected Terminal Value  CF_6 / (r - g_stable), g_stable = riskfree unless 'change growth in long term' = Yes
    Time to cashflow  tau = months left in the year / 12 (VLOOKUP on 'Answers', 2025+ layouts; older: tau = 1),
                      CF_1 x tau at tau, CF_t at tau + t - 1;  r = riskfree + ERP;  ERP = goal seek (index = value)
    ERP with default-risk adjusted US$ riskfree = ERP + US default spread (Aa1, since the May 2025 downgrade)
  blog/TariffERPbyday.xlsx, blog/ERPbyDayCOVID.xlsx, blog/ERPbyDay2008Crisis.xlsx -> Damodaran's daily ERP series (data)
  intrinsicvsactualrate23/24/25.xlsx, blog/IntrinsicRiskfree2026.xlsx, blog/intrinsicvsactualrateSept26.xlsx:
    intrinsic riskfree = inflation + real GDP growth; smoothed = 10-yr avg inflation + 10-yr avg real growth
  currencyriskfree2024.xls, CurrencyRiskfreeJuly2024.xlsx, currencyriskfree2025.xlsx: riskfree = govt bond - default spread
  Currency&Inflation2023.xlsx, CurrencyRiskfreeJuly2024.xlsx, DiffInflationRiskfree26.xlsx (compounded) and
  DiffInflRiskfree2026.xlsx (additive): riskfree_local = (1+rf_US$)(1+infl_local)/(1+infl_US$) - 1

Usage:
  python implied_erp.py erp inputs.json [--json]       implied ERP (+ index value at 2-6% ERPs, or at "erp")
  python implied_erp.py intrinsic inputs.json | --year 2025
  python implied_erp.py currency inputs.json | --lookup Brazil
  python implied_erp.py history 2025-04-08            Damodaran's published ERP (daily crisis series, monthly, annual)
  python implied_erp.py --from-case ERPJuly26         the source inputs for a date, as a starting JSON
  python implied_erp.py --selftest                    replay data/implied_erp_cases.json
  python implied_erp.py --build-cases <xlsx...>       maintenance (openpyxl/xlrd/pycel): rebuild cases + data

erp inputs (decimals; index units): index, riskfree, cash_flow (dividends+buybacks, trailing 12m) or cash_yield,
  growth (expected earnings growth, years 1-5); optional earnings_ltm (reported trailing 12m; sets payout) — if given,
  earnings_base is required (spreadsheet default: current-year estimated earnings; = earnings_ltm for "Reported LTM"), month (default "January") or months_to_year_end, stable_growth (default
  riskfree), sustainable_payout (default false) + roe, default_spread (US sovereign spread -> ERP vs default-free
  rate), erp (value the index at this ERP), value_at (list of ERPs; default 2%..6%).
intrinsic inputs: inflation, real_growth, treasury_rate; optional inflation_10y, real_growth_10y (10 values).
currency inputs: bond_rate + default_spread | rating | country (Jan 2026 spreads from ginzu_data.json), or
  inflation_local + inflation_base + base_riskfree [+ additive]. Keys starting with "_" are ignored.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
DATA = HERE.parent / "data" / "implied_erp_data.json"
CASES = HERE.parent / "data" / "implied_erp_cases.json"
MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august", "september", "october",
          "november", "december"]


# ---------------------------------------------------------------- implied ERP ('Impl premium calculator')
def prepare(inp: dict) -> dict:
    rf, g = inp["riskfree"], inp["growth"]
    cf0 = inp["cash_flow"] if "cash_flow" in inp else inp["cash_yield"] * inp["index"]
    e_ltm = inp.get("earnings_ltm")
    if e_ltm and "earnings_base" not in inp:  # the template's default base is the current-year *estimate*
        raise ValueError("earnings_base missing: the spreadsheet default is estimated earnings for the current year; pass that "
                         "estimate, or set earnings_base = earnings_ltm to use the spreadsheet's 'Reported LTM' option")
    e0 = inp["earnings_base"] if e_ltm else cf0
    p0 = cf0 / e_ltm if e_ltm else 1.0
    gs = inp.get("stable_growth", rf)
    sust = bool(inp.get("sustainable_payout"))
    p_t = 1 - gs / inp["roe"] if sust else p0
    left = inp.get("months_to_year_end") or 12 - MONTHS.index(str(inp.get("month", "January")).lower())
    tau = left / 12
    earn, pay, cash = [e0], [p0], [cf0]
    for t in range(1, 6):
        earn.append(earn[-1] * (1 + g))
        pay.append(p0 - (p0 - p_t) * t / 5 if sust else p0)
        cash.append(earn[-1] * pay[-1])
    earn.append(earn[5] * (1 + gs))
    pay.append(p_t)
    cash.append(earn[6] * p_t)
    return {"rf": rf, "gs": gs, "tau": tau, "earnings": earn, "payout": pay, "cash": cash}


def index_value(inp: dict, erp: float, p: dict | None = None) -> dict:
    p = p or prepare(inp)
    r, tau = p["rf"] + erp, p["tau"]
    times = [tau + t - 1 for t in range(1, 6)]
    tv = p["cash"][6] / (r - p["gs"])
    pv = [p["cash"][1] * tau / (1 + r) ** tau] + [p["cash"][t] / (1 + r) ** times[t - 1] for t in range(2, 6)]
    pv[4] += tv / (1 + r) ** times[4]
    return {"value": sum(pv), "terminal_value": tv, "pv": pv, "times": times}


def implied_erp(inp: dict) -> dict:
    """Goal seek: the ERP at which the present value of expected cash flows equals the index."""
    p = prepare(inp)
    lo, hi = p["gs"] - p["rf"] + 1e-9, 1.0
    if index_value(inp, lo, p)["value"] < inp["index"]:
        raise ValueError("no ERP makes value = index (cash flows too small for this stable growth)")
    for _ in range(200):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if index_value(inp, mid, p)["value"] > inp["index"] else (lo, mid)
    erp = (lo + hi) / 2
    out = {"erp": erp, "expected_return": p["rf"] + erp, "cash_yield": p["cash"][0] / inp["index"], **p,
           **index_value(inp, erp, p)}
    if inp.get("default_spread") is not None:
        out["default_free_riskfree"] = p["rf"] - inp["default_spread"]
        out["erp_vs_default_free"] = erp + inp["default_spread"]
    return out


def erp_report(inp: dict) -> str:
    o = implied_erp(inp)
    rows = [f"{'year':>5} {'earnings':>10} {'payout':>7} {'cash flow':>10} {'time':>6} {'PV':>10}"]
    for t in range(7):
        pv = f"{o['pv'][t - 1]:10.2f}" if 1 <= t <= 5 else ""
        tm = f"{o['times'][t - 1]:6.3f}" if 1 <= t <= 5 else ""
        rows.append(f"{['base', 1, 2, 3, 4, 5, 'term'][t]!s:>5} {o['earnings'][t]:10.2f} {o['payout'][t]:7.2%} "
                    f"{o['cash'][t]:10.2f} {tm:>6} {pv}")
    rows += ["", f"Terminal value (yr 5) {o['terminal_value']:,.2f}  (g stable {o['gs']:.2%}, r {o['expected_return']:.2%})",
             f"Index {inp['index']:,.2f}   cash yield {o['cash_yield']:.2%}   riskfree {o['rf']:.2%}   growth {inp['growth']:.2%}",
             f"IMPLIED ERP {o['erp']:.4%}   expected return on stocks {o['expected_return']:.4%}"]
    if "erp_vs_default_free" in o:
        rows.append(f"ERP vs default-free US$ rate ({o['default_free_riskfree']:.2%} = T.Bond - {inp['default_spread']:.2%}"
                    f" default spread) = {o['erp_vs_default_free']:.4%}; mature-market ERP (net of US default spread) = {o['erp']:.4%}")
    erps = [inp["erp"]] if "erp" in inp else inp.get("value_at", [0.02, 0.03, 0.04, 0.05, 0.06])
    rows.append("Index value at ERP: " + "  ".join(f"{e:.2%} -> {index_value(inp, e)['value']:,.0f}" for e in erps))
    return "\n".join(rows)


# ---------------------------------------------------------------- intrinsic riskfree rate
def intrinsic_riskfree(inflation: float, real_growth: float, treasury_rate: float | None = None,
                       inflation_10y: list | None = None, real_growth_10y: list | None = None) -> dict:
    out = {"intrinsic": inflation + real_growth}
    if treasury_rate is not None:
        out["treasury_minus_intrinsic"] = treasury_rate - out["intrinsic"]  # the sheet's "Fed effect" column
        out["treasury_over_intrinsic"] = treasury_rate / out["intrinsic"] if out["intrinsic"] > 0 else None
    if inflation_10y and real_growth_10y:
        out["smoothed"] = sum(inflation_10y) / len(inflation_10y) + sum(real_growth_10y) / len(real_growth_10y)
        if treasury_rate is not None:
            out["treasury_minus_smoothed"] = treasury_rate - out["smoothed"]
    return out


def intrinsic_for_year(data: dict, year: str) -> dict:
    rows = data["intrinsic"]["rows"]
    k = next(i for i, r in enumerate(rows) if str(r["year"]).startswith(str(year)))
    hist = (data["intrinsic"]["pre_1954"] + rows)[:len(data["intrinsic"]["pre_1954"]) + k + 1][-10:]
    r = rows[k]
    return {"year": r["year"], "inflation": r["inflation"], "real_growth": r["real_growth"], "treasury_rate": r["tbond"],
            "fed_funds": r.get("fed_funds"), **intrinsic_riskfree(r["inflation"], r["real_growth"], r["tbond"],
                                                                  [h["inflation"] for h in hist], [h["real_growth"] for h in hist])}


# ---------------------------------------------------------------- currency riskfree rates
def riskfree_from_bond(bond_rate: float, default_spread: float) -> float:
    return bond_rate - default_spread


def riskfree_from_inflation(base_riskfree: float, inflation_local: float, inflation_base: float,
                            additive: bool = False) -> float:
    if additive:
        return base_riskfree + inflation_local - inflation_base
    return (1 + base_riskfree) * (1 + inflation_local) / (1 + inflation_base) - 1


def sovereign_spread(inp: dict) -> tuple[float, str]:
    if "default_spread" in inp:
        return inp["default_spread"], "given"
    from cost_of_capital import load  # Damodaran's Jan 2026 country table (ginzu_data.json)
    countries = load()["countries"]
    if "country" in inp:
        c = countries[inp["country"]]
        return c["default_spread"], f"{inp['country']} ({c['rating']}), Jan 2026"
    hit = [c["default_spread"] for c in countries.values() if c.get("rating") == inp["rating"]]
    if not hit:
        raise KeyError(f"rating {inp['rating']} not in Damodaran's Jan 2026 country table")
    return hit[0], f"rating {inp['rating']}, Jan 2026"


def currency(inp: dict, data: dict) -> dict:
    if "bond_rate" in inp:
        s, src = sovereign_spread(inp)
        return {"method": "govt bond - default spread", "default_spread": s, "spread_source": src,
                "riskfree": riskfree_from_bond(inp["bond_rate"], s)}
    base = data["currency_inflation"][max(data["currency_inflation"])]  # Damodaran's latest table
    rf, ib = inp.get("base_riskfree", base["base_riskfree"]), inp.get("inflation_base", base["base_inflation"])
    return {"method": "differential inflation (" + ("additive" if inp.get("additive") else "compounded") + ")",
            "base_riskfree": rf, "inflation_base": ib,
            "riskfree": riskfree_from_inflation(rf, inp["inflation_local"], ib, bool(inp.get("additive")))}


# Damodaran's published rows kept as-is but flagged on lookup (see references/tool-implied-erp.md, Limits)
FLAGGED = {"2024-07": "bond rates and spreads repeat the Jan 2024 table (header there: 12/31/22)",
           **{f"2025-01 {c}": "row looks misaligned in the source file" for c in
              ("Brazilian Reai", "Bulgarian Lev", "Russian Ruble", "Romanian Lev")}}


def currency_lookup(data: dict, name: str) -> list[str]:
    out, n = [], name.lower()
    for v, t in sorted(data["currency_sovereign"].items()):
        for r in t["rows"]:
            if n in r[0].lower():
                flag = FLAGGED.get(f"{v} {r[0]}") or FLAGGED.get(v)
                out.append(f"{v} {r[0]:22s} govt bond {r[1]:.2%} ({r[2]}) - spread {r[3]:.2%} = riskfree {r[4]:.2%}"
                           + (f"  [caution: {flag}]" if flag else ""))
    for v, t in sorted(data["currency_inflation"].items()):
        for r in t["rows"]:
            if n in r[0].lower():
                f = lambda x: f"{x:.2%}" if isinstance(x, (int, float)) else str(x)
                out.append(f"{v} {r[0]:22s} expected inflation {f(r[2])} (actual {f(r[1])}) vs US {t['base_inflation']:.2%},"
                           f" US$ rf {t['base_riskfree']:.2%} -> riskfree {f(r[3])} [{t['method']}]")
    return out


# ---------------------------------------------------------------- Damodaran's published ERPs
def history(data: dict, when: str) -> list[str]:
    out = []
    for name, s in data["erp_daily"].items():
        for r in s["rows"]:
            if r["date"].startswith(when):
                extra = "".join(f"  {k} {v:.2%}" for k, v in r.items() if k.startswith("covid") and v is not None)
                out.append(f"daily ({name}) {r['date']}: S&P {r['sp500']:,.2f}  T.Bond {r['tbond']:.2%}  ERP {r['erp']:.2%}{extra}")
    for r in data["erp_monthly"]:
        if r["date"][:7] == when[:7] or (len(when) == 4 and r["date"].startswith(when)):
            vs = "  ".join(f"{k.replace('erp_', '')} {r[k]:.2%}" for k in r if k.startswith("erp_") and r[k] is not None)
            out.append(f"start of month {r['date']}: S&P {r['sp500']:,.0f}  T.Bond {r['tbond']:.2%}  ERP (T12m) {r['erp']:.2%}"
                       + (f"  [{vs}]" if vs else "") + (f"  growth {r['growth']:.2%}" if r.get("growth") else "")
                       + (f"  note: {r['note']}" if r.get("note") else ""))
    for r in data["erp_annual"]:
        if when[:4] in (str(r["year"]), str(r["year"] + 1)) and r.get("erp") is not None:
            out.append(f"annual (end of {r['year']} = start of {r['year'] + 1}): ERP {r['erp']:.2%}  T.Bond {r['tbond']:.2%}")
    return out


# ---------------------------------------------------------------- self test
def selftest() -> int:
    cases = json.loads(CASES.read_text(encoding="utf-8"))
    data = json.loads(DATA.read_text(encoding="utf-8"))
    bad = n = 0

    def check(name, his, ours, tol, rel=False):
        nonlocal bad, n
        err = abs(ours / his - 1) if rel else abs(ours - his)
        n += 1
        bad += err > tol
        print(f"{'ok ' if err <= tol else 'BAD'} {name:58s} Damodaran {his:14.8f} ours {ours:14.8f}  diff {err:.1e}")

    for c in cases["cases"]:
        k, inp = c["kind"], c.get("inputs", {})
        if k in ("erp", "erp_rounded"):
            o = implied_erp(inp)
            check(c["name"], c["expected"], o["erp"], 1e-6 if k == "erp" else 5e-5 + 1e-9)
            if "expected_adj" in c:
                check(c["name"] + " (vs default-free rf)", c["expected_adj"], o["erp_vs_default_free"], 1e-6)
        elif k == "value":
            check(c["name"], c["expected"], index_value(inp, c["erp"])["value"], 5e-4, rel=True)
        elif k == "intrinsic":
            o = intrinsic_riskfree(**inp)
            for key, his in c["expected"].items():
                check(f"{c['name']} {key}", his, o[key], 1e-9)
        elif k == "currency_table":
            t = data[c["table"]][c["vintage"]]
            worst, rows = 0.0, [r for r in t["rows"] if isinstance(r[-1], (int, float))]
            for r in rows:
                if c["table"] == "currency_sovereign":
                    ours = riskfree_from_bond(r[1], r[3])
                else:  # the sheet's fallback: actual inflation when the IMF forecast is missing (additive 2026 table only)
                    loc = r[2] if isinstance(r[2], (int, float)) else r[1]
                    ours = riskfree_from_inflation(t["base_riskfree"], loc, t["base_inflation"], t["method"] == "additive")
                worst = max(worst, abs(ours - r[-1]))
            n += 1
            bad += worst > 1e-9
            print(f"{'ok ' if worst <= 1e-9 else 'BAD'} {c['name']:58s} {len(rows)} rows, max diff {worst:.1e}")
    for s in cases.get("series", []):  # informational: Damodaran's daily ERPs rounded to 0.01%, index to whole points
        print(f"info {s['name']}: {s['summary']}")
    for d in cases.get("discrepancies", []):
        print(f"info not reproducible from the saved file: {d}")
    print(f"{n - bad}/{n} checks match Damodaran's spreadsheets")
    return 1 if bad else 0


# ---------------------------------------------------------------- maintenance: --build-cases
CALC = "Impl premium calculator"
CASH_CODES = ("C", "A10", "A5", "P10", "P5", "NORM", "NC")
VARIANTS = [("ERP (T12m)", "C", False), ("ERP (T12 m with sustainable payout)", "C", True),
            ("ERP (Smoothed)", "A10", False), ("ERP (Normalized)", "NORM", False), ("ERP (Net Cash Yield)", "NC", False)]
MONTHLY_KEYS = {"Start of month": "date", "S&P 500": "sp500", "T.Bond Rate": "tbond", "$ Riskfree Rate": "usd_riskfree",
                "Ten-year average CF": "cf_avg10", "CF (Trailing 12 month)": "cf_t12m", "Normalized CF": "cf_normalized",
                "Expected growth rate": "growth", "ERP (T12m)": "erp", "ERP (T12 m with sustainable payout)": "erp_sustainable",
                "ERP (T12m) with adj riskfree rate": "erp_adj_riskfree", "ERP (Smoothed)": "erp_smoothed",
                "ERP (Normalized)": "erp_normalized", "ERP (Net Cash Yield)": "erp_net_cash", "ERP (Covid Adjusted)": "erp_covid",
                "Expected Return": "expected_return", "Notes": "note"}
DAILY_KEYS = {"Date": "date", "S&P 500": "sp500", "T. Bond Rate": "tbond", "10-yr Treasuries": "tbond", "Implied ERP": "erp",
              "ERP": "erp", "Implied Premium": "erp", "COVID ERP": "covid_erp", "Earnings drop in 2020": "covid_earnings_drop",
              "Recovery by 2024": "covid_recovery", "Earnings g- next 5 years": "growth"}
DAILY_NAMES = {"ERPbyDay2008Crisis.xlsx": "crisis2008", "ERPbyDayCOVID.xlsx": "covid2020", "TariffERPbyday.xlsx": "tariffs2025"}
CURRENCY_VINTAGE = {"Currency&Inflation2023.xlsx": "2023-07", "currencyriskfree2024.xls": "2024-01",
                    "CurrencyRiskfreeJuly2024.xlsx": "2024-07", "currencyriskfree2025.xlsx": "2025-01",
                    "DiffInflRiskfree2026.xlsx": "2026-01", "DiffInflationRiskfree26.xlsx": "2026-07"}
# Damodaran's published numbers checked against the file's model (post text), and post numbers the files do not reproduce
POSTED = [("ERPOct25.xlsx", "erp_rounded", "ERPOct25 as posted (cash C, month October) -> 4.01%",
           {"cash": "C", "month": "October"}, 0.0401, None),
          ("ERPOct25.xlsx", "value", "ERPOct25 post: value at 4.25% ERP, sustainable payout -> 'about 5940'",
           {"cash": "C", "month": "October", "sustainable_payout": True}, 5940, 0.0425),
          ("ERPJan26.xlsx", "value", "ERPJan26 post: value at 6% ERP -> 4790", {}, 4790, 0.06)]
POST_ONLY = ["ERPAug23: the post says 4.44% for August 2023; the saved file gives 4.83% (T12m), 4.38% (sustainable "
             "payout), 4.29% (normalized) - none is 4.44%",
             "ERPJan26: the post says the index is worth 14834 at a 2% ERP; the file's model gives 14634 (the 6% value, "
             "4790, matches) - likely a typo in the post",
             "ERPJuly25: the file's US default spread cell is 0.40% (ERP vs default-free rate 4.61%); the post nets out "
             "0.27% (4.48%); the T.Bond-based (mature market) ERP of 4.21% agrees",
             "ERPJuly26: the post and the monthly table say 4.20% (4.42% vs the default-free rate); the saved file solves "
             "to 4.17% (4.39%); 4.20% needs growth of 13.84% vs the file's 13.69% - inputs revised after posting"]
BRANCH_FILES = ("ERPJan26.xlsx", "ERPJuly26.xlsx")


def _num(x):
    if isinstance(x, str):
        s = x.strip()
        s = s.replace(",", ".") if s.count(",") == 1 and "." not in s else s.replace(",", "")  # "5158,20" in the source file
        try:
            return float(s.rstrip("%*")) / (100 if s.endswith("%") else 1)
        except ValueError:
            return None
    return x


def _date(x):
    import datetime as dt
    if isinstance(x, dt.datetime):
        return x.date().isoformat()
    return dt.datetime.strptime(x.strip(), "%d-%b-%y").date().isoformat() if isinstance(x, str) else x


def _table(ws, header_row: int, keys: dict) -> list[dict]:
    cols = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(header_row, c).value
        if isinstance(h, str) and h.strip() in keys and keys[h.strip()] not in cols.values():
            cols[c] = keys[h.strip()]
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = {k: ws.cell(r, c).value for c, k in cols.items()}
        if row.get("date") is None or not isinstance(_num(row.get("erp")), float):
            continue
        out.append({k: (_date(v) if k == "date" else v if k == "note" else _num(v)) for k, v in row.items()})
    return out


def _erp_file(path: Path, openpyxl) -> dict:
    """Damodaran's inputs, cached ERP, cash-flow choices and monthly-table row from one 'Impl premium calculator' file."""
    import re
    wf, wv = openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)
    f, v = wf[CALC], wv[CALC]
    labels = [(r, str(f.cell(r, 1).value).strip()) for r in range(1, f.max_row + 1) if f.cell(r, 1).value]
    row = lambda p, after=0: next((r for r, a in labels if a.startswith(p) and r > after), None)
    B = lambda p: v.cell(row(p), 2).value if row(p) else None
    code = {v.cell(r, 8).value: v.cell(r, 7).value for r in range(1, 40) if isinstance(v.cell(r, 8).value, str)}
    blk = row("Implied Equity Risk Premium Calculator")
    pr, cr, ivr = row("Expected cash payout", blk), row("Expected Dividends + Buybacks", blk), row("Intrinsic Value of Index", blk)
    er = next(r for r in range(blk, blk + 4) if re.search("treasury|current level of Index", str(f.cell(r, 1).value)))
    ar = next((r for r in range(blk, blk + 4) if "default-risk" in str(f.cell(r, 1).value)), None)
    roe_cell = re.search(r"F(\d+)", f.cell(pr, 8).value).group(0)
    lt = row("Do you want to change the growth rate")
    inp = {"index": B("Enter current level of index"),
           "riskfree": v.cell(row("Enter current long term risk free rate") or row("Enter current 10-year"), 2).value,
           "cash_flow": v.cell(cr, 2).value, "growth": B("Expected growth rate in earnings for the next 5 years"),
           "earnings_ltm": code["Reported LTM"], "earnings_base": v.cell(row("Expected Earnings", blk), 2).value,
           "month": B("Estimation month") or "January", "roe": wv["Buyback & Dividend computation"][roe_cell].value}
    if lt and v.cell(lt, 2).value == "Yes":
        inp["stable_growth"] = v.cell(lt + 1, 2).value
    if B("Do you want to adjust the payout ratio") == "Yes":
        inp["sustainable_payout"] = True
    if B("Default spread") is not None:
        inp["default_spread"] = B("Default spread")
    inp["_cash_choices"] = {k: code[k] for k in CASH_CODES if k in code}
    m = wv["Implied ERP (Monthly from 9-08)"]
    hdr = {str(m.cell(1, c).value).strip(): c for c in range(1, m.max_column + 1) if m.cell(1, c).value}
    last = max(r for r in range(2, m.max_row + 1) if isinstance(_num(m.cell(r, hdr["ERP (T12m)"]).value), float))
    cells = {"month": f"B{row('Estimation month')}", "base": f"B{row('Base Year Earnings')}", "cash": f"B{row('What cash yield')}",
             "lt": f"B{lt}", "lt_rate": f"B{(lt or 0) + 1}", "sust": f"B{row('Do you want to adjust the payout ratio')}",
             "value": f"B{ivr}"}
    fb = row("Intrinsic Value of Index")  # first block: value at the sheet's chosen ERP
    return {"name": path.stem, "inputs": inp, "erp": v.cell(er, 3).value, "erp_adj": v.cell(ar, 3).value if ar else None,
            "value_erp": B("ERP used to compute fair value"), "value": v.cell(fb, 3).value, "cells": cells,
            "monthly_row": {k: m.cell(last, c).value for k, c in hdr.items()}, "wv": wv}


def _erp_data(e: dict, openpyxl) -> dict:
    wv = e["wv"]
    m = wv["Implied ERP (Monthly from 9-08)"]
    a = wv["Implied ERP- Annual since 1960"]
    annual = [{"year": a.cell(r, 1).value, "earnings_yield": a.cell(r, 2).value, "dividend_yield": a.cell(r, 3).value,
               "tbond": a.cell(r, 4).value, "erp": a.cell(r, 5).value}
              for r in range(2, a.max_row + 1) if isinstance(a.cell(r, 1).value, int)]
    c = wv[CALC]
    ref = {c.cell(r, 8).value: c.cell(r, 7).value for r in range(15, 35) if c.cell(r, 8).value in ("HUS", "HG", "ERP10", "ERPLong", "ERPHigh")}
    return {"erp_monthly": _table(m, 1, MONTHLY_KEYS), "erp_annual": annual,
            "erp_reference": {"from": e["name"], "historical_us": ref["HUS"], "historical_global": ref["HG"],
                              "avg_implied_last_decade": ref["ERP10"], "avg_implied_1960_now": ref["ERPLong"],
                              "max_implied_1960_now": ref["ERPHigh"]}}


def _branches(path: Path, e: dict) -> list[dict]:
    """Recalculate Damodaran's template (pycel) under options the saved files never use; value at the cached ERP."""
    from pycel import ExcelCompiler
    i, c = e["inputs"], e["cells"]
    scen = [("sustainable payout", {c["sust"]: "Yes"}, {"sustainable_payout": True}),
            ("stable growth 3%", {c["lt"]: "Yes", c["lt_rate"]: 0.03}, {"stable_growth": 0.03}),
            ("sustainable payout + stable growth 2.5%", {c["sust"]: "Yes", c["lt"]: "Yes", c["lt_rate"]: 0.025},
             {"sustainable_payout": True, "stable_growth": 0.025}),
            ("base earnings = reported LTM", {c["base"]: "Reported LTM"}, {"earnings_base": i["earnings_ltm"]}),
            ("month October", {c["month"]: "October"}, {"month": "October"}),
            ("cash flow = normalized (NORM)", {c["cash"]: "NORM"}, {"cash_flow": i["_cash_choices"]["NORM"]})]
    out = []
    for name, cells, over in scen:
        xl = ExcelCompiler(filename=str(path))
        xl.evaluate(f"{CALC}!{c['value']}")
        for k, val in cells.items():
            xl.set_value(f"{CALC}!{k}", val)
        xl.recalculate()
        out.append({"kind": "value", "name": f"{e['name']} template branch: {name}", "erp": e["erp"],
                    "inputs": {**i, **over}, "expected": xl.evaluate(f"{CALC}!{c['value']}")})
    return out


def _series_2008(rows: list[dict], cf: float) -> str:
    """Damodaran's Oct 7 - Dec 31 2008 daily ERPs vs the model with the sheet's annualized cash flow; tolerance = input rounding."""
    ok, off = 0, []
    use = [r for r in rows if r["date"] >= "2008-10-07"]
    for r in use:
        inp = {"index": r["sp500"], "riskfree": r["tbond"], "cash_flow": cf, "growth": r["growth"]}
        e = implied_erp(inp)["erp"]
        tol = 5e-5 + abs(implied_erp({**inp, "index": r["sp500"] + 0.5})["erp"] - e) \
            + abs(implied_erp({**inp, "riskfree": r["tbond"] + 5e-5})["erp"] - e)
        ok += abs(e - r["erp"]) <= tol
        if abs(e - r["erp"]) > tol:
            off.append(f"{r['date']} Damodaran {r['erp']:.2%} ours {e:.3%}")
    return (f"{ok}/{len(use)} days (Oct 7 - Dec 31 2008, cash flow {cf}) within input rounding; off: {', '.join(off)}; "
            "before Oct 7 Damodaran used a different base cash flow not stated in the file")


def _intrinsic_file(path: Path, openpyxl) -> dict:
    """Damodaran's Table1 (same columns A-L in all vintages): year, T.Bond, inflation, real GDP growth, intrinsic, Fed effect,
    ratio, Fed funds, 10-yr avg inflation, 10-yr avg growth, smoothed intrinsic, T.Bond - smoothed."""
    ws = openpyxl.load_workbook(path, data_only=True)["Sheet1"]
    rows = []
    for r in range(2, ws.max_row + 1):
        y = ws.cell(r, 1).value
        if not (isinstance(y, int) or hasattr(y, "year")):
            break
        g = lambda c: ws.cell(r, c).value
        rows.append({"year": y if isinstance(y, int) else y.strftime("%Y-%m"), "tbond": g(2), "inflation": g(3),
                     "real_growth": g(4), "fed_funds": g(8), "_intrinsic": g(5), "_fed": g(6), "_smoothed": g(11), "_ts": g(12)})
    hdr = {str(ws.cell(1, c).value).strip(): c for c in range(ws.max_column, 0, -1) if ws.cell(1, c).value}
    yc = hdr.get("Year")
    pre = [{"year": ws.cell(r, yc).value, "inflation": ws.cell(r, yc + 1).value, "real_growth": ws.cell(r, yc + 2).value}
           for r in range(2, 11)] if yc else []
    last, hist = rows[-1], rows[-10:]
    case = {"kind": "intrinsic", "name": f"{path.stem} {last['year']}",
            "inputs": {"inflation": last["inflation"], "real_growth": last["real_growth"], "treasury_rate": last["tbond"],
                       "inflation_10y": [h["inflation"] for h in hist], "real_growth_10y": [h["real_growth"] for h in hist]},
            "expected": {"intrinsic": last["_intrinsic"], "treasury_minus_intrinsic": last["_fed"],
                         "smoothed": last["_smoothed"], "treasury_minus_smoothed": last["_ts"]}}
    data = {"from": path.name, "pre_1954": pre, "rows": [{k: v for k, v in r.items() if not k.startswith("_")} for r in rows]}
    wb = openpyxl.load_workbook(path, data_only=True)
    if "GovtBondRates" in wb.sheetnames:
        g = wb["GovtBondRates"]
        yrs = [_date(g.cell(1, c).value) if not isinstance(g.cell(1, c).value, int) else g.cell(1, c).value for c in range(2, g.max_column + 1)]
        data["govt_bond_rates"] = {"years": yrs, "rows": {g.cell(r, 1).value: [g.cell(r, c).value for c in range(2, g.max_column + 1)]
                                                          for r in range(2, g.max_row + 1) if isinstance(g.cell(r, 2).value, float)}}
    return {"case": case, "data": data}


def _sheets(path: Path):
    """(name, rows as lists of values, formulas-or-None) for .xls (xlrd) and .xlsx (openpyxl)."""
    if path.suffix == ".xls":
        import xlrd
        for s in xlrd.open_workbook(str(path)).sheets():
            yield s.name, [s.row_values(r) for r in range(s.nrows)], None
        return
    import openpyxl
    wv, wf = openpyxl.load_workbook(path, data_only=True), openpyxl.load_workbook(path)
    for ws in wv.worksheets:
        yield ws.title, [[c.value for c in r] for r in ws.iter_rows()], [[c.value for c in r] for r in wf[ws.title].iter_rows()]


def _currency_file(path: Path) -> tuple[dict, dict]:
    import re
    sov, inf = {}, {}
    for name, rows, forms in _sheets(path):
        for h, head in enumerate(rows[:8]):
            heads = [str(x or "").strip() for x in head]
            if any(x.startswith("Govt Bond Rate") for x in heads):
                ix = lambda p: max(i for i, x in enumerate(heads) if x.startswith(p))
                b, rt, s, rf = ix("Govt Bond"), ix("Bond Rating"), ix("Default Spread"), ix("Risk")
                sov = {"from": f"{path.name} [{name}]", "as_of": heads[b].replace("Govt Bond Rate", "").strip(),
                       "rows": [[r[0], r[b], r[rt], r[s], r[rf]] for r in rows[h + 1:] if r and r[0] and isinstance(r[b], float)]}
            elif heads and heads[0] == "Country" and forms:
                first = forms[h + 1][3]
                refs = list(dict.fromkeys(re.findall(r"\$([A-Z])\$(\d+)", first)))
                val = lambda ref: rows[int(ref[1]) - 1][ord(ref[0]) - 65]
                inf = {"from": path.name, "method": "compounded" if "(1+" in first else "additive",
                       "base_riskfree": val(refs[0]), "base_inflation": val(refs[1]),
                       "columns": [str(x) for x in head[:4]],
                       "rows": [[r[0], r[1], r[2], r[3]] for r in rows[h + 1:] if r and r[0]]}
    return sov, inf


def build(paths: list[str]) -> int:
    import warnings
    import openpyxl
    warnings.filterwarnings("ignore")
    src = json.loads((Path(paths[0]).resolve().parent / "_sources.json").read_text(encoding="utf-8"))
    cases, series, disc, erps, intr = [], [], list(POST_ONLY), [], []
    data = {"source": "Damodaran's implied ERP, intrinsic riskfree and currency riskfree spreadsheets",
            "url": {}, "erp_daily": {}, "currency_sovereign": {}, "currency_inflation": {}}
    for p in map(Path, paths):
        data["url"][p.name] = src.get(p.name)
        print("reading", p.name)
        if p.name in DAILY_NAMES:
            ws = openpyxl.load_workbook(p, data_only=True).worksheets[0]
            hr = next(r for r in range(1, 10) if ws.cell(r, 1).value == "Date")
            rows = _table(ws, hr, DAILY_KEYS)
            data["erp_daily"][DAILY_NAMES[p.name]] = {"from": p.name, "rows": rows}
            if p.name == "ERPbyDay2008Crisis.xlsx":
                data["erp_daily"]["crisis2008"]["cash_flow_from_oct7"] = ws["E1"].value
                series.append({"name": "ERPbyDay2008Crisis", "summary": _series_2008(rows, ws["E1"].value)})
        elif p.name.startswith("ERP"):
            e = _erp_file(p, openpyxl)
            erps.append(e)
            i = {k: v for k, v in e["inputs"].items()}
            cases.append({"kind": "erp", "name": e["name"], "inputs": i, "expected": e["erp"],
                          **({"expected_adj": e["erp_adj"]} if e["erp_adj"] is not None else {})})
            cases.append({"kind": "value", "name": f"{e['name']} index value at Damodaran's ERP {e['value_erp']:.2%}",
                          "inputs": i, "erp": e["value_erp"], "expected": e["value"]})
            mr = e["monthly_row"]
            for col, cash, sust in VARIANTS:
                vi = {**i, "cash_flow": i["_cash_choices"][cash], **({"sustainable_payout": True} if sust else {})}
                ours = implied_erp(vi)["erp"]
                c = {"kind": "erp_rounded", "name": f"{e['name']} monthly table: {col}", "inputs": vi, "expected": mr[col]}
                if abs(ours - mr[col]) <= 5e-5:
                    cases.append(c)
                else:
                    disc.append(f"{c['name']}: Damodaran {mr[col]:.2%}, the file's inputs give {ours:.3%}")
            for fn, kind, name, over, exp, erp in POSTED:
                if fn == p.name:
                    o = {k: v for k, v in over.items() if k != "cash"}
                    vi = {**i, **o, **({"cash_flow": i["_cash_choices"][over["cash"]]} if "cash" in over else {})}
                    cases.append({"kind": kind, "name": name, "inputs": vi, "expected": exp, **({"erp": erp} if erp else {})})
            if p.name in BRANCH_FILES:
                cases += _branches(p, e)
        elif p.name.lower().startswith(("intrinsic",)):
            intr.append(_intrinsic_file(p, openpyxl))
            cases.append(intr[-1]["case"])
        elif p.name in CURRENCY_VINTAGE:
            v = CURRENCY_VINTAGE[p.name]
            sov, inf = _currency_file(p)
            for table, t in (("currency_sovereign", sov), ("currency_inflation", inf)):
                if t:
                    data[table][v] = t
                    cases.append({"kind": "currency_table", "name": f"{p.name} {table.split('_')[1]} ({v})",
                                  "table": table, "vintage": v})
    if erps:
        latest = max(erps, key=lambda e: str(e["monthly_row"]["Start of month"]))
        data.update(_erp_data(latest, openpyxl))
        data["as_of"] = str(latest["monthly_row"]["Start of month"])[:10]
    if intr:
        data["intrinsic"] = max(intr, key=lambda x: str(x["data"]["rows"][-1]["year"]))["data"]
    DATA.write_text(json.dumps(data, indent=0, default=str), encoding="utf-8")
    CASES.write_text(json.dumps({"cases": cases, "series": series, "discrepancies": disc}, indent=1, default=str),
                     encoding="utf-8")
    print(f"wrote {len(cases)} cases -> {CASES}\nwrote data -> {DATA}")
    return selftest()


# ---------------------------------------------------------------- CLI
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("args", nargs="*", help="[erp|intrinsic|currency|history] inputs.json | date")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--year")
    ap.add_argument("--lookup")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--from-case")
    ap.add_argument("--build-cases", nargs="+", metavar="XLSX")
    a = ap.parse_args()
    if a.selftest:
        return selftest()
    if a.build_cases:
        return build(a.build_cases)
    if a.from_case:
        cs = [c for c in json.loads(CASES.read_text(encoding="utf-8"))["cases"]
              if c["kind"] == "erp" and c["name"].lower().startswith(a.from_case.lower())]
        if not cs:
            print("cases:", ", ".join(c["name"] for c in json.loads(CASES.read_text(encoding="utf-8"))["cases"] if c["kind"] == "erp"))
            return 1
        print(json.dumps({"_source": f"Damodaran's own inputs, {cs[0]['name']} (sheet ERP {cs[0]['expected']:.4%})",
                          **cs[0]["inputs"]}, indent=1))
        return 0
    data = json.loads(DATA.read_text(encoding="utf-8"))
    cmd, rest = (a.args[0], a.args[1:]) if a.args and a.args[0] in ("erp", "intrinsic", "currency", "history") else ("", a.args)
    inp = {}
    if rest and rest[0].endswith(".json"):
        inp = {k: v for k, v in json.loads(Path(rest[0]).read_text(encoding="utf-8")).items() if not k.startswith("_")}
        cmd = cmd or inp.pop("mode", "erp")
    if cmd == "history":
        lines = history(data, rest[0]) if rest else []
        print("\n".join(lines) or "no published ERP for that date (daily series: Sep-Dec 2008, 2020, Dec 2024-May 2025;"
                                  " monthly from Sep 2008; annual 1960-2024)")
        return 0
    if cmd == "intrinsic":
        out = intrinsic_for_year(data, a.year) if a.year else intrinsic_riskfree(**inp)
    elif cmd == "currency":
        if a.lookup:
            print("\n".join(currency_lookup(data, a.lookup)) or "not found")
            return 0
        out = currency(inp, data)
    elif cmd == "erp" and inp:
        if a.json:
            out = implied_erp(inp)
        else:
            print(erp_report(inp))
            return 0
    else:
        ap.print_help()
        return 1
    print(json.dumps(out, indent=1) if a.json else "\n".join(
        f"{k:28s} {v:.4%}" if isinstance(v, float) and abs(v) < 5 else f"{k:28s} {v}" for k, v in out.items()))
    return 0


if __name__ == "__main__":
    sys.exit(main())
