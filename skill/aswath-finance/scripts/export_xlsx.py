"""Export a ginzu.py valuation to a shareable Excel workbook (investment-banking conventions, live formulas).

Every calculation cell is an Excel formula that mirrors ginzu.py / cost_of_capital.py line by line, so the
workbook recalculates when an input changes; a Checks tab ties it to the engine's own output. No circular
references (employee options and IPO shares at value are solved by an unrolled iteration on a Dilution tab);
the sensitivity and scenario data tables are native Excel What-If tables that Excel computes on open.

Usage:
  python export_xlsx.py inputs.json [-o out.xlsx] [--verify]   same inputs file as ginzu.py
  python export_xlsx.py --selftest     rebuild all 45 calibration/branch cases + 21 cost-of-capital cases and
                                       recalculate each workbook with pycel: must equal the engine

--verify (Windows + Excel) recalculates in Excel, saves the cached values and prints the Checks result.

Optional presentation block in the inputs JSON (ignored by ginzu.py, like every "_" key):
  "_workbook": {"company": "...", "ticker": "...", "valuation_date": "2026-09-25", "currency": "$",
                "units": "$ in millions except per share",
                "market": {"week52_low": 0, "week52_high": 0, "analyst_target": 0},   football field only
                "scenarios": [{"name": "Target margin 30%", "set": {"target_margin": 0.30}}],   scalar inputs only
                "notes": ["judgment call ..."]}
  "_sources": {"<input key>": "where the number comes from", ...}   shown next to each input and on Sources
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import DataTableFormula

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from cost_of_capital import industry, load as load_data  # noqa: E402
from ginzu import breakeven, segment_paths, value  # noqa: E402
from xlsx_common import (BOX, F, FILL_FLEX, FILL_LIGHT, FILL_NAVY, FONT, FX, GREEN_FILL, ITER, LIGHT, LINK, M,  # noqa: E402
                         NOTE, NUM2, NUM4, PCT, PCT2, RED_FILL, TOP, col, font, hdr, label, printable, put, section,
                         title, widths)
from xlsx_wacc import WaccSheet  # noqa: E402

# Ginzu defaults for inputs the Inputs tab always shows (so scenarios can vary them even when the JSON omits them)
DEFAULTS = {"nol": 0.0, "ipo_proceeds": 0.0, "ipo_shares_at_value": False, "keep_effective_tax": False, "reinvestment_lag": 1,
            "failure_probability": 0.0, "failure_proceeds_basis": "V", "failure_proceeds_pct": 0.0}
SEGMENT_FED = {"growth_next_year", "growth_years_2_5", "margin_next_year", "target_margin", "margin_convergence_year",
               "riskfree_rate", "riskfree_after_10", "terminal_growth"}  # inputs ginzu.segment_paths reads


class Book:
    """One valuation -> one workbook. `build()` writes it and returns the key cell addresses."""

    def __init__(self, raw: dict, meta: dict | None = None):
        self.meta = {**raw.get("_workbook", {}), **(meta or {})}
        self.sources = raw.get("_sources", {})
        self.data = load_data()
        inp = {k: v for k, v in raw.items() if not k.startswith("_")}
        inp.setdefault("mature_erp", self.data["mature_erp"])  # ginzu.value() default
        self.notes = list(self.meta.get("notes", []))
        if inp.get("segments") and not inp.get("revenue_growth_path"):
            fed = set(self.meta.get("scenarios") and [k for x in self.meta["scenarios"] for k in x["set"]] or []) & SEGMENT_FED
            if fed:
                raise ValueError(f"scenario keys {sorted(fed)} feed the segments build, which the workbook carries as fixed "
                                 "paths: vary the segments in the JSON and export once per case instead")
            rf10 = inp.get("riskfree_after_10", inp["riskfree_rate"])
            gp, mp = segment_paths(inp, inp.get("terminal_growth", rf10))
            inp = {**inp, "revenue_growth_path": gp, "margin_path": mp}
            self.notes.append("Revenue growth and margin paths are derived from the segments block (ginzu.segment_paths, "
                              "market x share); edit the segments in the JSON and re-export to change them. The growth, "
                              "margin and riskfree flexes move these fixed paths, not the segment build.")
        self.scen = self.meta.get("scenarios", [])
        self.skeys = list(dict.fromkeys(k for s in self.scen for k in s["set"]))
        inp = {**{k: DEFAULTS[k] for k in self.skeys if k in DEFAULTS}, **inp}  # value-neutral: the Ginzu defaults
        for k in self.skeys:
            if k not in inp or isinstance(inp[k], (dict, list)):
                raise ValueError(f"scenario key '{k}' must be a scalar input present in the base case")
        self.inp, self.out = inp, value(inp)
        self.cur = self.meta.get("currency", "$")
        self.PS = f'"{self.cur}"#,##0.00_);("{self.cur}"#,##0.00)'
        self.units = self.meta.get("units", f"{self.cur} in millions except per share")
        self.eng_scen = [value({**inp, **s["set"]})["value_per_share"] for s in self.scen]
        self.dil = bool(inp.get("options") or inp.get("ipo_shares_at_value"))
        self.R: dict[str, int] = {}
        self.labels: dict[str, str] = {}

    # ------------------------------------------------------------ references
    def ref(self, key):
        return f"Inputs!${'D' if key in self.skeys else 'C'}${self.R[key]}"

    def has(self, key):
        return key in self.R

    def path(self, key, t):
        return f"Inputs!${get_column_letter(2 + t)}${self.R['path:' + key]}"

    def src(self, key, default=""):
        s = self.sources
        if isinstance(s, dict):
            return s[key] if key in s else next((v for k, v in s.items() if k.split(" ")[0].split("(")[0] == key), default)
        return default

    # ------------------------------------------------------------ build
    def build(self, path: Path, tables: bool = True) -> dict:
        self.tables = tables
        self.wb = wb = Workbook()
        names = ["Cover", "Summary", "Inputs"] + (["Scenarios"] if self.scen else []) + [
            "Sensitivity", "WACC", "DCF"] + (["Dilution"] if self.dil else []) + ["Diagnostics", "Checks", "Sources"]
        wb.active.title = names[0]
        self.ws = {n: (wb.active if n == "Cover" else wb.create_sheet(n)) for n in names}
        self.inputs_sheet()
        self.flex_cells()
        if self.scen:
            self.scenarios_sheet()
        self.wacc_cell = WaccSheet(self).build(self.ws["WACC"])
        self.dcf_sheet()
        if self.dil:
            self.dilution_sheet()
        self.sensitivity_tables()
        self.diagnostics_sheet()
        self.summary_sheet()
        self.checks_sheet()
        self.sources_sheet()
        self.cover_sheet()
        for nm, ref in (("ValuePerShare", "DCF!$C$49"), ("SharePrice", self.ref("price")),
                        ("OperatingAssets", "DCF!$C$39"), ("WACC_Initial", self.wacc_cell)):
            wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
        for ws in wb.worksheets:  # unstyled cells default to Calibri
            for row in ws.iter_rows():
                for c in row:
                    if c.font is None or c.font.name != F:
                        c.font = font(10, bold=bool(c.font and c.font.bold), italic=bool(c.font and c.font.italic),
                                      color=c.font.color if c.font else None)
        wb.calculation.fullCalcOnLoad = True
        wb.active = 0
        wb.save(path)
        return {"value": "DCF!C49", "wacc": self.wacc_cell.replace("$", ""), "checks": self.checks_cell}

    def eval_order(self) -> list[str]:
        """Model cells in dependency order: lets a lazy recursive evaluator (pycel) walk the long year-by-year and
        iteration chains one step at a time instead of recursing through all of them at once."""
        wacc = [f"WACC!{c.coordinate}" for row in self.ws["WACC"].iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
        rd = [f"DCF!C{r}" for r in (55, 56, 57)] if self.inp.get("rd") else []
        years = [f"DCF!{col(t)}{r}" for t in range(12) for r in range(7, 24)]
        dil = [f"Dilution!{c}{r}" for r in range(18, 19 + ITER) for c in "CDEFGHIJ"] if self.dil else []
        return (wacc + rd + [f"DCF!C{r}" for r in range(26, 33)] + years + [f"DCF!C{r}" for r in range(33, 47)] + dil
                + [f"DCF!C{r}" for r in range(47, 50)])

    # ------------------------------------------------------------ inputs
    def block(self, ws, heading, rows):
        section(ws, self.r, heading)
        for key, lab, v, fmt, *note in rows:
            self.r += 1
            label(ws, f"B{self.r}", lab)
            put(ws, f"C{self.r}", int(v) if isinstance(v, bool) else v, fmt)
            c = put(ws, f"N{self.r}", note[0] if note else self.src(key), kind="label")
            c.font, c.alignment = NOTE, Alignment(wrap_text=True, vertical="top")
            self.R[key], self.labels[key] = self.r, lab
        self.r += 2

    def inputs_sheet(self):
        ws, i = self.ws["Inputs"], self.inp
        title(ws, "Inputs & Assumptions", f"{self.units} | Blue = hard-coded input (source in column N)"
              + (" | Column D = live value of the active scenario" if self.skeys else ""))
        widths(ws, {"B": 46, **{get_column_letter(c): 11 for c in range(3, 13)}, "M": 2, "N": 80})
        if self.skeys:
            hdr(ws, 4, {"C": "Base", "D": "Live"})
        self.r, mk, dflt = 5, self.meta.get("market", {}), "Ginzu default"
        self.block(ws, "Market data", [
            ("price", "Share price", i["price"], self.PS), ("shares_outstanding", "Shares outstanding", i["shares_outstanding"], M),
            *[(f"m:{k}", lab, mk[k], self.PS, "Market data (third party)") for k, lab in (
                ("week52_low", "52-week low"), ("week52_high", "52-week high"),
                ("analyst_target", "Analyst average price target (third party)")) if k in mk]])
        base = [("revenues", "Revenues (base year)", M), ("ebit", "Operating income (EBIT, base year)", M),
                ("book_equity", "Book value of equity", M), ("book_debt", "Book value of debt (incl. leases)", M),
                ("cash", "Cash & marketable securities", M), ("non_operating_assets", "Non-operating assets (cross holdings)", M),
                ("minority_interests", "Minority interests", M)]
        rows = [(k, lab, i[k], fmt) for k, lab, fmt in base]
        rows += [(k, lab, i[k], fmt) for k, lab, fmt in (("revenues_last_year", "Revenues, prior year", M),
                                                          ("trapped_cash", "Trapped cash", M),
                                                          ("trapped_cash_tax_rate", "Tax rate already paid on trapped cash", PCT)) if k in i]
        rows += [("nol", "Net operating loss carried forward", i.get("nol", 0), M, self.src("nol", dflt + " (none)")),
                 ("ipo_proceeds", "IPO proceeds retained by the firm", i.get("ipo_proceeds", 0), M, self.src("ipo_proceeds", "None")),
                 ("ipo_shares_at_value", "New IPO shares issued at estimated value (1 = yes)", bool(i.get("ipo_shares_at_value")), "0",
                  self.src("ipo_shares_at_value", "No"))]
        self.block(ws, "Base year (trailing twelve months)", rows)
        drv = [(k, lab, i[k], fmt) for k, lab, fmt in (
            ("growth_next_year", "Revenue growth, next year", PCT), ("growth_years_2_5", "Revenue growth, years 2-5 (CAGR)", PCT),
            ("margin_next_year", "Operating margin, next year", PCT), ("target_margin", "Target pre-tax operating margin", PCT),
            ("margin_convergence_year", "Year of margin convergence", "0"), ("sales_to_capital_1_5", "Sales to capital, years 1-5", NUM2),
            ("sales_to_capital_6_10", "Sales to capital, years 6-10", NUM2)) if k in i]
        self.block(ws, "Value drivers (the story)", drv)
        opt = lambda k, lab, fmt, note: (k, lab, i[k], fmt) if k in i else None  # noqa: E731
        rates = [("effective_tax_rate", "Effective tax rate", i["effective_tax_rate"], PCT),
                 ("marginal_tax_rate", "Marginal tax rate", i["marginal_tax_rate"], PCT),
                 ("keep_effective_tax", "Keep effective tax rate forever (1 = yes)", bool(i.get("keep_effective_tax")), "0",
                  self.src("keep_effective_tax", dflt + " (tax drifts to marginal)")),
                 ("riskfree_rate", "Riskfree rate", i["riskfree_rate"], PCT2),
                 opt("riskfree_after_10", "Riskfree rate after year 10", PCT2, ""),
                 ("mature_erp", "Mature market equity risk premium", i["mature_erp"], PCT2,
                  self.src("mature_erp", f"Damodaran data ({self.data['updated']})")),
                 opt("initial_cost_of_capital", "Initial cost of capital (direct input)", PCT2, "")
                 if not i.get("cost_of_capital") else None,
                 opt("terminal_cost_of_capital", "Terminal cost of capital (override)", PCT2, ""),
                 opt("terminal_growth", "Terminal growth rate (override)", PCT2, ""),
                 opt("terminal_roic", "Terminal return on capital (override)", PCT, ""),
                 ("reinvestment_lag", "Reinvestment lag (years, 0-3)", i.get("reinvestment_lag", 1), "0",
                  self.src("reinvestment_lag", dflt + " (1 year)")),
                 ("failure_probability", "Probability of failure", i.get("failure_probability", 0), PCT,
                  self.src("failure_probability", dflt + " (no failure)")),
                 ("failure_proceeds_basis", 'Failure proceeds tied to ("B" book / "V" fair value)',
                  i.get("failure_proceeds_basis", "V"), None, self.src("failure_proceeds_basis", dflt)),
                 ("failure_proceeds_pct", "Failure proceeds, % of that basis", i.get("failure_proceeds_pct", 0), PCT,
                  self.src("failure_proceeds_pct", dflt))]
        self.block(ws, "Rates, terminal value and default-assumption overrides", [x for x in rates if x])
        paths = [k for k in ("revenue_growth_path", "margin_path", "reinvestment_path") if k in i]
        if paths:
            section(ws, self.r, "Driver paths (replace the generated paths)")
            hdr(ws, self.r + 1, {get_column_letter(2 + t): f"Year {t}" for t in range(1, 11)})
            self.r += 1
            for k in paths:
                self.r += 1
                label(ws, f"B{self.r}", {"revenue_growth_path": "Revenue growth", "margin_path": "Operating margin",
                                         "reinvestment_path": "Reinvestment"}[k])
                for t, v in enumerate(i[k], start=1):
                    put(ws, f"{get_column_letter(2 + t)}{self.r}", v, M if k == "reinvestment_path" else PCT)
                c = put(ws, f"N{self.r}", self.src(k), kind="label")
                c.font, c.alignment = NOTE, Alignment(wrap_text=True, vertical="top")
                self.R["path:" + k] = self.r
            self.r += 2
        if i.get("rd"):
            rd = i["rd"]
            self.block(ws, "R&D capitalization (Ginzu R&D converter)", [
                ("rd.current", "R&D expense, current year", rd["current"], M, self.src("rd")),
                ("rd.life", "Amortizable life (years)", rd["life"], "0"),
                *[(f"rd.past{k}", f"R&D expense, year -{k}", v, M) for k, v in enumerate(rd.get("past", []), start=1)]])
        if i.get("options"):
            o = i["options"]
            self.block(ws, "Employee options (dilution-adjusted Black-Scholes)", [
                ("opt.number", "Options outstanding", o["number"], M, self.src("options")), ("opt.strike", "Average strike price", o["strike"], self.PS),
                ("opt.maturity", "Average maturity (years)", o["maturity"], NUM2), ("opt.volatility", "Standard deviation of stock", o["volatility"], PCT),
                ("opt.price_basis", 'Value options at ("price" or "value")', o.get("price_basis", "price"), None)])
        if i.get("industry_us"):
            ind = industry(self.data, i["industry_us"], "US")
            note = f"Damodaran industry averages, US {i['industry_us']} ({self.data['updated']})"
            self.block(ws, f"Industry benchmarks for Diagnostics: US {i['industry_us']}", [
                ("ind.growth", "Revenue growth (5-year average)", ind["Annual Average Revenue growth - Last 5 years"], PCT, note),
                ("ind.margin", "Pre-tax operating margin (lease & R&D adjusted)", ind["Pre-tax Operating Margin (Lease & R&D adjusted)"], PCT, note),
                ("ind.s2c", "Sales / capital", ind["Sales/Capital"], NUM2, note), ("ind.coc", "Cost of capital", ind["Cost of capital"], PCT, note)])
        for k in self.skeys:
            ws[f"D{self.R[k]}"].number_format = ws[f"C{self.R[k]}"].number_format
        ws.freeze_panes = "C5"
        printable(ws, one_page=True)

    def flex_cells(self):
        ws = self.ws["Sensitivity"]
        title(ws, "Sensitivity Analysis", "Live Excel data tables on value per share. Flex cells (yellow) must stay neutral "
              "for the base case; the data tables vary them.")
        widths(ws, {"B": 44, **{c: 13 for c in "CDEFGH"}, "I": 36})
        section(ws, 4, "Flex cells (neutral values = base case)", last="I")
        gl = "Change in revenue growth, years 1-5" + (" (path)" if "revenue_growth_path" in self.inp else "")
        ml = "Parallel change in margin path" if "margin_path" in self.inp else "Change in target operating margin"
        for r, (lab, fmt) in enumerate((("Change in initial cost of capital", PCT2), ("Change in terminal cost of capital", PCT2),
                                        (gl, PCT), (ml, PCT), ("Change in riskfree rate", PCT2)), start=5):
            label(ws, f"B{r}", lab)
            put(ws, f"C{r}", 0, fmt).fill = FILL_FLEX

    def scenarios_sheet(self):
        ws, n, keys = self.ws["Scenarios"], len(self.scen) + 1, self.skeys
        title(ws, "Scenario Analysis", "Set the active scenario in D5: the whole model runs on it (Inputs column D). "
              "The value column is a live data table over all scenarios.")
        vc = get_column_letter(4 + len(keys))
        pc, ic, oc = (get_column_letter(5 + len(keys) + k) for k in (0, 2, 3))
        widths(ws, {"B": 5, "C": 48, **{get_column_letter(4 + k): 13 for k in range(len(keys) + 2)}, ic: 6, oc: 12})
        label(ws, "C5", f"Active scenario (1-{n})", bold=True)
        s = put(ws, "D5", 1, "0")
        s.fill, s.border = FILL_FLEX, BOX
        put(ws, "E5", f"=INDEX($C$8:$C${7 + n},$D$5)")
        hdr(ws, 7, {"B": "#", "C": "Scenario", **{get_column_letter(4 + k): self.labels.get(key, key) for k, key in enumerate(keys)},
                    vc: "Value / share", pc: "vs price"})
        ws.row_dimensions[7].height = 66
        for k, sc in enumerate([{"name": "Base case", "set": {}}] + self.scen, start=1):
            r = 7 + k
            put(ws, f"B{r}", k, "0", kind="label")
            put(ws, f"C{r}", sc["name"], kind="label")
            for j, key in enumerate(keys):
                c, fmt = get_column_letter(4 + j), self.ws["Inputs"][f"C{self.R[key]}"].number_format
                v = f"=Inputs!$C${self.R[key]}" if k == 1 else sc["set"].get(key, f"={c}$8")
                put(ws, f"{c}{r}", v, fmt)
            put(ws, f"{vc}{r}", f"={oc}{r}", self.PS)
            put(ws, f"{pc}{r}", f"={vc}{r}/SharePrice-1", PCT)
            put(ws, f"{ic}{r}", k, "0", kind="label")
        put(ws, f"{ic}6", "Data table (What-If; do not edit)", kind="label").font = NOTE
        put(ws, f"{oc}7", "=ValuePerShare", self.PS)
        if self.tables:
            ws[f"{oc}8"] = DataTableFormula(ref=f"{oc}8:{oc}{7 + n}", r1="D5")
        for r in range(8, 8 + n):
            ws[f"{oc}{r}"].number_format = self.PS
        for j, key in enumerate(keys):
            c = get_column_letter(4 + j)
            self.ws["Inputs"][f"D{self.R[key]}"] = f"=INDEX(Scenarios!${c}$8:${c}${7 + n},Scenarios!$D$5)"
            self.ws["Inputs"][f"D{self.R[key]}"].font = FONT
        self.scen_rows = (8, 7 + n, vc)
        ws.freeze_panes = "D8"
        printable(ws)

    # ------------------------------------------------------------ DCF
    def dcf_sheet(self):
        ws, i, R = self.ws["DCF"], self.inp, self.ref
        title(ws, "Discounted Cash Flow (FCFF)", "Damodaran fcffsimpleginzu mechanics: tax drifts to marginal in years 6-10, NOL "
              f"shield, cost of capital drifts to terminal, terminal reinvestment = g / ROIC. {self.units}.")
        widths(ws, {"B": 46, **{col(t): 12 for t in range(12)}})
        hdr(ws, 4, {"C": "Base year", **{col(t): f"Year {t}" for t in range(1, 11)}, "N": "Terminal"})
        label(ws, "B5", "Year (t)")
        for t in range(12):
            put(ws, f"{col(t)}5", t, "0", kind="label").font = NOTE
        path_r = "reinvestment_path" in i
        for r, t in {7: "Revenue growth", 8: "Revenue", 9: "Operating margin", 10: "Operating income (EBIT)", 11: "Tax rate",
                     12: "Net operating loss carried forward", 13: "After-tax operating income",
                     14: "" if path_r else "Sales to capital", 15: "Less: reinvestment", 16: "Free cash flow to firm (FCFF)",
                     18: "Cost of capital", 19: "Cumulative discount factor", 20: "Present value of FCFF",
                     22: "Invested capital", 23: "Return on invested capital"}.items():
            label(ws, f"B{r}", t, bold=r in (8, 16))
        TG, RT, RO = "$C$31", "$C$30", "$C$32"
        tgt = f"({R('target_margin')}+{FX['d_margin']})" if "target_margin" in i else None
        tterm = f"IF({R('keep_effective_tax')}=1,{R('effective_tax_rate')},{R('marginal_tax_rate')})"
        put(ws, "C8", f"={R('revenues')}", M, bold=True)
        put(ws, "C9", "=C10/C8", PCT)
        put(ws, "C10", f"={R('ebit')}" + ("+C57" if i.get("rd") else ""), M)
        put(ws, "C11", f"={R('effective_tax_rate')}", PCT)
        put(ws, "C12", f"={R('nol')}", M)
        put(ws, "C13", "=IF(C10>0,C10*(1-C11),C10)", M)
        put(ws, "C22", f"={R('book_equity')}+{R('book_debt')}-{R('cash')}" + ("+C55" if i.get("rd") else ""), M)
        put(ws, "C23", "=C13/C22", PCT)
        for t in range(1, 11):
            c, p = col(t), col(t - 1)
            if "revenue_growth_path" in i:
                g = f"={self.path('revenue_growth_path', t)}" + (f"+{FX['d_growth']}" if t <= 5 else "")
            elif t <= 5:
                g = f"={R('growth_next_year' if t == 1 else 'growth_years_2_5')}+{FX['d_growth']}"
            else:
                g = f"=$H$7-($H$7-{TG})/5*{t - 5}"
            put(ws, f"{c}7", g, PCT)
            put(ws, f"{c}8", f"={p}8*(1+{c}7)", M, bold=True)
            if "margin_path" in i:
                m = f"={self.path('margin_path', t)}+{FX['d_margin']}"
            elif t == 1:
                m = f"={R('margin_next_year')}"
            else:
                cv = R("margin_convergence_year")
                m = f"=IF({c}$5>{cv},{tgt},{tgt}-({tgt}-$D$9)/{cv}*({cv}-{c}$5))"
            put(ws, f"{c}9", m, PCT)
            put(ws, f"{c}10", f"={c}9*{c}8", M)
            put(ws, f"{c}11", f"={p}11" if t <= 5 else f"={p}11+({tterm}-$H$11)/5", PCT)
            put(ws, f"{c}12", f"=IF({c}10<0,{p}12-{c}10,IF({p}12>{c}10,{p}12-{c}10,0))", M)
            put(ws, f"{c}13", f"=IF({c}10>0,IF({c}10<{p}12,{c}10,{c}10-({c}10-{p}12)*{c}11),{c}10)", M)
            if path_r:
                put(ws, f"{c}15", f"={self.path('reinvestment_path', t)}", M)
            else:
                put(ws, f"{c}14", f"={R('sales_to_capital_1_5' if t <= 5 else 'sales_to_capital_6_10')}", NUM2)
                rv = lambda k: f"{col(k)}8"  # noqa: E731
                gt = f"$N$8*$N$7/{c}14"  # terminal-year revenue growth (lags 2-3 near the end)
                lags = [f"({rv(t)}-{rv(t - 1)})/{c}14", f"({rv(t + 1)}-{rv(t)})/{c}14",
                        f"({rv(t + 2)}-{rv(t + 1)})/{c}14" if t <= 9 else gt,
                        f"({rv(t + 3)}-{rv(t + 2)})/{c}14" if t <= 8 else (gt if t == 9 else f"L15*(1+{TG})")]
                put(ws, f"{c}15", f"=CHOOSE({R('reinvestment_lag')}+1,{','.join(lags)})", M)
            put(ws, f"{c}16", f"={c}13-{c}15", M, bold=True).border = TOP
            put(ws, f"{c}18", "=$C$27" if t <= 5 else f"={p}18-($H$18-{RT})/5", PCT2)
            put(ws, f"{c}19", "=1/(1+D18)" if t == 1 else f"={p}19/(1+{c}18)", NUM4)
            put(ws, f"{c}20", f"={c}16*{c}19", M)
            put(ws, f"{c}22", f"={p}22+{c}15", M)
            put(ws, f"{c}23", f"={c}13/{p}22", PCT)
        for r, f, fmt in ((7, f"={TG}", PCT), (8, "=M8*(1+N7)", M), (9, "=M9", PCT), (10, "=N9*N8", M), (11, f"={tterm}", PCT),
                          (13, "=N10*(1-N11)", M), (15, f"=IF({TG}>0,{TG}/{RO}*N13,0)", M), (16, "=N13-N15", M),
                          (18, f"={RT}", PCT2), (23, f"={RO}", PCT)):
            put(ws, f"N{r}", f, fmt, bold=r in (8, 16))
        self.valuation_block(ws)
        ws.freeze_panes = "C6"
        printable(ws, one_page=True)

    def valuation_block(self, ws):
        i, R = self.inp, self.ref
        section(ws, 25, "Valuation")
        opt = lambda k, f, d: f"={R(k)}{f}" if k in i else d  # noqa: E731
        opts, shares = ("=Dilution!$I$" + str(18 + ITER), "=Dilution!$C$" + str(18 + ITER)) if self.dil else (0, f"={R('shares_outstanding')}")
        trapped = f"={R('trapped_cash')}*({R('marginal_tax_rate')}-{R('trapped_cash_tax_rate')})" if "trapped_cash" in i else 0
        rows = [
            (26, "Initial cost of capital (WACC tab)", f"={self.wacc_cell}", PCT2),
            (27, "Initial cost of capital (incl. flex)", f"=C26+{FX['d_init']}", PCT2),
            (28, "Riskfree rate after year 10", opt("riskfree_after_10", f"+{FX['d_rf']}", "=WACC!$C$5"), PCT2),
            (29, "Mature market ERP", f"={R('mature_erp')}", PCT2),
            (30, "Terminal cost of capital", opt("terminal_cost_of_capital", f"+{FX['d_rf']}+{FX['d_term']}", f"=C28+C29+{FX['d_term']}"), PCT2),
            (31, "Terminal growth rate", opt("terminal_growth", "", "=C28"), PCT2),
            (32, "Terminal return on capital", opt("terminal_roic", "", "=C30"), PCT),
            (33, "Terminal value  = FCFF(t+1) / (r - g)", "=N16/(C30-C31)", M),
            (34, "PV of terminal value", "=C33*M19", M), (35, "PV of FCFF, years 1-10", "=SUM(D20:M20)", M),
            (36, "Sum of present values", "=C34+C35", M), (37, "Probability of failure", f"={R('failure_probability')}", PCT),
            (38, "Proceeds if the firm fails",
             f"=IF(EXACT({R('failure_proceeds_basis')},\"B\"),{R('book_equity')}+{R('book_debt')},C36)*{R('failure_proceeds_pct')}", M),
            (39, "Value of operating assets", "=C36*(1-C37)+C38*C37", M), (40, "Less: debt", f"={R('book_debt')}", M),
            (41, "Less: minority interests", f"={R('minority_interests')}", M), (42, "Plus: cash", f"={R('cash')}", M),
            (43, "Less: tax on repatriating trapped cash", trapped, M),
            (44, "Plus: non-operating assets", f"={R('non_operating_assets')}", M),
            (45, "Plus: IPO proceeds retained", f"={R('ipo_proceeds')}", M),
            (46, "Value of equity", "=C39-C40-C41+C42-C43+C44+C45", M),
            (47, "Less: value of employee options", opts, M), (48, "Shares outstanding (incl. IPO shares at value)", shares, M),
            (49, "Estimated value per share", "=(C46-C47)/C48", self.PS), (50, "Current share price", f"={R('price')}", self.PS),
            (51, "Price as % of value", "=C50/C49", PCT), (52, "Upside / (downside) to value", "=C49/C50-1", PCT)]
        for r, lab, f, fmt in rows:
            label(ws, f"B{r}", lab, bold=r in (39, 46, 49))
            put(ws, f"C{r}", f, fmt, bold=r in (39, 46, 49))
        for r in (39, 46):
            ws[f"C{r}"].border = TOP
        ws["C49"].border, ws["C49"].fill = BOX, FILL_LIGHT
        if i.get("rd"):
            section(ws, 54, "R&D capitalization (Ginzu R&D converter)")
            life, past = R("rd.life"), [R(f"rd.past{k}") for k in range(1, len(i["rd"].get("past", [])) + 1)]
            asset = "".join(f"+{p}*MAX(0,{life}-{k})/{life}" for k, p in enumerate(past, start=1))
            amort = "+".join(f"{p}*({k}<={life})" for k, p in enumerate(past, start=1)) or "0"
            for r, lab, f in ((55, "Value of research asset", f"={R('rd.current')}{asset}"),
                              (56, "Amortization of research asset (this year)", f"=({amort})/{life}"),
                              (57, "Adjustment to operating income (R&D - amortization)", f"={R('rd.current')}-C56")):
                label(ws, f"B{r}", lab)
                put(ws, f"C{r}", f, M)

    def dilution_sheet(self):
        ws, i, R = self.ws["Dilution"], self.inp, self.ref
        title(ws, "Share Count & Employee Options", "Solves the Ginzu circular reference (options valued at estimated value, "
              f"IPO shares issued at value) by {ITER} explicit fixed-point steps: no iterative calculation needed.")
        widths(ws, {"B": 44, **{c: 13 for c in "CDEFGHIJ"}})
        o = bool(i.get("options"))
        params = [("Value of equity (DCF)", "=DCF!$C$46", M), ("Shares outstanding", f"={R('shares_outstanding')}", M),
                  ("IPO proceeds retained", f"={R('ipo_proceeds')}", M), ("IPO shares issued at value (1 = yes)", f"={R('ipo_shares_at_value')}", "0"),
                  ("Riskfree rate (live)", "=WACC!$C$5", PCT2), ("Share price", f"={R('price')}", self.PS)]
        if o:
            params += [("Options outstanding", f"={R('opt.number')}", M), ("Strike", f"={R('opt.strike')}", self.PS),
                       ("Maturity (years)", f"={R('opt.maturity')}", NUM2), ("Volatility", f"={R('opt.volatility')}", PCT),
                       ("Value options at", f"={R('opt.price_basis')}", None)]
        for r, (lab, f, fmt) in enumerate(params, start=5):
            label(ws, f"B{r}", lab)
            put(ws, f"C{r}", f, fmt)
        hdr(ws, 17, {"B": "Step", "C": "Shares", "D": "Stock price used", "E": "Diluted price", "F": "d1", "G": "d2",
                     "H": "Value per option", "I": "Options value", "J": "Value per share"})
        put(ws, "B18", 0, "0", kind="label")
        put(ws, "C18", "=$C$6", M)
        put(ws, "H18", 0, NUM4)
        put(ws, "I18", 0, M)
        put(ws, "J18", "=$C$5/$C$6", self.PS)
        for r in range(19, 19 + ITER):
            put(ws, f"B{r}", r - 18, "0", kind="label")
            put(ws, f"C{r}", f"=$C$6+IF($C$8=1,$C$7/J{r - 1},0)", M)
            if o:
                put(ws, f"D{r}", f'=IF(EXACT($C$15,"value"),J{r - 1},$C$10)', self.PS)
                put(ws, f"E{r}", f"=(D{r}*$C$6+H{r - 1}*$C$11)/($C$6+$C$11)", self.PS)
                put(ws, f"F{r}", f"=(LN(E{r}/$C$12)+($C$9+$C$14^2/2)*$C$13)/($C$14*SQRT($C$13))", NUM4)
                put(ws, f"G{r}", f"=F{r}-$C$14*SQRT($C$13)", NUM4)
                put(ws, f"H{r}", f"=E{r}*NORMSDIST(F{r})-$C$12*EXP(-$C$9*$C$13)*NORMSDIST(G{r})", NUM4)
                put(ws, f"I{r}", f"=H{r}*$C$11", M)
            else:
                put(ws, f"I{r}", 0, M)
            put(ws, f"J{r}", f"=($C$5-I{r})/C{r}", self.PS)
        printable(ws, landscape=False)

    # ------------------------------------------------------------ sensitivity tables
    def sensitivity_tables(self):
        ws = self.ws["Sensitivity"]

        def grid(r0, heading, note, rows, cols, rfmt, cfmt, r_in, c_in):
            section(ws, r0, heading, last="I")
            put(ws, f"B{r0 + 1}", note, kind="label").font = NOTE
            put(ws, f"B{r0 + 2}", "=ValuePerShare", self.PS)
            for c, v in zip("CDEFG", cols):
                put(ws, f"{c}{r0 + 2}", v, cfmt).fill = FILL_LIGHT
            for k, v in enumerate(rows, start=r0 + 3):
                put(ws, f"B{k}", v, rfmt).fill = FILL_LIGHT
                for c in "CDEFG":
                    ws[f"{c}{k}"].number_format = self.PS
            ref = f"C{r0 + 3}:G{r0 + 2 + len(rows)}"
            if self.tables:
                ws[f"C{r0 + 3}"] = DataTableFormula(ref=ref, dt2D=True, r1=c_in, r2=r_in)
            ws.conditional_formatting.add(ref, CellIsRule(operator="greaterThanOrEqual", formula=["$I$15"], fill=GREEN_FILL))
            return ref, f"Sensitivity!{'CDEFG'[cols.index(0)]}{r0 + 3 + rows.index(0)}"

        self.T1, self.C1 = grid(11, "Value per share: initial vs terminal cost of capital (change from base)",
                       "Rows: change in initial cost of capital | Columns: change in terminal cost of capital",
                       (-0.02, -0.015, -0.01, -0.005, 0, 0.005, 0.01, 0.015, 0.02), (-0.01, -0.005, 0, 0.005, 0.01), PCT2, PCT2, "C5", "C6")
        self.T2, self.C2 = grid(25, "Value per share: revenue growth vs operating margin (change from base)",
                       f"Rows: {ws['B7'].value.lower()} | Columns: {ws['B8'].value.lower()}",
                       (-0.10, -0.05, -0.025, 0, 0.025, 0.05, 0.10), (-0.05, -0.025, 0, 0.025, 0.05), PCT, PCT, "C7", "C8")
        section(ws, 37, "Riskfree rate: value per share and cost of capital", last="I")
        for c, t in (("B", "Change in riskfree rate"), ("C", "Value / share"), ("D", "Initial WACC")):
            label(ws, f"{c}38", t, bold=True)
        put(ws, "C39", "=ValuePerShare", self.PS)
        put(ws, "D39", "=WACC_Initial", PCT2)
        for k, v in enumerate((-0.015, -0.01, -0.005, 0, 0.005, 0.01, 0.015), start=40):
            put(ws, f"B{k}", v, PCT2).fill = FILL_LIGHT
            ws[f"C{k}"].number_format, ws[f"D{k}"].number_format = self.PS, PCT2
        if self.tables:
            ws["C40"] = DataTableFormula(ref="C40:D46", r1="C9")
        ws.conditional_formatting.add("C40:C46", CellIsRule(operator="greaterThanOrEqual", formula=["$I$15"], fill=GREEN_FILL))
        put(ws, "I13", "Shaded green: value at or above the current share price", kind="label").font = NOTE
        label(ws, "I14", "Current share price", bold=True)
        put(ws, "I15", "=SharePrice", self.PS)
        printable(ws, one_page=True)

    # ------------------------------------------------------------ diagnostics
    def diagnostics_sheet(self):
        ws, R = self.ws["Diagnostics"], self.ref
        title(ws, "Diagnostics", "The Ginzu 'Diagnostics' sheet: the forecast next to industry averages, with the questions Damodaran suggests asking")
        widths(ws, {"B": 30, **{c: 13 for c in "CDEFG"}, "H": 70})
        ind = lambda k: f"={R(k)}" if self.has(k) else "n/a"  # noqa: E731
        last = f"=DCF!C8/{R('revenues_last_year')}-1" if self.has("revenues_last_year") else "n/a"
        blocks = [
            ("1. Revenue growth", [("Industry (5-yr avg)", ind("ind.growth"), PCT), ("Last year", last, PCT), ("Next year", "=DCF!D7", PCT),
                                   ("CAGR yrs 2-5", "=(DCF!H8/DCF!D8)^(1/4)-1", PCT)],
             "Growth above industry: is the company small? Very different from last year's growth: why?"),
            ("2. Dollar revenues", [("Now", "=DCF!C8", M), ("Year 1", "=DCF!D8", M), ("Year 5", "=DCF!H8", M), ("Year 10", "=DCF!M8", M)],
             "How big is the total market, what do its biggest firms earn, what share does year 10 imply?"),
            ("3. Operating margin", [("Industry", ind("ind.margin"), PCT), ("Now", "=DCF!C9", PCT), ("Year 5", "=DCF!H9", PCT), ("Year 10", "=DCF!M9", PCT)],
             "Industry margins? Unit economics (the cost of the extra unit)? What does competition look like?"),
            ("4. Reinvestment & ROIC", [("Industry sales/capital", ind("ind.s2c"), NUM2), ("Sales/capital now", "=DCF!C8/DCF!C22", NUM2),
                                        ("ROIC year 10", "=DCF!M23", PCT), ("Stable ROIC", "=DCF!N23", PCT)],
             "Bounce-back or new growth? Excess capacity? Does efficiency change with scale? Comfortable with year-10 ROIC?"),
            ("5. Cost of capital", [("Industry", ind("ind.coc"), PCT), ("Years 1-5", "=DCF!D18", PCT),
                                    ("Compounded", "=(1/DCF!M19)^(1/10)-1", PCT), ("Stable", "=DCF!N18", PCT)],
             "How does it compare with the industry? Why does it change over time? Is the failure rate consistent with the company?"),
            ("6. Value vs price", [("Value / price", "=DCF!C49/DCF!C50", PCT),
                                   ("Flag", '=IF(DCF!C49/DCF!C50<0.5,"Value low: revisit",IF(DCF!C49/DCF!C50>2,"Value high: revisit","Within 0.5x-2x"))', None)],
             "Ginzu flags: value below 50% or above 200% of price -> revisit growth, target margin, sales/capital and terminal ROIC."),
        ]
        r = 5
        for head, cells, ask in blocks:
            section(ws, r, head, last="H")
            hdr(ws, r + 1, {"CDEFG"[j]: lab for j, (lab, _, _) in enumerate(cells)})
            for j, (_, f, fmt) in enumerate(cells):
                put(ws, f"{'CDEFG'[j]}{r + 2}", f, fmt, kind="label" if f == "n/a" else None)
            put(ws, f"H{r + 1}", ask, kind="label").alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(f"H{r + 1}:H{r + 2}")
            ws.row_dimensions[r + 1].height = 30
            r += 4
        printable(ws)

    # ------------------------------------------------------------ summary
    def summary_sheet(self):
        ws, i = self.ws["Summary"], self.inp
        name = self.meta.get("company", "Company") + (f" ({self.meta['ticker']})" if self.meta.get("ticker") else "")
        title(ws, f"Valuation Summary: {name}", "Intrinsic value, Damodaran FCFF (Ginzu) framework"
              + (f" | Valuation date {self.meta['valuation_date']}" if self.meta.get("valuation_date") else "") + f" | {self.units}")
        widths(ws, {"B": 50, "C": 14, "D": 12, "E": 3, "F": 40, "G": 14, "H": 3, "T": 36, "U": 11, "V": 11, "W": 11})
        section(ws, 5, "Key outputs", last="D")
        kp = [("Intrinsic value per share", "=ValuePerShare", self.PS), ("Current share price", "=SharePrice", self.PS),
              ("Upside / (downside) to value", "=DCF!C52", PCT), ("Price as % of value", "=DCF!C51", PCT),
              ("Market capitalization", f"=SharePrice*{self.ref('shares_outstanding')}", M),
              ("Intrinsic value of operating assets", "=OperatingAssets", M), ("Intrinsic value of equity", "=DCF!C46", M),
              ("Initial / terminal cost of capital", '=TEXT(DCF!C27,"0.00%")&" / "&TEXT(DCF!C30,"0.00%")', None),
              ("Terminal growth", "=DCF!C31", PCT2), ("Year-10 revenue", "=DCF!M8", M)]
        if "target_margin" in i and "margin_path" not in i:
            be = breakeven(i, "target_margin", -0.5, 0.95)
            if be is not None:
                kp.append(("Breakeven target margin for value = price (engine)", be, PCT))
        for r, (lab, f, fmt) in enumerate(kp, start=6):
            label(ws, f"B{r}", lab, bold=r == 6)
            put(ws, f"C{r}", f, fmt, bold=r == 6)
        ws["C6"].fill, ws["C6"].border = FILL_LIGHT, BOX
        ws["C13"].alignment = Alignment(horizontal="right")
        section(ws, 5, "Enterprise value to equity bridge", first="F", last="G")
        bridge = [("Value of operating assets", "=DCF!C39", M), ("Less: debt", "=-DCF!C40", M), ("Less: minority interests", "=-DCF!C41", M),
                  ("Plus: cash", "=DCF!C42", M), ("Less: tax on trapped cash", "=-DCF!C43", M), ("Plus: non-operating assets", "=DCF!C44", M),
                  ("Plus: IPO proceeds retained", "=DCF!C45", M), ("Value of equity", "=DCF!C46", M),
                  ("Less: employee options", "=-DCF!C47", M), ("Shares outstanding", "=DCF!C48", M), ("Value per share", "=DCF!C49", self.PS)]
        for r, (lab, f, fmt) in enumerate(bridge, start=6):
            label(ws, f"F{r}", lab, bold=r in (13, 16))
            put(ws, f"G{r}", f, fmt, bold=r in (13, 16))
        ws["G13"].border = TOP
        ff = [("DCF: cost of capital sensitivity", f"=MIN(Sensitivity!{self.T1})", f"=MAX(Sensitivity!{self.T1})"),
              ("DCF: growth x margin sensitivity", f"=MIN(Sensitivity!{self.T2})", f"=MAX(Sensitivity!{self.T2})")]
        if self.scen:
            a, b, vc = self.scen_rows
            section(ws, 19, "Scenario analysis (live; set the active scenario on the Scenarios tab)", last="D")
            hdr(ws, 20, {"B": "Scenario", "C": "Value / share", "D": "vs price"})
            for k in range(b - a + 1):
                put(ws, f"B{21 + k}", f"=Scenarios!C{a + k}")
                put(ws, f"C{21 + k}", f"=Scenarios!{vc}{a + k}", self.PS)
                put(ws, f"D{21 + k}", f"=C{21 + k}/SharePrice-1", PCT)
            ff.append(("DCF: scenario range", f"=MIN(Scenarios!{vc}{a}:{vc}{b})", f"=MAX(Scenarios!{vc}{a}:{vc}{b})"))
        mk = {k: self.ref(f"m:{k}") for k in ("week52_low", "week52_high", "analyst_target") if self.has(f"m:{k}")}
        if "week52_low" in mk and "week52_high" in mk:
            ff.insert(0, ("52-week trading range", f"={mk['week52_low']}", f"={mk['week52_high']}"))
        if "analyst_target" in mk:
            ff.insert(1, ("Analyst average target (third party)", f"={mk['analyst_target']}*0.995", f"={mk['analyst_target']}*1.005"))
        ff += [("DCF base case", "=ValuePerShare*0.995", "=ValuePerShare*1.005"), ("Current share price", "=SharePrice*0.995", "=SharePrice*1.005")]
        section(ws, 5, "Chart data: football field", first="T", last="W")
        hdr(ws, 6, {"T": "Reference", "U": "Low", "V": "High - low", "W": "High"})
        for r, (lab, lo, hi) in enumerate(ff, start=7):
            label(ws, f"T{r}", lab)
            put(ws, f"U{r}", lo, self.PS)
            put(ws, f"W{r}", hi, self.PS)
            put(ws, f"V{r}", f"=W{r}-U{r}", self.PS)
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap = "bar", "stacked", 100
        ch.title = "Football field: value per share"
        ch.add_data(Reference(ws, min_col=21, max_col=22, min_row=6, max_row=6 + len(ff)), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=20, min_row=7, max_row=6 + len(ff)))
        ch.series[0].graphicalProperties.noFill = True
        ch.series[0].graphicalProperties.line.noFill = True
        ch.series[1].graphicalProperties.solidFill = "13315C"
        ch.legend, ch.y_axis.majorGridlines = None, None
        ch.y_axis.number_format = f'"{self.cur}"#,##0'
        ch.x_axis.delete = ch.y_axis.delete = False
        ch.height, ch.width = 8.2, 17.5
        ws.add_chart(ch, "I5")
        printable(ws, one_page=True)
        ws.print_area = "A1:S40"

    # ------------------------------------------------------------ checks, sources, cover
    def checks_sheet(self):
        ws, eng = self.ws["Checks"], self.out
        title(ws, "Model Checks", "Ties the workbook to the aswath-finance engine (ginzu.py, calibrated to Damodaran's spreadsheets). "
              "Engine values are hard-coded at export; they are only comparable in the base state.")
        widths(ws, {"B": 6, "C": 62, "D": 16, "E": 16, "F": 12, "G": 60})
        base = "AND(" + ",".join(f"{v}=0" for v in FX.values()) + (",Scenarios!$D$5=1" if self.scen else "") + ")"
        wacc0 = eng["years"][0]["cost_of_capital"]
        rows = [("Model in base state (flex cells neutral" + (", scenario 1)" if self.scen else ")"),
                 f'=IF({base},"Base","Flexed")', '="Base"', "=D{r}=E{r}", "Engine tie-outs apply in the base state only"),
                ("Value per share ties to engine", "=ValuePerShare", eng["value_per_share"], f"=IF({base},ABS(D{{r}}-E{{r}})<0.005,TRUE)",
                 "ginzu.py value_per_share"),
                ("Initial cost of capital ties to engine", "=WACC_Initial", wacc0, f"=IF({base},ABS(D{{r}}-E{{r}})<0.000001,TRUE)",
                 "cost_of_capital.py / input")]
        if self.scen:
            a, b, vc = self.scen_rows
            diff = ",".join(f"ABS(Scenarios!{vc}{a + 1 + k}-{v!r})" for k, v in enumerate(self.eng_scen))
            rows.append(("Scenario values tie to engine (max abs difference)", f"=MAX({diff})", 0.0,
                         f"=IF({base},D{{r}}<0.01,TRUE)", "ginzu.py on each scenario's inputs"))
        rows += [("Sensitivity table 1 centre = base value", f"={self.C1}", "=ValuePerShare", f"=IF({base},ABS(D{{r}}-E{{r}})<0.005,TRUE)",
                  "Excel data table (recalculates on open)"),
                 ("Sensitivity table 2 centre = base value", f"={self.C2}", "=ValuePerShare", f"=IF({base},ABS(D{{r}}-E{{r}})<0.005,TRUE)", ""),
                 ("Equity bridge adds up", "=DCF!C46", "=DCF!C39-DCF!C40-DCF!C41+DCF!C42-DCF!C43+DCF!C44+DCF!C45",
                  "=ABS(D{r}-E{r})<0.001", "")]
        if self.dil:
            L = 18 + ITER
            rows.append(("Dilution iteration converged", f"=Dilution!J{L}", f"=Dilution!J{L - 1}",
                         "=ABS(D{r}-E{r})<0.000001*MAX(1,ABS(D{r}))", "Last two fixed-point steps agree"))
        hdr(ws, 5, {"B": "#", "C": "Check", "D": "Workbook", "E": "Engine / rule", "F": "Result", "G": "Note"})
        r = 6
        for k, (lab, wv, ev, res, note) in enumerate(rows, start=1):
            fmt = PCT2 if "cost of capital" in lab else (self.PS if "value" in lab.lower() or "table" in lab else M)
            put(ws, f"B{r}", k, "0", kind="label")
            label(ws, f"C{r}", lab)
            put(ws, f"D{r}", wv, fmt)
            put(ws, f"E{r}", ev, fmt)
            put(ws, f"F{r}", res.format(r=r))
            put(ws, f"G{r}", note, kind="label").font = NOTE
            r += 1
        last = r - 1
        for op, fill in (("TRUE", GREEN_FILL), ("FALSE", RED_FILL)):
            ws.conditional_formatting.add(f"F6:F{last}", CellIsRule(operator="equal", formula=[op], fill=fill))
        label(ws, f"C{r + 1}", "ALL CHECKS PASS", bold=True)
        put(ws, f"F{r + 1}", f"=COUNTIF(F6:F{last},FALSE)=0", bold=True).border = BOX
        self.checks_cell = f"Checks!F{r + 1}"
        self.wb.defined_names["AllChecksPass"] = DefinedName("AllChecksPass", attr_text=f"Checks!$F${r + 1}")
        r += 3
        section(ws, r, "Ginzu rules (warnings, not errors: Damodaran states them outright)", last="G")
        for k, (lab, wv, ev, res) in enumerate((
                ("Terminal growth <= riskfree rate after year 10", "=DCF!C31", "=DCF!C28", "=D{r}<=E{r}+0.0000001"),
                ("Terminal ROIC >= terminal cost of capital (else growth destroys value)", "=DCF!C32", "=DCF!C30", "=D{r}>=E{r}-0.0000001"),
                ("Value / price within the Ginzu 0.5x-2x Diagnostics flag", "=DCF!C49/DCF!C50", "0.5x-2x", "=AND(D{r}>=0.5,D{r}<=2)")), start=r + 1):
            label(ws, f"C{k}", lab)
            put(ws, f"D{k}", wv, PCT2)
            put(ws, f"E{k}", ev, None if ev == "0.5x-2x" else PCT2)
            put(ws, f"F{k}", f'=IF({res.format(r=k)[1:]},"OK","WARNING")')
        printable(ws)

    def sources_sheet(self):
        ws = self.ws["Sources"]
        title(ws, "Sources, Methodology & Notes", "Every input traces to a filing, market data, or Damodaran's own model/data")
        widths(ws, {"B": 4, "C": 34, "D": 110})
        s = self.sources
        rows = list(s.items()) if isinstance(s, dict) else [("", x) for x in s]
        rows += [("Notes" if k == 0 else "", n) for k, n in enumerate(self.notes)]
        rows += [("Methodology", "FCFF DCF transcribed from Damodaran's fcffsimpleginzu spreadsheet (ginzu.py, which reproduces "
                                 "31 published Damodaran valuations); cost of capital from Damodaran's worksheet (cost_of_capital.py)."),
                 ("", f"Data: {self.data['source']}, {self.data['updated']}."),
                 ("Disclaimer", "Educational application of Damodaran's public framework and data. Not investment advice, not a "
                                "recommendation, and not affiliated with or endorsed by Damodaran or any bank.")]
        section(ws, 5, "Sources and notes", first="B", last="D")
        for r, (k, v) in enumerate(rows, start=6):
            label(ws, f"C{r}", k, bold=bool(k)).alignment = Alignment(vertical="top")
            c = put(ws, f"D{r}", str(v), kind="label")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            m = re.search(r"https?://\S+", str(v))
            if m:
                c.hyperlink, c.font = m.group(0).rstrip(").,"), LINK
        printable(ws)

    def cover_sheet(self):
        ws = self.ws["Cover"]
        ws.sheet_view.showGridLines = False
        widths(ws, {"A": 3, "B": 4, "C": 38, "D": 70})
        for r in range(2, 9):
            for c in "BCD":
                ws[f"{c}{r}"].fill = FILL_NAVY
        ws["C3"] = self.meta.get("company", "Company") + (f" ({self.meta['ticker']})" if self.meta.get("ticker") else "")
        ws["C3"].font = font(20, bold=True, color="FFFFFF")
        ws["C5"], ws["C5"].font = "Intrinsic Valuation: Damodaran FCFF (Ginzu) Framework", font(13, color="FFFFFF")
        ws["C7"] = " | ".join(x for x in (f"Valuation date: {self.meta['valuation_date']}" if self.meta.get("valuation_date") else "",
                                          self.units, "For internal discussion") if x)
        ws["C7"].font = font(italic=True, color=LIGHT)
        label(ws, "C10", "Headline", bold=True)
        for r, (lab, f, fmt) in enumerate((("Intrinsic value per share", "=ValuePerShare", self.PS), ("Share price", "=SharePrice", self.PS),
                                           ("Price as % of value", "=DCF!C51", PCT),
                                           ("Model checks", '=IF(AllChecksPass,"All checks pass","CHECK FAILURES - see Checks tab")', None)), start=11):
            label(ws, f"C{r}", lab)
            put(ws, f"D{r}", f, fmt, bold=True).alignment = Alignment(horizontal="left")
        label(ws, "C16", "Contents", bold=True)
        desc = {"Summary": "Key outputs, EV-to-equity bridge, scenarios, football field", "Inputs": "Every hard-coded input with its source (blue)",
                "Scenarios": "Scenario selector: the whole model runs on the active scenario", "Sensitivity": "Live data tables",
                "WACC": "Cost of capital (Damodaran's worksheet)", "DCF": "FCFF projection, terminal value, equity bridge",
                "Dilution": "Employee options and IPO shares at value (explicit iteration)", "Diagnostics": "Ginzu diagnostic questions vs industry",
                "Checks": "Tie-out to the engine and the Ginzu rules", "Sources": "Sources, methodology, disclaimer"}
        r = 17
        for sh in self.ws:
            if sh == "Cover":
                continue
            c = ws[f"C{r}"]
            c.value, c.hyperlink, c.font = sh, f"#'{sh}'!A1", LINK
            label(ws, f"D{r}", desc[sh])
            r += 1
        label(ws, f"C{r + 1}", "Colour convention", bold=True)
        put(ws, f"C{r + 2}", "Blue: hard-coded input", kind="input")
        put(ws, f"C{r + 3}", "Black: formula", kind="calc")
        put(ws, f"C{r + 4}", "Green: link to another sheet", kind="link")
        label(ws, f"C{r + 5}", "Yellow fill: flex / selector cell (keep at base)").fill = FILL_FLEX
        put(ws, f"C{r + 7}", "Educational application of Damodaran's public framework and data; not investment advice.", kind="label").font = NOTE
        printable(ws, landscape=False)


# ---------------------------------------------------------------- verify (Excel) and selftest (pycel)
def excel_recalc(path: Path) -> str | None:
    """Recalculate + save in Excel (Windows COM) so the file carries cached values. Returns an error or None."""
    if sys.platform != "win32":
        return "needs Excel on Windows; open the file in Excel instead (it recalculates on open)"
    p = str(path.resolve()).replace("'", "''")
    ps = ("$ErrorActionPreference='Stop'; $xl=New-Object -ComObject Excel.Application; $xl.Visible=$false; $xl.DisplayAlerts=$false; "
          f"try {{ $wb=$xl.Workbooks.Open('{p}'); $xl.CalculateFull(); $wb.Save(); $wb.Close($false) }} finally {{ $xl.Quit(); "
          "[Runtime.Interopservices.Marshal]::ReleaseComObject($xl) | Out-Null }")
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
    return r.stderr.strip()[-300:] or "Excel failed" if r.returncode else None


def read_cells(path: Path, cells: dict) -> dict:
    wb = load_workbook(path, data_only=True)
    return {k: wb[a.split("!")[0]][a.split("!")[1]].value for k, a in cells.items()}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("inputs", nargs="?")
    ap.add_argument("-o", "--out")
    ap.add_argument("--verify", action="store_true", help="recalculate in Excel (Windows) and report the Checks tab")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        from xlsx_selftest import selftest
        return selftest()
    if not a.inputs:
        ap.error("inputs JSON required")
    src = Path(a.inputs)
    out = Path(a.out) if a.out else src.with_name(src.stem + "_valuation.xlsx")
    b = Book(json.loads(src.read_text(encoding="utf-8")))
    cells = b.build(out)
    print(f"wrote {out}\nengine value per share {b.out['value_per_share']:,.4f} (price {b.inp['price']:,.2f})")
    if a.verify:
        err = excel_recalc(out)
        if err:
            print(f"verify skipped: {err}")
            return 0
        got = read_cells(out, {"value": cells["value"], "checks": cells["checks"]})
        print(f"Excel: value per share {got['value']:,.4f} | all checks pass: {got['checks']}")
        return 0 if got["checks"] is True else 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
