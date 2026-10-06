"""Damodaran's market-level tools: intrinsic value of the S&P 500, implied ERP, breakeven revenues for a
market cap / enterprise value, PE history of the index, and the market-cap share of the market leaders.

Damodaran's spreadsheets (https://pages.stern.nyu.edu/~adamodar/pc/blog/<file>), transcribed cell by cell:
  index "S&P 500 Valuation" sheet (B3:H38), one template per version, keyed by `template`:
        2023-01  S&P500ValueJan2023.xlsx   2 forecast years, payout % input, full-year discounting
        2023-07  S&P500ValueJuly2023.xlsx  3 forecast years, (1+r)^0.5 first period
        2023-10  S&P500ValueOct2023.xlsx   3 forecast years, (1+r)^0.25 first period
        2024-01  S&P500ValueJan2024.xlsx, S&PValueJan2025.xlsx  2 years, growth fades to riskfree by year 5
        2025-09  IntrinsicValueS&P500Sept25.xlsx "Impl premium calculator" B3:H54 (same sheet as the
                 monthly ERPOct25 / ERPJan26 files, pc/implprem/): 5-year growth, payout fading to 1-g/ROE,
                 flat rf+ERP, partial first year; plus 'Expected growth rate'!G11 (top-down growth)
  breakeven  MktCapBreakeven.xlsx (equity: net margin, ROE, cost of equity), BreakevenAIBusiness.xlsx,
             BreakevenAnthropic.xlsx, BreakevenEnterprisevalue.xlsx (firm: after-tax operating margin, ROIC,
             cost of capital, optional finite life): 'Breakeven Calculation', 'Revenue Sensitivity Table',
             'Waiting Table'
  pe         PEHistory2026.xlsx / PEHistorical2025.xlsx Sheet1 (PE, normalized PE, CAPE, T.Bond PE, ...)
  leaders    FearsomeFiveMktCap.xlsx (Apple, Microsoft, Alphabet, Amazon, Meta vs all US equities)

Usage:
  python market.py inputs.json [--json]     # {"model": "index" | "breakeven", ...}; index also solves implied ERP
  python market.py --pe [--json]            # Damodaran's PE table 1960-2025 + period averages
  python market.py --leaders [--json]       # Fearsome Five share of US market cap, 2009-2024
  python market.py --from-case BreakevenAnthropic   # the source inputs for a case, as a starting JSON
  python market.py --selftest               # replay data/market_cases.json
  python market.py --build-cases <xlsx...>  # rebuild cases/data from the source files (needs openpyxl, pycel)

Index inputs (index points, rates as decimals): index_level, riskfree_rate, erp (number or a spreadsheet code ERP10,
  ERPLong, ERPHigh, ERPLow, HUS, HG), roe, cape_denominator (optional, the spreadsheet's 10-yr earnings average)
  2023-01..2024-01: base_earnings, earnings_forecasts [E1, E2(, E3)], base_payout (2023-01) or
    base_cash_return (dividends+buybacks), earnings_adjustment (omit = "No"), riskfree_change (bool),
    riskfree_after_5, riskfree_change_mode ("gradual" | "immediate" | "end")
  2025-09: base_earnings, trailing_earnings, base_cash_flow, growth (or consensus [E0, E1, E2] -> the spreadsheet's
    top-down growth), adjust_payout (default true), long_term_growth (default riskfree), month (start of),
    default_spread (optional)
Breakeven inputs: basis "equity" | "firm", value_to_match, margin, return_on_capital, risk_class (D1..D5) or
  discount_rate, stable_growth, years_to_maturity, cash_yield, current_revenues, finite_years (optional),
  replicate_sheet_bugs (true = the spreadsheet's finite-life formula as written; default = corrected growing annuity).
"""

from __future__ import annotations

import argparse
import json
import math
import sys
from functools import lru_cache
from pathlib import Path

HERE = Path(__file__).resolve().parent
DATA = HERE.parent / "data" / "market_data.json"
CASES = HERE.parent / "data" / "market_cases.json"
FIRST_PERIOD = {"2023-01": 1.0, "2023-07": 0.5, "2023-10": 0.25, "2024-01": 1.0}  # exponent on (1+r) in year 1


@lru_cache(maxsize=1)
def load() -> dict:
    return json.loads(DATA.read_text(encoding="utf-8"))


# ---------------------------------------------------------------- index value: "S&P 500 Valuation" sheet
def _earnings_2023(t: str, e0: float, f: list, adj: float | None, g: float) -> list:
    """Row 'Expected Earnings', years 1-5, as each version of the sheet builds it."""
    if t == "2023-01":
        if adj is not None:  # C24, D24 adjusted; F24 = avg(yr1, yr2) grown 3 years; E24 halfway
            c, d = f[0] * (1 + adj), f[1] * (1 + adj / 2)
            y4 = (c + d) / 2 * (1 + g) ** 3
            return [c, d, d + (y4 - d) / 2, y4, y4 * (1 + g)]
        return [f[0], f[1]] + [f[1] * (1 + g) ** k for k in (1, 2, 3)]
    if t == "2023-07":
        c, d = (f[0] * (1 + adj), f[1] * (1 + adj / 2)) if adj is not None else (f[0], f[1])
        e = f[2]
        if adj is not None:
            y4 = (c + d) / 2 * (1 + g) ** 3
        else:
            ge = e / d - 1
            y4 = e * (1 + ge - (ge - g) / 2)  # F26 = E26 - (E26 - H26)/2
        return [c, d, e, y4, y4 * (1 + g)]
    if t == "2023-10":
        d, e = (f[1] * (1 + adj), f[2] * (1 + adj / 2)) if adj is not None else (f[1], f[2])
        return [f[0], d, e, e * (1 + g), e * (1 + g) ** 2]
    if t == "2024-01":  # growth fades from year-2 growth to g: 2/3, 1/3, 0 of the gap
        d = f[1] * (1 + adj) if adj is not None else f[1]
        gd = d / f[0] - 1
        e = d * (1 + 2 * (gd - g) / 3 + g)
        y4 = e * (1 + (gd - g) / 3 + g)
        return [f[0], d, e, y4, y4 * (1 + g)]
    raise ValueError(f"unknown template {t}")


def _erp(inp: dict) -> float:
    e = inp["erp"]
    return load()["erp_choices"]["choices"][e]["value"] if isinstance(e, str) else e


def value_2023(inp: dict) -> dict:
    t, rf, erp, roe = inp["template"], inp["riskfree_rate"], _erp(inp), inp["roe"]
    change = inp.get("riskfree_change", False)
    rf_after = inp.get("riskfree_after_5", rf)
    mode = inp.get("riskfree_change_mode", "gradual")
    g = rf_after if change else rf  # B12: stable growth = riskfree rate (after year 5)
    rfs = [rf if not change else rf_after if mode == "immediate" else
           rf + (rf_after - rf) * k / 5 if mode == "gradual" else rf for k in range(1, 6)]
    e0 = inp["base_earnings"]
    earn = _earnings_2023(t, e0, inp["earnings_forecasts"], inp.get("earnings_adjustment"), g)
    earn.append(earn[4] * (1 + g))  # terminal year
    p0 = inp["base_payout"] if t == "2023-01" else inp["base_cash_return"] / e0
    pt = 1 - g / roe  # sustainable payout
    pay = [p0 - (p0 - pt) * k / 4 for k in range(5)] + [pt]
    cf = [e * p for e, p in zip(earn, pay)]
    r = [x + erp for x in rfs]
    tv = cf[5] / erp  # H28/(H31-H30): terminal return (riskfree + ERP) minus growth (= riskfree)
    disc, pv = (1 + r[0]) ** FIRST_PERIOD[t], []
    for k in range(5):
        disc = disc if k == 0 else disc * (1 + r[k])
        pv.append(cf[k] / disc)
    val = sum(pv) + tv / disc
    years = [{"t": k + 1, "earnings": earn[k], "payout": pay[k], "cash_flow": cf[k], "riskfree": rfs[k],
              "required_return": r[k], "pv": pv[k]} for k in range(5)]
    years.append({"t": "terminal", "earnings": earn[5], "payout": pt, "cash_flow": cf[5], "riskfree": g,
                  "required_return": g + erp, "pv": tv / disc})
    return _outputs(inp, val, erp, tv, years, forward=earn[0], stable_growth=g)


# ---------------------------------------------------------------- index value: 2025 estimator / ERP calculator
def topdown_growth(e0: float, e1: float, e2: float, rf: float) -> float:
    """'Expected growth rate'!G11: consensus growth for 2 years, fading over 2 years to the riskfree rate."""
    g1, g2 = e1 / e0 - 1, e2 / e1 - 1
    g3 = g2 - (g2 - rf) / 3
    path = [g1, g2, g3, g2 - 2 * (g2 - rf) / 3, g3 - 2 * (g3 - rf) / 2]
    return math.prod(1 + x for x in path) ** 0.2 - 1


def value_2025(inp: dict) -> dict:
    rf, erp, e0 = inp["riskfree_rate"], _erp(inp), inp["base_earnings"]
    g = inp["growth"] if inp.get("growth") is not None else topdown_growth(*inp["consensus"], rf)
    glt = inp.get("long_term_growth", rf)  # B14 "No" -> riskfree rate
    earn = [e0 * (1 + g) ** k for k in range(1, 6)]
    earn.append(earn[4] * (1 + glt))
    p0 = inp["base_cash_flow"] / inp["trailing_earnings"]  # B34 = B35 / trailing reported earnings
    adjust = inp.get("adjust_payout", True)
    pt = 1 - glt / inp["roe"] if adjust else p0
    pay = [p0 - (p0 - pt) * k / 5 if adjust else p0 for k in range(1, 6)] + [pt]
    cf = [e * p for e, p in zip(earn, pay)]
    r = rf + erp
    tv = cf[5] / (r - glt)
    f = load()["months_to_year_end"]["table"][inp.get("month", "October")] / 12  # Answers!G2:H13
    times = [f + k for k in range(5)]
    pv = [cf[0] * f / (1 + r) ** f] + [cf[k] / (1 + r) ** times[k] for k in range(1, 5)]
    val = sum(pv) + tv / (1 + r) ** times[4]
    years = [{"t": round(times[k], 4), "earnings": earn[k], "payout": pay[k], "cash_flow": cf[k],
              "riskfree": rf, "required_return": r, "pv": pv[k]} for k in range(5)]
    years.append({"t": "terminal", "earnings": earn[5], "payout": pt, "cash_flow": cf[5], "riskfree": rf,
                  "required_return": r, "pv": tv / (1 + r) ** times[4]})
    out = _outputs(inp, val, erp, tv, years, forward=None, stable_growth=glt)
    out["growth"] = g
    out["trailing_pe"] = val / inp["trailing_earnings"]  # C40 uses reported trailing earnings
    return out


def _outputs(inp: dict, val: float, erp: float, tv: float, years: list, forward: float | None,
             stable_growth: float) -> dict:
    lvl = inp["index_level"]
    out = {"value": val, "index_level": lvl, "over_under": lvl / val - 1, "erp": erp,
           "expected_return": inp["riskfree_rate"] + erp, "terminal_value": tv, "stable_growth": stable_growth,
           "trailing_pe": val / inp["base_earnings"], "years": years}
    if forward:
        out["forward_pe"] = val / forward
    if inp.get("cape_denominator"):
        out["cape"] = val / inp["cape_denominator"]
    return out


def value_index(inp: dict) -> dict:
    return value_2025(inp) if inp.get("template", "2025-09") == "2025-09" else value_2023(inp)


def implied_erp(inp: dict) -> dict:
    """The spreadsheet's goal seek: the ERP that makes the intrinsic value equal the index level."""
    lo = max(1e-6, inp.get("long_term_growth", inp["riskfree_rate"]) - inp["riskfree_rate"] + 1e-6)
    hi = 0.5
    f = lambda e: value_index({**inp, "erp": e})["value"] - inp["index_level"]
    if f(lo) * f(hi) > 0:
        return {"implied_erp": None}
    for _ in range(200):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if f(mid) > 0 else (lo, mid)
        if hi - lo < 1e-13:
            break
    e = (lo + hi) / 2
    out = {"implied_erp": e, "implied_expected_return": inp["riskfree_rate"] + e}
    if "default_spread" in inp:  # C44: ERP measured against the default-free (Aa1-adjusted) riskfree rate
        out["implied_erp_default_adj"] = e + inp["default_spread"]
    return out


# ---------------------------------------------------------------- breakeven revenues
def breakeven(inp: dict) -> dict:
    basis = inp.get("basis", "equity")
    d = load()["breakeven_defaults"][basis]
    g = inp.get("stable_growth", d["stable_growth"])
    r = inp.get("discount_rate") or next(v for k, v in d["risk_classes"].items()
                                         if k.startswith(inp.get("risk_class", d["risk_class"])[:2]))
    v, m, roc = inp["value_to_match"], inp["margin"], inp["return_on_capital"]
    n, y = inp.get("years_to_maturity", d["years_to_maturity"]), inp.get("cash_yield", d["cash_yield"])
    per_rev = lambda mm, rr: mm * (1 + g) * (1 - g / rr) / (r - g)  # value of $1 of mature revenue
    k = his_k = per_rev(m, roc)
    if inp.get("finite_years"):  # the 'Finite' branch (see tool-market.md): the source sheet multiplies by the deferral
        f = ((1 + g) / (1 + r)) ** inp["finite_years"]  # factor; default is the growing-annuity factor Damodaran's method implies
        his_k = k * f
        k = his_k if inp.get("replicate_sheet_bugs") else k * (1 - f)
    now = v / k
    wait = now * (1 + r - y) ** n
    ax = {a: inp.get(a, d[a]) for a in ("grid_margins", "grid_returns", "wait_years", "wait_yields")}
    ms, rs = list(ax["grid_margins"]), list(ax["grid_returns"])
    if d["append_own"]:  # the firm sheets add your own margin (row 14) and ROIC (column K) to the grid
        ms, rs = ms + [m], rs + [roc]
    out = {"basis": basis, "discount_rate": r, "stable_growth": g, "breakeven_revenue_now": now,
           "breakeven_revenue_wait": wait, "years_to_maturity": n, "cash_yield": y, "margins": ms, "returns": rs,
           "grid": [[v / per_rev(mm, rr) for rr in rs] for mm in ms],  # perpetuity formula, as in the source sheet
           "wait_years": ax["wait_years"], "wait_yields": ax["wait_yields"],
           "wait_grid": [[now * (1 + r - yy) ** nn for yy in ax["wait_yields"]] for nn in ax["wait_years"]]}
    if inp.get("current_revenues"):
        out["cagr_needed"] = (wait / inp["current_revenues"]) ** (1 / n) - 1
    if inp.get("finite_years"):
        out |= {"finite_years": inp["finite_years"], "replicate_sheet_bugs": bool(inp.get("replicate_sheet_bugs")),
                "his_sheet_breakeven_revenue_now": v / his_k}
    return out


# ---------------------------------------------------------------- PE history and market leaders
def pe_history(rows: list) -> list:
    """Sheet1 columns B-I: PE, normalized PE (10-yr avg earnings), CAPE (10-yr avg inflation-adjusted
    earnings), Shiller PE, T.Bond PE, Shiller/T.Bond PE, price/dividends, price/cash returned."""
    out, infl = [], []
    for i, r in enumerate(rows):
        infl.append(1.0 if i == 0 else infl[-1] * (1 + r["inflation"]))
        ey = r.get("earnings_yield") or r["earnings"] / r["index"]
        o = {"year": r["year"], "pe": 1 / ey, "earnings_yield": ey, "tbond": r["tbond"], "tbond_pe": 1 / r["tbond"],
             "ey_minus_tbond": ey - r["tbond"], "price_dividends": r["index"] / r["dividends"]}
        if i >= 9:
            last = rows[i - 9:i + 1]
            o["norm_pe"] = r["index"] / (sum(x["earnings"] for x in last) / 10)
            o["cape"] = r["index"] / (sum(x["earnings"] * infl[i] / infl[i - 9 + j] for j, x in enumerate(last)) / 10)
        o["shiller_pe"] = r.get("shiller_pe") or o.get("cape")
        o["shiller_to_tbond_pe"] = o["shiller_pe"] / o["tbond_pe"]
        if r.get("cash_returned"):
            o["price_cash_returned"] = r["index"] / r["cash_returned"]
        out.append(o)
    return out


def period_average(table: list, key: str, y0, y1) -> float:
    xs = [r[key] for r in table if y0 <= _yr(r["year"]) <= y1 and r.get(key) is not None]  # AVERAGE skips blanks
    return sum(xs) / len(xs)


def _yr(y) -> int:
    return y if isinstance(y, int) else int(str(y)[:4])


def leaders_share(rows: list) -> list:
    """Fearsome Five market cap, share of US equities, and share of the year's change in market cap."""
    out = []
    for i, r in enumerate(rows):
        five = sum(r[c] for c in ("apple", "microsoft", "alphabet", "amazon", "meta"))
        o = {"year": r["year"], "five": five, "all_us": r["all_us"], "share": five / r["all_us"]}
        if i:
            o["change_five"], o["change_all"] = five - out[-1]["five"], r["all_us"] - out[-1]["all_us"]
            o["share_of_change"] = o["change_five"] / o["change_all"]
        out.append(o)
    return out


# ---------------------------------------------------------------- reports
def report_index(inp: dict, out: dict) -> str:
    rows = [f"Template {inp.get('template', '2025-09')} | index {out['index_level']:,.2f} | riskfree "
            f"{inp['riskfree_rate']:.2%} | ERP {out['erp']:.2%} | expected return {out['expected_return']:.2%}"
            + (f" | 5-yr growth {out['growth']:.2%}" if "growth" in out else ""),
            "year      earnings  payout  cash flow  required r        PV"]
    for y in out["years"]:
        rows.append(f"{str(y['t']):8s} {y['earnings']:9.2f} {y['payout']:7.2%} {y['cash_flow']:10.2f} "
                    f"{y['required_return']:10.2%} {y['pv']:10.2f}")
    rows += [f"Terminal value (end yr 5) {out['terminal_value']:,.2f}  stable growth {out['stable_growth']:.2%}",
             f"INTRINSIC VALUE OF INDEX  {out['value']:,.2f}   index {out['index_level']:,.2f}   "
             f"{'over' if out['over_under'] > 0 else 'under'}valued by {abs(out['over_under']):.1%}",
             f"Intrinsic trailing PE {out['trailing_pe']:.2f}" +
             (f" | forward PE {out['forward_pe']:.2f}" if "forward_pe" in out else "") +
             (f" | CAPE {out['cape']:.2f}" if "cape" in out else "")]
    ie = out.get("implied")
    if ie and ie["implied_erp"] is not None:
        rows.append(f"IMPLIED ERP (value = index) {ie['implied_erp']:.2%}, expected return "
                    f"{ie['implied_expected_return']:.2%}" + (f", vs default-adjusted riskfree "
                                                             f"{ie['implied_erp_default_adj']:.2%}"
                                                             if "implied_erp_default_adj" in ie else ""))
    ch = load()["erp_choices"]["choices"]
    rows.append("Damodaran's ERP yardsticks: " + ", ".join(f"{k} {v['value']:.2%}" for k, v in ch.items()))
    return "\n".join(rows)


def report_breakeven(out: dict) -> str:
    lab = ("market cap", "net margin", "ROE", "cost of equity") if out["basis"] == "equity" else \
        ("enterprise value", "after-tax operating margin", "ROIC", "cost of capital")
    rows = [f"Basis {out['basis']} ({lab[0]}) | {lab[3]} {out['discount_rate']:.2%} | stable growth "
            f"{out['stable_growth']:.2%}",
            f"BREAKEVEN REVENUE if mature today   {out['breakeven_revenue_now']:,.0f}",
            f"BREAKEVEN REVENUE after {out['years_to_maturity']} yrs wait "
            f"(cash yield {out['cash_yield']:.1%})  {out['breakeven_revenue_wait']:,.0f}"]
    if "cagr_needed" in out:
        rows.append(f"Revenue CAGR needed               {out['cagr_needed']:.2%}")
    if out.get("finite_years"):
        rows.append(f"NOTE finite life {out['finite_years']} yrs: SHEET ERROR - the source sheet multiplies by ((1+g)/(1+r))^N "
                    "(a perpetuity deferred N years): the sheet's breakeven revenue if mature today "
                    f"{out['his_sheet_breakeven_revenue_now']:,.0f}. This run uses "
                    + ("the source sheet's factor (replicate_sheet_bugs=true)" if out["replicate_sheet_bugs"] else
                       "the growing-annuity factor 1-((1+g)/(1+r))^N; set replicate_sheet_bugs=true to reproduce the source sheet"))
    rows.append(f"\nSensitivity (mature today): rows {lab[1]}, columns {lab[2]}")
    rows.append("        " + "".join(f"{x:>11.0%}" for x in out["returns"]))
    rows += [f"{m:6.0%}  " + "".join(f"{v:11,.0f}" for v in row) for m, row in zip(out["margins"], out["grid"])]
    rows.append("\nWaiting table: rows years to maturity, columns cash yield during wait")
    rows.append("        " + "".join(f"{x:>11.1%}" for x in out["wait_yields"]))
    rows += [f"{n:6d}  " + "".join(f"{v:11,.0f}" for v in row) for n, row in zip(out["wait_years"], out["wait_grid"])]
    return "\n".join(rows)


def report_pe(table: list, periods: list) -> str:
    keys = ["pe", "norm_pe", "cape", "tbond_pe", "shiller_to_tbond_pe", "earnings_yield", "tbond"]
    fmt = lambda k, v: "      -" if v is None else (f"{v:7.2%}" if k in ("earnings_yield", "tbond") else f"{v:7.2f}")
    rows = ["year        PE  normPE    CAPE  TBondPE  Shil/TB     E/P   TBond"]
    rows += [f"{str(r['year']):8s}" + " ".join(fmt(k, r.get(k)) for k in keys) for r in table]
    rows.append("\nperiod averages")
    for y0, y1 in periods:
        rows.append(f"{y0}-{y1}  " + " ".join(fmt(k, period_average(table, k, y0, y1)) for k in keys))
    return "\n".join(rows)


# ---------------------------------------------------------------- selftest
def run_case(c: dict) -> dict:
    k, inp = c["kind"], c["inputs"]
    if k == "index":
        return {**value_index(inp), **implied_erp(inp)} if c["expected"].get("implied_erp") else value_index(inp)
    if k == "topdown":
        return {"growth": topdown_growth(*inp["consensus"], inp["riskfree_rate"])}
    if k == "breakeven":
        return breakeven(inp)
    if k == "pe":
        t = pe_history(inp["rows"])
        out = {f"{r['year']}.{f}": r.get(f) for r in t for f in c["expected_fields"]}
        for lab, (y0, y1), f in inp["periods"]:
            out[f"{lab}.{f}"] = period_average(t, f, y0, y1)
        return out
    if k == "leaders":
        return {f"{r['year']}.{f}": r[f] for r in leaders_share(inp["rows"]) for f in r if f != "year"}
    raise ValueError(k)


def _pairs(his, ours):
    """Walk the sheet's expected value and ours in parallel (numbers or nested lists)."""
    if isinstance(his, list):
        for h, o in zip(his, ours, strict=True):
            yield from _pairs(h, o)
    else:
        yield his, ours


def selftest() -> int:
    bad = 0
    cases = json.loads(CASES.read_text(encoding="utf-8"))
    for c in cases:
        got, worst, n = run_case(c), 0.0, 0
        for key, his in c["expected"].items():
            for h, o in _pairs(his, got[key]):
                d = abs(o - h)
                worst = max(worst, 0.0 if d <= 1e-6 else d / abs(h))
                n += 1
        ok = worst <= 5e-4
        bad += not ok
        print(f"{'ok ' if ok else 'BAD'} {c['name'][:60]:60s} {n:4d} values  max rel diff {worst:.1e}")
    print(f"{len(cases) - bad}/{len(cases)} cases match Damodaran's spreadsheets (0.05% relative or 1e-6 absolute)")
    return 1 if bad else 0


# ---------------------------------------------------------------- build cases from the source files (openpyxl, pycel)
URL = "https://pages.stern.nyu.edu/~adamodar/pc/"
AVG = r"AVERAGE\('Buyback & Dividend computation'!([A-Z]+)(\d+):[A-Z]+(\d+)\)"
MODES = {"Gradual (over 5 years)": "gradual", "Immediate (starting now)": "immediate",
         "End (at the end of 5 years)": "end"}


def _src(p: Path) -> str:
    return URL + ("implprem/" if p.stem.startswith("ERP") else "blog/") + p.name


def _avg_range(ws_f, ws_bb, cell: str) -> float:
    """The spreadsheet's CAPE denominator: the AVERAGE(...) range named in the formula, over cached values."""
    import re  # noqa: PLC0415
    col, r0, r1 = re.search(AVG, ws_f[cell].value).groups()
    return sum(ws_bb[f"{col}{r}"].value for r in range(int(r0), int(r1) + 1)) / (int(r1) - int(r0) + 1)


def _b_sp500(p: Path, wf, wv) -> list:
    f, v, bb = wf["S&P 500 Valuation"], wv["S&P 500 Valuation"], wv["Buyback & Dividend computation"]
    t = ("2023-01" if str(f["A18"].value).startswith("How much") else "2023-07" if "^0.5" in f["C32"].value
         else "2023-10" if "^0.25" in f["C32"].value else "2024-01")
    o = 32 if t == "2023-01" else 33  # row of 'Intrinsic Value of Index'
    inp = {"template": t, "index_level": v["B3"].value, "riskfree_rate": v["B5"].value, "erp": v["B11"].value,
           "roe": v["B13"].value, "base_earnings": v["B15"].value, "riskfree_change": v["B6"].value == "Yes",
           "riskfree_after_5": v["B7"].value, "riskfree_change_mode": MODES.get(v["B8"].value, "end"),
           "cape_denominator": _avg_range(f, bb, f"B{o + 3}")}
    if t == "2023-01":
        inp |= {"earnings_forecasts": [v["B16"].value, v["B17"].value], "base_payout": v["B18"].value}
        flag, adj = v["B19"].value, v["B20"].value
    else:
        n = 2 if t == "2024-01" else 3
        inp |= {"base_cash_return": v["B16"].value, "earnings_forecasts": [v[f"B{17 + k}"].value for k in range(n)]}
        flag, adj = v["B20"].value, v["B21"].value
    if flag == "Yes":
        inp["earnings_adjustment"] = adj
    exp = {"value": v[f"B{o}"].value, "trailing_pe": v[f"B{o + 1}"].value, "forward_pe": v[f"B{o + 2}"].value,
           "cape": v[f"B{o + 3}"].value, "over_under": v[f"B{o + 5}"].value}
    return [{"name": p.stem, "source": _src(p), "kind": "index", "inputs": inp, "expected": exp}]


def _b_erpcalc(p: Path, wf, wv) -> list:
    f, v = wf["Impl premium calculator"], wv["Impl premium calculator"]
    bb, eg = wv["Buyback & Dividend computation"], wv["Expected growth rate"]
    base_row = int(f["G29"].value.rsplit("H", 1)[1])  # ='Expected growth rate'!H4 -> consensus base year row
    cons = [eg[f"H{base_row + k}"].value for k in range(3)]
    inp = {"template": "2025-09", "index_level": v["B5"].value, "riskfree_rate": v["B10"].value,
           "erp": v["B13"].value, "default_spread": v["B12"].value, "month": v["B3"].value,
           "base_earnings": v["B33"].value, "trailing_earnings": v["G28"].value, "base_cash_flow": v["B35"].value,
           "adjust_payout": v["B16"].value == "Yes", "roe": bb["F26"].value,
           "cape_denominator": _avg_range(f, bb, "C41")}
    inp |= {"consensus": cons} if v["B8"].value == "TD" else {"growth": v["B9"].value}
    if v["B14"].value == "Yes":
        inp["long_term_growth"] = v["B15"].value
    cases = [{"name": p.stem, "source": _src(p), "kind": "index", "inputs": inp,
              "expected": {"value": v["C39"].value, "trailing_pe": v["C40"].value, "cape": v["C41"].value}},
             {"name": f"{p.stem} top-down growth", "source": _src(p), "kind": "topdown",
              "inputs": {"consensus": cons, "riskfree_rate": inp["riskfree_rate"]},
              "expected": {"growth": eg["G11"].value}}]
    exp = {"value": v["B54"].value}
    if abs(v["B54"].value / inp["index_level"] - 1) < 1e-6:  # the sheet's goal seek was run on these inputs
        exp |= {"implied_erp": v["C45"].value, "implied_expected_return": v["C43"].value,
                "implied_erp_default_adj": v["C44"].value}
    cases.append({"name": f"{p.stem} at Damodaran's implied ERP", "source": _src(p), "kind": "index",
                  "inputs": {**inp, "erp": v["C45"].value}, "expected": exp})
    return cases


def _b_breakeven(p: Path, wv, data: dict) -> list:
    ws, st, wt = wv["Breakeven Calculation"], wv["Revenue Sensitivity Table"], wv["Waiting Table"]
    keys = [("market cap to match", "value_to_match"), ("enterprise value", "value_to_match"),
            ("riskiness", "risk_class"), ("stable nominal", "stable_growth"), ("net margin", "margin"),
            ("operating margin", "margin"), ("return on", "return_on_capital"), ("number of years before", "years"),
            ("average cash yield", "cash_yield"), ("revenues", "current_revenues"), ("number of years after", "life"),
            ("if finite", "finite_years"), ("breakeven revenue if", "now"), ("breakeven revenue with", "wait"),
            ("cagr", "cagr"), ("cost of", "rate")]
    got, row = {}, {}
    for r in range(1, 40):
        a = str(ws.cell(r, 1).value or "").strip().lower()
        k = next((k for pre, k in keys if a.startswith(pre) and (pre != "revenues" or a == "revenues")), None)
        if k and k not in got:
            got[k], row[k] = ws.cell(r, 2).value, f"B{r}"
    basis = "equity" if str(ws["A2"].value).startswith("Market Cap") else "firm"
    risk = {ws[f"F{r}"].value: ws[f"G{r}"].value for r in range(2, 8)}
    axes = {"grid_margins": [st[f"B{r}"].value for r in range(4, 14)],
            "grid_returns": [st.cell(3, c).value for c in range(3, 11)],
            "wait_years": [wt[f"C{r}"].value for r in range(4, 12)],
            "wait_yields": [wt.cell(3, c).value for c in range(4, 11)]}
    if basis not in data.setdefault("breakeven_defaults", {}):
        data["breakeven_defaults"][basis] = {
            "source": p.name, "url": _src(p), "as_of": "2025-12 post" if basis == "equity" else "2026-08 post",
            "risk_classes": risk, "risk_class": got["risk_class"][:2],
            "stable_growth": got["stable_growth"], "years_to_maturity": got["years"], "cash_yield": got["cash_yield"],
            "append_own": basis == "firm", **axes,
            "note": "inflation input in the source sheet feeds no formula; D6 'Direct input' = pass discount_rate"}
    inp = {"basis": basis, "value_to_match": got["value_to_match"], "margin": got["margin"],
           "return_on_capital": got["return_on_capital"], "risk_class": got["risk_class"][:2],
           "stable_growth": got["stable_growth"], "years_to_maturity": got["years"], "cash_yield": got["cash_yield"],
           **axes}
    if got["risk_class"].startswith("D6"):
        inp["discount_rate"] = got["rate"]
    if got.get("current_revenues"):
        inp["current_revenues"] = got["current_revenues"]
    if str(got.get("life", "")).startswith("Finite"):
        inp["finite_years"] = got["finite_years"]
        inp["replicate_sheet_bugs"] = True  # calibration reproduces the source sheet as written
    n_r, n_c = (11, 9) if basis == "firm" else (10, 8)
    exp = {"discount_rate": got["rate"], "breakeven_revenue_now": got["now"], "breakeven_revenue_wait": got["wait"],
           "grid": [[st.cell(4 + i, 3 + j).value for j in range(n_c)] for i in range(n_r)],
           "wait_grid": [[wt.cell(4 + i, 4 + j).value for j in range(7)] for i in range(8)]}
    if got.get("cagr") is not None:
        exp["cagr_needed"] = got["cagr"]
    return [{"name": p.stem, "source": _src(p), "kind": "breakeven", "inputs": inp, "expected": exp, "_cells": row}]


def _b_pe(p: Path, wf, wv, data: dict) -> list:
    import re  # noqa: PLC0415
    f, v = wf.active, wv.active
    isf = lambda c: isinstance(f[c].value, str) and f[c].value.startswith("=")
    rows, year_of, exp = [], {}, {}
    fields = {"B": "pe", "C": "norm_pe", "D": "cape", "E": "shiller_pe", "F": "tbond_pe",
              "G": "shiller_to_tbond_pe", "H": "price_dividends", "I": "price_cash_returned"}
    for r in range(2, 200):
        y = v[f"A{r}"].value
        if not isinstance(y, (int, str)) or not str(y)[:4].isdigit():
            break
        year_of[r] = y
        row = {"year": y, "index": v[f"P{r}"].value, "earnings": v[f"Q{r}"].value, "dividends": v[f"R{r}"].value,
               "tbond": v[f"L{r}"].value, "inflation": v[f"U{r}"].value}
        if v[f"S{r}"].value:
            row["cash_returned"] = v[f"S{r}"].value
        if not isf(f"K{r}"):
            row["earnings_yield"] = v[f"K{r}"].value
        if not isf(f"E{r}"):
            row["shiller_pe"] = v[f"E{r}"].value
        rows.append(row)
        exp |= {f"{y}.{n}": v[f"{c}{r}"].value for c, n in fields.items() if v[f"{c}{r}"].value is not None}
    col_field = {**fields, "K": "earnings_yield", "L": "tbond"}
    periods = []
    for r in range(r + 1, r + 12):
        for c in "CDEFGHIJKLMN":
            m = re.fullmatch(r"=AVERAGE\(([A-Z])(\d+):[A-Z](\d+)\)", str(f[f"{c}{r}"].value))
            if m and m[1] in col_field:
                y0, y1 = year_of[int(m[2])], year_of[int(m[3])]
                key = f"sheet {c}{r} {v[f'B{r}'].value if c <= 'H' else v[f'J{r}'].value} = {m[1]}{m[2]}:{m[1]}{m[3]}"
                periods.append([key, [_yr(y0), _yr(y1)], col_field[m[1]]])
                exp[f"{key}.{col_field[m[1]]}"] = v[f"{c}{r}"].value
    if not data.get("pe_history") or _yr(rows[-1]["year"]) >= _yr(data["pe_history"]["rows"][-1]["year"]):
        data["pe_history"] = {"source": p.name, "url": _src(p), "as_of": str(rows[-1]["year"]),
                              "note": "earnings: S&P 500 operating earnings since 1988, reported before; "
                                      "inflation: CPI (FRED); shiller_pe given = Shiller's series", "rows": rows}
    return [{"name": p.stem, "source": _src(p), "kind": "pe", "expected_fields": list(fields.values()),
             "inputs": {"rows": rows, "periods": periods}, "expected": exp}]


def _b_leaders(p: Path, wv, data: dict) -> list:
    ws, rows, exp = wv.active, [], {}
    for r in range(2, 40):
        y = ws[f"A{r}"].value
        if y is None:
            break
        y = y if isinstance(y, int) else str(y)[:10]
        rows.append({"year": y, **{k: ws[f"{c}{r}"].value or 0 for c, k in zip("BCDEFH", (
            "apple", "microsoft", "alphabet", "amazon", "meta", "all_us"))}})
        for c, k in zip("GIKLM", ("five", "share", "change_five", "change_all", "share_of_change")):
            if ws[f"{c}{r}"].value is not None:
                exp[f"{y}.{k}"] = ws[f"{c}{r}"].value
    data["leaders"] = {"source": p.name, "url": _src(p), "as_of": str(rows[-1]["year"]),
                       "units": "US$ billions, end of year", "rows": rows}
    return [{"name": p.stem, "source": _src(p), "kind": "leaders", "inputs": {"rows": rows}, "expected": exp}]


def _branches(p: Path, base: dict) -> list:
    """Damodaran's template recalculated by pycel under input branches the saved files never use."""
    from pycel import ExcelCompiler  # noqa: PLC0415
    t = base["inputs"].get("template")
    if t and t != "2025-09":
        sheet, tgt = "S&P 500 Valuation", "B32" if t == "2023-01" else "B33"
        flag = ("B19", "No", None) if t == "2023-01" else ("B20", "Yes", -0.1)
        sc = [(f"rf to 3% {m}", {"B6": "Yes", "B7": 0.03, "B8": lab},
               {"riskfree_change": True, "riskfree_after_5": 0.03, "riskfree_change_mode": m}) for lab, m in MODES.items()]
        sc.append((f"earnings adjustment {flag[1]}", {flag[0]: flag[1], ("B20" if t == "2023-01" else "B21"): -0.1},
                   {"earnings_adjustment": flag[2]}))
        # tgt = SUM(PV row); pycel's cycle mode leaves that SUM stale when the sheet's E24<->F24 IF-loop flips,
        # so the branch value is the sum of the recalculated PV cells (C..G of the row above)
        exp = {"value": [f"{c}{int(tgt[1:]) - 1}" for c in "CDEFG"]}
    elif t == "2025-09":
        sheet, exp = "Impl premium calculator", {"value": "C39"}
        sc = [("long-term growth 3%", {"B14": "Yes", "B15": 0.03}, {"long_term_growth": 0.03}),
              ("payout not adjusted", {"B16": "No"}, {"adjust_payout": False}),
              ("base = reported LTM", {"B4": "Reported LTM"}, {"base_earnings": base["inputs"]["trailing_earnings"]}),
              ("month June", {"B3": "June"}, {"month": "June"})]
    else:
        sheet, row = "Breakeven Calculation", base["_cells"]
        exp = {k: row[c] for k, c in (("breakeven_revenue_now", "now"), ("breakeven_revenue_wait", "wait"),
                                      ("cagr_needed", "cagr")) if c in row}
        rc = row["risk_class"]
        sc = [("risk D4", {rc: "D4. Below-average risk"}, {"risk_class": "D4"}),
              ("direct rate 11%", {rc: "D6. Direct input", "G7": 0.11}, {"discount_rate": 0.11})]
        if "life" in row:
            sc.append(("finite life 20 yrs", {row["life"]: "Finite (specify years in cell below)",
                                              row["finite_years"]: 20}, {"finite_years": 20, "replicate_sheet_bugs": True}))
    out = []
    for name, cells, ours in sc:
        xl = ExcelCompiler(filename=str(p), cycles={"iterations": 100, "tolerance": 1e-10})  # the source files have cycles
        ev = lambda c: sum(float(xl.evaluate(f"{sheet}!{x}")) for x in c) if isinstance(c, list)             else float(xl.evaluate(f"{sheet}!{c}"))
        k0 = next(iter(exp))
        print(f"  {name}: pycel baseline {k0} {ev(exp[k0]):.6f} vs cached {base['expected'][k0]:.6f}", flush=True)
        for c in cells:
            xl.evaluate(f"{sheet}!{c}")
        for c, val in cells.items():
            xl.set_value(f"{sheet}!{c}", val)
        xl.recalculate()
        inp = {k: val for k, val in {**base["inputs"], **ours}.items() if val is not None}
        out.append({"name": f"{p.stem} branch: {name}", "source": _src(p) + " (pycel recalculation)",
                    "kind": base["kind"], "inputs": inp, "expected": {k: ev(c) for k, c in exp.items()}})
    return out


def build_cases(paths: list) -> int:
    """pycel recursion through the sheet's 'Historical ERP' chain needs a deep stack: run in a big-stack thread."""
    import threading  # noqa: PLC0415
    sys.setrecursionlimit(60_000)
    threading.stack_size(128 * 1024 * 1024)  # Windows limit
    res = []
    th = threading.Thread(target=lambda: res.append(_build(paths)))
    th.start()
    th.join()
    return res[0] if res else 1


def _build(paths: list) -> int:
    import warnings  # noqa: PLC0415

    import openpyxl  # noqa: PLC0415
    warnings.filterwarnings("ignore")
    data = json.loads(DATA.read_text(encoding="utf-8")) if DATA.exists() else {}
    data.pop("breakeven_defaults", None)
    cases = []
    for p in sorted(map(Path, paths), key=lambda x: x.name):
        wf, wv = openpyxl.load_workbook(p), openpyxl.load_workbook(p, data_only=True)
        names = wv.sheetnames
        if "S&P 500 Valuation" in names:
            new = _b_sp500(p, wf, wv)
        elif "Impl premium calculator" in names:
            new = _b_erpcalc(p, wf, wv)
            iv, ans = wv["Impl premium calculator"], wv["Answers"]
            data["months_to_year_end"] = {"source": f"{p.name} Answers!G2:H13 (months left in the year)",
                                          "url": _src(p), "as_of": p.stem,
                                          "table": {ans[f"G{r}"].value: ans[f"H{r}"].value for r in range(2, 14)}}
            if p.stem == "ERPJan26" or "erp_choices" not in data:
                hist = [wv["Implied ERP- Annual since 1960"][f"E{r}"].value for r in range(3, 100)]
                ch = {iv[f"H{r}"].value: {"value": iv[f"G{r}"].value, "label": iv[f"D{r}"].value} for r in range(22, 27)}
                ch["ERPLow"] = {"value": min(x for x in hist if isinstance(x, float)),
                                "label": "Historical Low ERP (1960-Current)"}
                data["erp_choices"] = {"source": f"{p.name} 'Impl premium calculator'!D22:H26; low = MIN of "
                                       "'Implied ERP- Annual since 1960'!E", "url": _src(p), "as_of": p.stem,
                                       "choices": ch}
        elif "Breakeven Calculation" in names:
            new = _b_breakeven(p, wv, data)
        elif p.stem.startswith("PEHist"):
            new = _b_pe(p, wf, wv, data)
        elif p.stem.startswith("Fearsome"):
            new = _b_leaders(p, wv, data)
        else:
            print("skipped (not a market template):", p.name)
            continue
        if new[0]["kind"] in ("index", "breakeven"):
            print("pycel branches:", p.name, flush=True)
            branch = _branches(p, new[0])
            new[0].pop("_cells", None)
            new += branch
        cases += new
    cases.append({"name": "Nvidia $5T market cap, 2025-12 post text", "kind": "breakeven",
                  "source": "https://aswathdamodaran.blogspot.com/2025/12/trillion-dollar-market-caps-fairy-tale.html",
                  "inputs": {"basis": "equity", "value_to_match": 5e6, "margin": 0.5301, "return_on_capital": 0.6444,
                             "risk_class": "D2", "years_to_maturity": 5, "cash_yield": 0.02},
                  "expected": {"breakeven_revenue_now": 483380, "breakeven_revenue_wait": 677970}})
    DATA.write_text(json.dumps(data, indent=1, default=str), encoding="utf-8")
    CASES.write_text(json.dumps(cases, indent=1, default=str), encoding="utf-8")
    print(f"wrote {len(cases)} cases -> {CASES}\nwrote data -> {DATA}")
    return selftest()


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("inputs", nargs="?")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--pe", action="store_true")
    ap.add_argument("--leaders", action="store_true")
    ap.add_argument("--from-case")
    ap.add_argument("--build-cases", nargs="+")
    a = ap.parse_args()
    if a.selftest:
        return selftest()
    if a.build_cases:
        return build_cases(a.build_cases)
    if a.from_case:
        cases = json.loads(CASES.read_text(encoding="utf-8"))
        hit = [c for c in cases if c["name"].lower().startswith(a.from_case.lower()) and c["kind"] in ("index", "breakeven")]
        if not hit:
            print("cases:", ", ".join(c["name"] for c in cases if c["kind"] in ("index", "breakeven")))
            return 1
        c = hit[0]
        print(json.dumps({"_source": c["source"], "model": c["kind"], **c["inputs"]}, indent=1))
        return 0
    if a.pe:
        d = load()["pe_history"]
        t = pe_history(d["rows"])
        last = _yr(t[-1]["year"])
        periods = [(1970, 1990), (1991, 2000), (2001, 2010), (2011, 2020), (2021, last), (1960, last)]
        print(json.dumps({"table": t}, indent=1) if a.json else report_pe(t, periods))
        return 0
    if a.leaders:
        t = leaders_share(load()["leaders"]["rows"])
        print(json.dumps(t, indent=1) if a.json else "\n".join(
            f"{str(r['year'])[:10]:10s} five {r['five']:9,.0f}  all US {r['all_us']:9,.0f}  share {r['share']:6.2%}"
            + (f"  share of change {r['share_of_change']:7.2%}" if "share_of_change" in r else "") for r in t))
        return 0
    inp = {k: v for k, v in json.loads(Path(a.inputs).read_text(encoding="utf-8")).items() if not k.startswith("_")}
    if inp.get("model") == "breakeven":
        out = breakeven(inp)
        print(json.dumps(out, indent=1) if a.json else report_breakeven(out))
        return 0
    out = value_index(inp)
    out["implied"] = implied_erp(inp)
    print(json.dumps(out, indent=1) if a.json else report_index(inp, out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
