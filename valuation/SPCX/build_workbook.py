"""Build the SPCX valuation workbook (investment-banking conventions) from the aswath-finance valuation.

Every calculation cell is a live Excel formula mirroring the verified engines (ginzu.py + Damodaran's SpaceX segment
mechanics + cost_of_capital.py); the Checks tab ties the workbook to the engine's own outputs.
Run with finance/.venv python, then finalize_workbook.ps1 (Excel) adds the data tables and recalculates.

Usage: ../../.venv/Scripts/python build_workbook.py
"""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from spacex_model import breakeven_multiplier, build, per_share, scenarios  # noqa: E402

OUT = HERE / "SPCX_Valuation_2026-09-25.xlsx"
BASE = json.loads((HERE / "sep_2026_inputs.json").read_text(encoding="utf-8"))
WALK = json.loads((HERE / "value_walk.json").read_text(encoding="utf-8"))

# ---------------------------------------------------------------- engine reference values (for Checks)
ENG_VPS = per_share(BASE)
ENG_WACC = build(BASE)["initial_cost_of_capital"]
ENG_SCEN = [(n, per_share(d)) for n, d in scenarios(BASE)]
ENG_BE = breakeven_multiplier(BASE)

# ---------------------------------------------------------------- style
NAVY, MID, LIGHT, GREY = "0B2545", "13315C", "DCE6F1", "F2F2F2"
F = "Arial"
FONT = Font(name=F, size=10)
BLUE = Font(name=F, size=10, color="0000FF")
GREEN = Font(name=F, size=10, color="008000")
BOLD = Font(name=F, size=10, bold=True)
WHITE_B = Font(name=F, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=F, size=16, bold=True, color=NAVY)
SUB = Font(name=F, size=10, italic=True, color="595959")
NOTE = Font(name=F, size=8, italic=True, color="595959")
FILL_NAVY, FILL_LIGHT, FILL_GREY = (PatternFill("solid", fgColor=c) for c in (NAVY, LIGHT, GREY))
FILL_FLEX = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="808080")
TOP = Border(top=THIN)
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
M = '#,##0_);(#,##0);"–"_)'          # $ millions
PS = '$#,##0.00_);($#,##0.00)'        # per share
PCT = '0.0%_);(0.0%);"–"_)'
PCT2 = '0.00%_);(0.00%);"–"_)'
X = '0.00"x"'
NUM2 = '0.00_);(0.00)'
NUM4 = '0.0000'


def title(ws, text, sub):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws["B2"], ws["B3"] = text, sub
    ws["B2"].font, ws["B3"].font = TITLE, SUB


def section(ws, row, text, first="B", last="N"):
    for col in range(ord(first), ord(last) + 1):
        c = ws[f"{chr(col)}{row}"]
        c.fill = FILL_NAVY
    ws[f"{first}{row}"] = text
    ws[f"{first}{row}"].font = WHITE_B


def hdr(ws, row, labels: dict):
    for ref, text in labels.items():
        c = ws[f"{ref}{row}"]
        c.value, c.font, c.fill = text, BOLD, FILL_LIGHT
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")


def put(ws, ref, value, kind="calc", fmt=None, bold=False):
    c = ws[ref]
    c.value = value
    c.font = {"input": BLUE, "link": GREEN, "label": FONT, "calc": FONT}[kind]
    if bold:
        c.font = Font(name=F, size=10, bold=True, color=c.font.color)
    if fmt:
        c.number_format = fmt
    return c


def label(ws, ref, text, bold=False, indent=0):
    c = put(ws, ref, text, "label", bold=bold)
    c.alignment = Alignment(indent=indent)
    return c


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def printable(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


YEARS = [chr(c) for c in range(ord("D"), ord("M") + 1)]  # years 1..10 in columns D..M (C = base, N = terminal)

wb = Workbook()
cover = wb.active
cover.title = "Cover"
summ, inp, scen, sens, wacc, segs, dcf, diag, chk, src = (wb.create_sheet(n) for n in (
    "Summary", "Inputs", "Scenarios", "Sensitivity", "WACC", "Segments", "DCF", "Diagnostics", "Checks", "Sources"))

# ================================================================= INPUTS
title(inp, "Inputs & Assumptions", "$ in millions except per-share data | Blue = hard-coded input (source in column J) | "
      "Valuation date 25-Sep-2026")
widths(inp, {"B": 46, "C": 14, "D": 14, "E": 13, "F": 13, "G": 12, "H": 12, "I": 12, "J": 70})
S = BASE["_sources"]
r = 5
section(inp, r, "Market data", last="J")
rows_mkt = [
    ("Valuation date", date(2026, 9, 25), "dd-mmm-yyyy", "Close of 25-Sep-2026 (last trading day before analysis)"),
    ("Share price ($)", 148.68, PS, "StockAnalysis, close 25-Sep-2026 4:00pm ET: https://stockanalysis.com/stocks/spcx/"),
    ("Shares outstanding (m)", 13181.78, M, "10-Q cover 28-Jul-2026: Class A 7,696.29m + Class B 5,485.49m (StockAnalysis shows 13.57bn)"),
    ("52-week low ($)", 104.83, PS, "StockAnalysis, 25-Sep-2026"),
    ("52-week high ($)", 225.64, PS, "StockAnalysis, 25-Sep-2026"),
    ("Analyst average price target ($)", 222.42, PS, "StockAnalysis: consensus Buy, 35 analysts (third-party aggregate)"),
    ("Damodaran June 2026 value per share ($)", 97.83, PS, "Damodaran, SpaceX2026IPOUpdated.xlsx (reproduced exactly by the engine)"),
]
R = {}
for i, (lab, val, fmt, source) in enumerate(rows_mkt, start=r + 1):
    label(inp, f"B{i}", lab)
    put(inp, f"C{i}", val, "input", fmt)
    put(inp, f"J{i}", source, "label").font = NOTE
    R[lab] = i
r = r + len(rows_mkt) + 2
section(inp, r, "Balance sheet & income (Q2-2026 10-Q; TTM to 30-Jun-2026)", last="J")
rows_bs = [
    ("Cash & marketable securities", 100009, S["cash"]),
    ("Debt incl. finance leases", 39364, S["book_debt"]),
    ("Book value of equity", 127224, S["book_equity"]),
    ("Minority interests", 0, "Not identified in the 10-Q balance sheet"),
    ("Non-operating assets (cross holdings)", 0, "None identified"),
    ("IPO proceeds not yet received", 0, S["cash"]),
    ("Operating income (EBIT), TTM", -3732, S["ebit (TTM)"]),
    ("Interest expense, TTM", 2380, "FILING: FY2025 1,945 + H1-2026 1,293 - H1-2025 858"),
    ("Revenue, FY2025", 18674, "Damodaran's June 2026 model (prospectus FY2025 segment revenues)"),
    ("Net operating loss carried forward", 0, "Ginzu default (no NOL override)"),
    ("Value of employee options", 0, S["tax rates, options"]),
]
for i, (lab, val, source) in enumerate(rows_bs, start=r + 1):
    label(inp, f"B{i}", lab)
    put(inp, f"C{i}", val, "input", M)
    put(inp, f"J{i}", source, "label").font = NOTE
    R[lab] = i
r = r + len(rows_bs) + 2
section(inp, r, "R&D capitalization (Ginzu R&D converter)", last="J")
rd = [("R&D expense, current (TTM)", 12190, S["rd.current (TTM)"])] + [
    (f"R&D expense, year -{k}", v, "Damodaran's June 2026 model inputs") for k, v in enumerate(BASE["rd"]["past"], start=1)]
for i, (lab, val, source) in enumerate(rd, start=r + 1):
    label(inp, f"B{i}", lab)
    put(inp, f"C{i}", val, "input", M)
    put(inp, f"J{i}", source, "label").font = NOTE
R["rd_cur"], R["rd_past0"] = r + 1, r + 2
r = r + len(rd) + 1
label(inp, f"B{r}", "Amortizable life (years)")
put(inp, f"C{r}", 5, "input", "0")
put(inp, f"J{r}", "Damodaran's June 2026 model", "label").font = NOTE
R["rd_life"] = r
r += 2
section(inp, r, "Taxes and discount-rate inputs", last="J")
rows_tax = [
    ("Effective tax rate", 0.10, PCT, S["tax rates, options"]),
    ("Marginal tax rate", 0.25, PCT, S["tax rates, options"]),
    ("Riskfree rate (US 10-year Treasury)", 0.0517, PCT2, S["riskfree_rate"]),
    ("Mature market ERP", 0.042, PCT2, S["mature_erp"]),
    ("Terminal cost of capital", 0.0886, PCT2, S["terminal_cost_of_capital"]),
    ("Terminal return on capital (ROIC)", 0.15, PCT, "Damodaran's June 2026 override (15%)"),
    ("Probability of failure", 0.0, PCT, "Damodaran's June 2026 model (no failure)"),
    ("Distress proceeds (% of fair value)", 0.5, PCT, "Ginzu template default"),
    ("June 2026 model riskfree rate (reference)", 0.0456, PCT2, "Damodaran's June 2026 model"),
    ("June 2026 model terminal cost of capital (reference)", 0.0825, PCT2, "Damodaran's June 2026 model"),
]
for i, (lab, val, fmt, source) in enumerate(rows_tax, start=r + 1):
    label(inp, f"B{i}", lab)
    put(inp, f"C{i}", val, "input", fmt)
    put(inp, f"J{i}", source, "label").font = NOTE
    R[lab] = i
r = r + len(rows_tax) + 2
section(inp, r, "Segment story (Damodaran's June 2026 story; base year rolled to TTM)", last="J")
hdr(inp, r + 1, {"B": "Segment", "C": "Base revenue (TTM)", "D": "2036 revenue", "E": "Base margin",
                 "F": "Target margin", "G": "Sales/capital yrs 1-5", "H": "Sales/capital yrs 6-10", "I": "Starts after yr 5"})
R["seg0"] = r + 2
for i, sgm in enumerate(BASE["segments"]):
    rr = r + 2 + i
    label(inp, f"B{rr}", sgm["name"])
    put(inp, f"C{rr}", sgm.get("base_revenue", 0), "input", M)
    put(inp, f"D{rr}", sgm["revenue_2036"], "input", M)
    put(inp, f"E{rr}", sgm["base_margin"], "input", PCT)
    put(inp, f"F{rr}", sgm["target_margin"], "input", PCT)
    put(inp, f"G{rr}", sgm["sales_to_capital_1_5"], "input", NUM2)
    put(inp, f"H{rr}", sgm["sales_to_capital_6_10"], "input", NUM2)
    put(inp, f"I{rr}", 1 if sgm.get("starts_after_year5") else 0, "input", "0")
put(inp, f"J{r + 2}", S["segment base revenues (TTM to 2026-06-30)"], "label").font = NOTE
put(inp, f"J{r + 3}", S["story (2036 segment revenues, target margins, sales-to-capital, convergence year, base segment margins, terminal ROIC 15%)"], "label").font = NOTE
r = r + 2 + len(BASE["segments"])
label(inp, f"B{r}", "Year of margin convergence")
put(inp, f"C{r}", 10, "input", "0")
R["conv"] = r
r += 2
section(inp, r, "Cost of capital inputs (Damodaran's method; Damodaran Jan-2026 global industry data)", last="J")
hdr(inp, r + 1, {"B": "Business (industry)", "C": "Revenue weight (TTM)", "D": "Unlevered beta", "E": "EV/Sales",
                 "F": "June-model weight"})
biz = [("Aerospace/Defense", 4056, 1.166684299, 3.385246124, 4100), ("Telecom. Services", 13873, 0.496255717, 2.263007784, 11400),
       ("Software (System & Application)", 5115, 1.330950255, 9.912802758, 100)]
R["biz0"] = r + 2
for i, (n, w, b, evs, jw) in enumerate(biz):
    rr = r + 2 + i
    label(inp, f"B{rr}", n)
    put(inp, f"C{rr}", w, "input", M)
    put(inp, f"D{rr}", b, "input", NUM4)
    put(inp, f"E{rr}", evs, "input", NUM2)
    put(inp, f"F{rr}", jw, "input", M)
put(inp, f"J{r + 2}", "Industry betas & EV/Sales: Damodaran fcffsimpleginzu.xlsx global industry averages (Jan-2026)", "label").font = NOTE
put(inp, f"J{r + 3}", "Revenue weights = TTM segments (Space, Connectivity, AI); June weights = the stale April mix in the June model", "label").font = NOTE
r = r + 2 + len(biz) + 1
hdr(inp, r, {"B": "Operating country", "C": "Revenue weight", "D": "ERP (Jul-2026)", "E": "ERP (Jan-2026)"})
ctry = [("United States", 15207, 0.0442, 0.0446), ("China", 6662, 0.0507, 0.0514), ("Rest of the world", 9667, 0.0618, 0.0618)]
R["ctry0"] = r + 1
for i, (n, w, e1, e0) in enumerate(ctry):
    rr = r + 1 + i
    label(inp, f"B{rr}", n)
    put(inp, f"C{rr}", w, "input", M)
    put(inp, f"D{rr}", e1, "input", PCT2)
    put(inp, f"E{rr}", e0, "input", PCT2)
put(inp, f"J{r + 1}", "Weights: Damodaran's June sheet (10-Q has no geographic split). ERPs: Damodaran's ctrypremJuly26.xlsx (country_risk.py); rest of world = 6.18% (June sheet)", "label").font = NOTE
r = r + 1 + len(ctry) + 1
for lab, val, fmt, source in [("Bond rating", "Ba2/BB", None, "Damodaran's June 2026 model"),
                              ("Default spread for rating", 0.018395, PCT2, "Damodaran's Jan-2026 rating/spread table"),
                              ("Average debt maturity (years)", 3, "0", "Damodaran's June 2026 model")]:
    label(inp, f"B{r}", lab)
    put(inp, f"C{r}", val, "input", fmt)
    put(inp, f"J{r}", source, "label").font = NOTE
    R[lab] = r
    r += 1
r += 1
section(inp, r, "Industry benchmarks for Diagnostics (US Aerospace/Defense, Damodaran Jan-2026)", last="J")
for i, (lab, val, fmt) in enumerate([("Revenue growth (5-yr average)", 0.110999, PCT), ("Pre-tax operating margin (R&D adjusted)", 0.087654, PCT),
                                     ("Sales / capital", 2.743410, NUM2), ("Cost of capital", 0.076015, PCT)], start=r + 1):
    label(inp, f"B{i}", lab)
    put(inp, f"C{i}", val, "input", fmt)
    put(inp, f"J{i}", "Damodaran fcffsimpleginzu.xlsx, Industry Averages (US), Jan-2026", "label").font = NOTE
    R[lab] = i
inp.freeze_panes = "C5"
printable(inp)

I = lambda key, col="C": f"Inputs!${col}${R[key]}"  # noqa: E731
SEG = lambda i, col: f"Inputs!${col}${R['seg0'] + i}"  # noqa: E731

# ================================================================= SENSITIVITY flex cells (referenced by the model)
title(sens, "Sensitivity Analysis", "Live Excel data tables on value per share ($). Flex cells (yellow) must stay neutral "
      "for the base case; the data tables vary them.")
widths(sens, {"B": 40, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 13, "I": 40})
section(sens, 4, "Flex cells (neutral values = base case)", last="I")
flex = [("Change in initial cost of capital", 0, PCT2), ("Change in terminal cost of capital", 0, PCT2),
        ("2036 revenue multiplier (all segments)", 1, X), ("Change in target margin (all segments)", 0, PCT),
        ("Change in riskfree rate", 0, PCT2)]
for i, (lab, val, fmt) in enumerate(flex, start=5):
    label(sens, f"B{i}", lab)
    c = put(sens, f"C{i}", val, "input", fmt)
    c.fill = FILL_FLEX
FX = {k: f"Sensitivity!$C${i}" for i, k in enumerate(("d_init", "d_term", "rev_mult", "d_margin", "d_rf"), start=5)}

# ================================================================= SCENARIOS
title(scen, "Scenario Analysis", "One change at a time vs the base case (same set as spacex_model.py --sensitivity). "
      "Set the active scenario in D5; column K is a live data table.")
widths(scen, {"B": 5, "C": 50, "D": 11, "E": 11, "F": 11, "G": 12, "H": 11, "I": 11, "J": 11, "K": 12, "L": 11, "N": 8, "O": 12})
label(scen, "C5", "Active scenario (enter 1-12)", bold=True)
sel = put(scen, "D5", 1, "input", "0")
sel.fill, sel.border = FILL_FLEX, BOX
put(scen, "E5", '=INDEX($C$8:$C$19,$D$5)', "calc")
hdr(scen, 7, {"B": "#", "C": "Scenario", "D": "Riskfree", "E": "Terminal cost of capital", "F": "Beta weight set (1=TTM, 2=June)",
              "G": "xAI 2036 revenue", "H": "Starlink target margin", "I": "xAI sales/capital yrs 1-5",
              "J": "2036 revenue multiplier", "K": "Value / share", "L": "vs price"})
scen.row_dimensions[7].height = 42
base_vals = {"D": ("link", f"={I('Riskfree rate (US 10-year Treasury)')}", PCT2),
             "E": ("link", f"={I('Terminal cost of capital')}", PCT2), "F": ("input", 1, "0"),
             "G": ("link", f"={SEG(2, 'D')}", M), "H": ("link", f"={SEG(1, 'F')}", PCT),
             "I": ("link", f"={SEG(2, 'G')}", NUM2), "J": ("input", 1, X)}
overrides = {2: {"F": ("input", 2)}, 3: {"D": ("link", f"={I('June 2026 model riskfree rate (reference)')}"),
                                         "E": ("link", f"={I('June 2026 model terminal cost of capital (reference)')}")},
             4: {"E": ("link", f"={I('June 2026 model terminal cost of capital (reference)')}")},
             5: {"E": ("calc", f"={I('Riskfree rate (US 10-year Treasury)')}+{I('Mature market ERP')}")},
             6: {"G": ("input", 80000)}, 7: {"G": ("input", 240000)}, 8: {"H": ("input", 0.5)}, 9: {"H": ("input", 0.7)},
             10: {"I": ("input", 0.5)}, 11: {"J": ("input", 0.75)}, 12: {"J": ("input", 1.25)}}
for k, (nm, _) in enumerate(ENG_SCEN, start=1):
    rr = 7 + k
    put(scen, f"B{rr}", k, "label", "0")
    put(scen, f"C{rr}", nm.replace(" (sep_2026_inputs.json)", ""), "label")
    for col, (kind, val, fmt) in base_vals.items():
        if k == 1:
            put(scen, f"{col}{rr}", val, kind, fmt)
        elif col in overrides.get(k, {}):
            kind2, v2 = overrides[k][col]
            put(scen, f"{col}{rr}", v2, kind2, fmt)
        else:
            put(scen, f"{col}{rr}", f"=${col}$8", "calc", fmt)
    put(scen, f"K{rr}", f"=O{rr}", "calc", PS)
    put(scen, f"L{rr}", f"=K{rr}/{I('Share price ($)')}-1", "calc", PCT)
put(scen, "N6", "Data table (Excel What-If; do not edit)", "label").font = NOTE
put(scen, "O7", "=ValuePerShare", "calc", PS)
for k in range(1, 13):
    put(scen, f"N{7 + k}", k, "label", "0")
    scen[f"O{7 + k}"].number_format = PS
section(scen, 21, "Live drivers (active scenario, feed the model)", last="L")
live = {"D": "Riskfree", "E": "Terminal cost of capital", "F": "Beta weight set", "G": "xAI 2036 revenue",
        "H": "Starlink target margin", "I": "xAI sales/capital yrs 1-5", "J": "2036 revenue multiplier"}
label(scen, "C22", "Active scenario values", bold=True)
for col, fmt in zip("DEFGHIJ", (PCT2, PCT2, "0", M, PCT, NUM2, X)):
    put(scen, f"{col}22", f"=INDEX({col}$8:{col}$19,$D$5)", "calc", fmt)
LV = {k: f"Scenarios!${c}$22" for k, c in zip(("rf", "term", "betaset", "xai_rev", "sl_margin", "xai_sc", "rev_mult"), "DEFGHIJ")}
put(scen, "C24", "Engine check (spacex_model.py --sensitivity): scenario values must equal column K; see Checks.",
    "label").font = NOTE
scen.freeze_panes = "D8"
printable(scen)

# ================================================================= WACC
title(wacc, "Cost of Capital", "Damodaran's method: multi-business unlevered beta (revenue x industry EV/Sales weights), ERP by "
      "operating country, rating-based cost of debt, market-value weights")
widths(wacc, {"B": 38, "C": 13, "D": 13, "E": 13, "F": 12, "G": 13, "H": 13})
section(wacc, 5, "Unlevered beta: multi-business", last="H")
hdr(wacc, 6, {"B": "Business", "C": "Revenue (live)", "D": "EV/Sales", "E": "Estimated value", "F": "Weight",
              "G": "Unlevered beta", "H": "Weighted beta"})
for i in range(3):
    rr, ir = 7 + i, R["biz0"] + i
    put(wacc, f"B{rr}", f"=Inputs!B{ir}", "link")
    put(wacc, f"C{rr}", f"=IF({LV['betaset']}=2,Inputs!$F${ir},Inputs!$C${ir})", "calc", M)
    put(wacc, f"D{rr}", f"=Inputs!$E${ir}", "link", NUM2)
    put(wacc, f"E{rr}", f"=C{rr}*D{rr}", "calc", M)
    put(wacc, f"F{rr}", f"=E{rr}/$E$10", "calc", PCT)
    put(wacc, f"G{rr}", f"=Inputs!$D${ir}", "link", NUM4)
    put(wacc, f"H{rr}", f"=F{rr}*G{rr}", "calc", NUM4)
label(wacc, "B10", "Total / unlevered beta", bold=True)
for col, fmt in (("C", M), ("E", M), ("F", PCT), ("H", NUM4)):
    c = put(wacc, f"{col}10", f"=SUM({col}7:{col}9)", "calc", fmt, bold=True)
    c.border = TOP
section(wacc, 12, "Equity risk premium: operating-country weights", last="H")
hdr(wacc, 13, {"B": "Country", "C": "Revenue", "D": "Weight", "E": "ERP", "F": "Weighted ERP"})
for i in range(3):
    rr, ir = 14 + i, R["ctry0"] + i
    put(wacc, f"B{rr}", f"=Inputs!B{ir}", "link")
    put(wacc, f"C{rr}", f"=Inputs!$C${ir}", "link", M)
    put(wacc, f"D{rr}", f"=C{rr}/$C$17", "calc", PCT)
    put(wacc, f"E{rr}", f"=Inputs!$D${ir}", "link", PCT2)
    put(wacc, f"F{rr}", f"=D{rr}*E{rr}", "calc", PCT2)
label(wacc, "B17", "Total / weighted ERP", bold=True)
put(wacc, "C17", "=SUM(C14:C16)", "calc", M, bold=True).border = TOP
put(wacc, "D17", "=SUM(D14:D16)", "calc", PCT, bold=True).border = TOP
put(wacc, "F17", "=SUM(F14:F16)", "calc", PCT2, bold=True).border = TOP
section(wacc, 19, "Cost of debt (actual rating) and market value of debt", last="H")
rows_w = [
    ("Riskfree rate (live)", f"={LV['rf']}+{FX['d_rf']}", PCT2, "calc"),
    ("Bond rating", f"={I('Bond rating')}", None, "link"),
    ("Default spread", f"={I('Default spread for rating')}", PCT2, "link"),
    ("Pre-tax cost of debt", "=C20+C22", PCT2, "calc"),
    ("Book value of debt", f"={I('Debt incl. finance leases')}", M, "link"),
    ("Interest expense (TTM)", f"={I('Interest expense, TTM')}", M, "link"),
    ("Average maturity (years)", f"={I('Average debt maturity (years)')}", "0", "link"),
    ("Market value of debt", "=C25*(1-(1+C23)^(-C26))/C23+C24/(1+C23)^C26", M, "calc"),
]
for i, (lab, f, fmt, kind) in enumerate(rows_w, start=20):
    label(wacc, f"B{i}", lab)
    put(wacc, f"C{i}", f, kind, fmt)
section(wacc, 29, "Capital structure (market values)", last="H")
rows_cs = [("Share price", f"={I('Share price ($)')}", PS, "link"), ("Shares outstanding (m)", f"={I('Shares outstanding (m)')}", M, "link"),
           ("Market value of equity", "=C30*C31", M, "calc"), ("Market value of debt", "=C27", M, "calc"),
           ("Debt / equity", "=C33/C32", PCT2, "calc"), ("Debt / capital", "=C33/(C32+C33)", PCT2, "calc"),
           ("Equity / capital", "=C32/(C32+C33)", PCT2, "calc")]
for i, (lab, f, fmt, kind) in enumerate(rows_cs, start=30):
    label(wacc, f"B{i}", lab)
    put(wacc, f"C{i}", f, kind, fmt)
section(wacc, 38, "Cost of capital", last="H")
rows_cc = [("Unlevered beta", "=H10", NUM4), ("Marginal tax rate", f"={I('Marginal tax rate')}", PCT),
           ("Levered beta  = Bu x (1 + (1 - t) x D/E)", "=C39*(1+(1-C40)*C34)", NUM4), ("Equity risk premium", "=F17", PCT2),
           ("Cost of equity  = Rf + Bl x ERP", "=C20+C41*C42", PCT2), ("After-tax cost of debt", "=C23*(1-C40)", PCT2),
           ("Cost of capital (initial)", "=C36*C43+C35*C44", PCT2)]
for i, (lab, f, fmt) in enumerate(rows_cc, start=39):
    label(wacc, f"B{i}", lab, bold=(i == 45))
    put(wacc, f"C{i}", f, "link" if f.startswith("=Inputs") else "calc", fmt, bold=(i == 45))
wacc["C45"].border = BOX
name(wb, "WACC_Initial", "WACC!$C$45")
printable(wacc, landscape=False)

# ================================================================= SEGMENTS
title(segs, "Segment Build", "Damodaran's June 2026 SpaceX mechanics: revenue steps toward year 5 then the 2036 target; margins "
      "converge linearly; reinvestment = revenue change / segment sales-to-capital. $ in millions.")
widths(segs, {"B": 34, **{c: 12 for c in "CDEFGHIJKLMN"}})
hdr(segs, 4, {"B": "", "C": "Base (TTM)", **{c: f"Year {k}" for k, c in enumerate(YEARS, start=1)}})
put(segs, "B5", "Year (t)", "label")
for k, c in enumerate(["C"] + YEARS):
    put(segs, f"{c}5", k, "label", "0").font = NOTE
label(segs, "B6", "Step toward anchor (June 2026 model path)")
for c, v in zip(("D", "E", "F", "G", "I", "J", "K", "L"), (0.2, 0.3, 0.4, 0.5) * 2):
    put(segs, f"{c}6", v, "input", "0.0")
section(segs, 8, "Live segment drivers", last="M")
hdr(segs, 9, {"B": "Segment", "C": "Base revenue", "D": "2036 revenue", "E": "Base margin", "F": "Target margin",
              "G": "S/C yrs 1-5", "H": "S/C yrs 6-10", "I": "Starts after yr 5"})
label(segs, "B15", "Revenue multiplier (scenario x flex)")
put(segs, "C15", f"={LV['rev_mult']}*{FX['rev_mult']}", "calc", X)
label(segs, "B16", "Target margin change (flex)")
put(segs, "C16", f"={FX['d_margin']}", "link", PCT)
label(segs, "B17", "Year of margin convergence")
put(segs, "C17", f"={I('conv') if False else 'Inputs!$C$' + str(R['conv'])}", "link", "0")
for i in range(4):
    rr = 10 + i
    put(segs, f"B{rr}", f"={SEG(i, 'B')}", "link")
    put(segs, f"C{rr}", f"={SEG(i, 'C')}", "link", M)
    rev_src = LV["xai_rev"] if i == 2 else SEG(i, "D")
    put(segs, f"D{rr}", f"={rev_src}*$C$15", "calc", M)
    put(segs, f"E{rr}", f"={SEG(i, 'E')}", "link", PCT)
    mg_src = LV["sl_margin"] if i == 1 else SEG(i, "F")
    put(segs, f"F{rr}", f"={mg_src}+$C$16", "calc", PCT)
    put(segs, f"G{rr}", f"={LV['xai_sc'] if i == 2 else SEG(i, 'G')}", "link", NUM2)
    put(segs, f"H{rr}", f"={SEG(i, 'H')}", "link", NUM2)
    put(segs, f"I{rr}", f"={SEG(i, 'I')}", "link", "0")
blocks = []
top = 20
for i in range(4):
    d = 10 + i  # driver row
    section(segs, top, f"=B{d}", last="M")
    segs[f"B{top}"].font = WHITE_B
    rv, mg, eb, ri = top + 1, top + 2, top + 3, top + 4
    label(segs, f"B{rv}", "Revenue", indent=1)
    label(segs, f"B{mg}", "Operating margin", indent=1)
    label(segs, f"B{eb}", "Operating income (EBIT)", indent=1)
    label(segs, f"B{ri}", "Reinvestment", indent=1)
    put(segs, f"C{rv}", f"=IF($I${d}=1,0,$C${d})", "calc", M)
    put(segs, f"C{mg}", f"=$E${d}", "calc", PCT)
    for k, c in enumerate(YEARS, start=1):
        p = chr(ord(c) - 1)
        if k <= 4:
            step = f"{p}{rv}+($H{rv}-{p}{rv})*{c}$6"
        elif k == 5:
            step = f"$C{rv}+($D${d}-$C{rv})/3"
        elif k <= 9:
            step = f"{p}{rv}+($D${d}-{p}{rv})*{c}$6"
        else:
            step = f"$D${d}"
        put(segs, f"{c}{rv}", f"=IF($I${d}=1,IF({c}$5>5,({c}$5-5)*$D${d}/5,0),{step})", "calc", M)
        put(segs, f"{c}{mg}", f"=IF({c}$5>$C$17,$F${d},$F${d}-($F${d}-$E${d})/$C$17*($C$17-{c}$5))", "calc", PCT)
        put(segs, f"{c}{eb}", f"={c}{rv}*{c}{mg}", "calc", M)
        put(segs, f"{c}{ri}", f"=({c}{rv}-{p}{rv})/IF({c}$5<=5,$G${d},$H${d})", "calc", M)
    blocks.append((rv, eb, ri))
    top += 6
section(segs, top, "Total company", last="M")
TR, TE, TM, TRI = top + 1, top + 2, top + 3, top + 4
label(segs, f"B{TR}", "Total revenue", bold=True)
label(segs, f"B{TE}", "Total operating income (EBIT)", bold=True)
label(segs, f"B{TM}", "Operating margin")
label(segs, f"B{TRI}", "Total reinvestment", bold=True)
for c in ["C"] + YEARS:
    put(segs, f"{c}{TR}", "=" + "+".join(f"{c}{b[0]}" for b in blocks), "calc", M, bold=True).border = TOP
    if c != "C":
        put(segs, f"{c}{TE}", "=" + "+".join(f"{c}{b[1]}" for b in blocks), "calc", M, bold=True)
        put(segs, f"{c}{TM}", f"={c}{TE}/{c}{TR}", "calc", PCT)
        put(segs, f"{c}{TRI}", "=" + "+".join(f"{c}{b[2]}" for b in blocks), "calc", M, bold=True)
segs.freeze_panes = "C6"
printable(segs)

# ================================================================= DCF
title(dcf, "Discounted Cash Flow (FCFF)", "Damodaran fcffsimpleginzu mechanics: tax drifts to marginal in years 6-10, NOL shield, "
      "cost of capital drifts to terminal, terminal reinvestment = g / ROIC. $ in millions except per share.")
widths(dcf, {"B": 44, **{c: 12 for c in "CDEFGHIJKLMN"}})
hdr(dcf, 4, {"C": "Base (TTM)", **{c: f"Year {k}" for k, c in enumerate(YEARS, start=1)}, "N": "Terminal"})
put(dcf, "B5", "Year (t)", "label")
for k, c in enumerate(["C"] + YEARS):
    put(dcf, f"{c}5", k, "label", "0").font = NOTE
lines = {7: "Revenue", 8: "Revenue growth", 9: "Operating margin", 10: "Operating income (EBIT)", 11: "Tax rate",
         12: "Net operating loss (NOL)", 13: "After-tax operating income", 14: "Less: reinvestment",
         15: "Free cash flow to firm (FCFF)", 17: "Cost of capital", 18: "Cumulative discount factor",
         19: "Present value of FCFF", 21: "Invested capital", 22: "Return on invested capital"}
for rr, t in lines.items():
    label(dcf, f"B{rr}", t, bold=rr in (7, 15))
TG = "$C$31"  # terminal growth cell (valuation block)
put(dcf, "C7", f"=Segments!C{TR}", "link", M)
put(dcf, "C9", "=C10/C7", "calc", PCT)
put(dcf, "C10", f"={I('Operating income (EBIT), TTM')}+C57", "calc", M)
put(dcf, "C11", f"={I('Effective tax rate')}", "link", PCT)
put(dcf, "C12", f"={I('Net operating loss carried forward')}", "link", M)
put(dcf, "C13", "=IF(C10>0,C10*(1-C11),C10)", "calc", M)
put(dcf, "C21", f"={I('Book value of equity')}+{I('Debt incl. finance leases')}-{I('Cash & marketable securities')}+C55", "calc", M)
put(dcf, "C22", "=C13/C21", "calc", PCT)
for k, c in enumerate(YEARS, start=1):
    p = chr(ord(c) - 1)
    put(dcf, f"{c}7", f"=Segments!{c}{TR}", "link", M, bold=True)
    put(dcf, f"{c}8", f"={c}7/{p}7-1", "calc", PCT)
    put(dcf, f"{c}9", f"={c}10/{c}7", "calc", PCT)
    put(dcf, f"{c}10", f"=Segments!{c}{TE}", "link", M)
    put(dcf, f"{c}11", f"={p}11" if k <= 5 else f"={p}11+({I('Marginal tax rate')}-{I('Effective tax rate')})/5", "calc", PCT)
    put(dcf, f"{c}12", f"=IF({c}10<0,{p}12-{c}10,IF({p}12>{c}10,{p}12-{c}10,0))", "calc", M)
    put(dcf, f"{c}13", f"=IF({c}10>0,IF({c}10<{p}12,{c}10,{c}10-({c}10-{p}12)*{c}11),{c}10)", "calc", M)
    put(dcf, f"{c}14", f"=Segments!{c}{TRI}", "link", M)
    put(dcf, f"{c}15", f"={c}13-{c}14", "calc", M, bold=True).border = TOP
    put(dcf, f"{c}17", "=$C$27" if k <= 5 else f"={p}17-($H$17-$N$17)/5", "calc", PCT2)
    put(dcf, f"{c}18", f"=1/(1+{c}17)" if k == 1 else f"={p}18/(1+{c}17)", "calc", NUM4)
    put(dcf, f"{c}19", f"={c}15*{c}18", "calc", M)
    put(dcf, f"{c}21", f"={p}21+{c}14", "calc", M)
    put(dcf, f"{c}22", f"={c}13/{p}21", "calc", PCT)
put(dcf, "N7", f"=M7*(1+{TG})", "calc", M, bold=True)
put(dcf, "N8", "=N7/M7-1", "calc", PCT)
put(dcf, "N9", "=M9", "calc", PCT)
put(dcf, "N10", "=N9*N7", "calc", M)
put(dcf, "N11", f"={I('Marginal tax rate')}", "link", PCT)
put(dcf, "N13", "=N10*(1-N11)", "calc", M)
put(dcf, "N14", f"=IF({TG}>0,{TG}/$C$32*N13,0)", "calc", M)
put(dcf, "N15", "=N13-N14", "calc", M, bold=True).border = TOP
put(dcf, "N17", "=$C$28", "calc", PCT2)
put(dcf, "N22", "=$C$32", "calc", PCT)
section(dcf, 25, "Valuation", last="N")
vrows = [
    (26, "Cost of capital (from WACC tab)", "=WACC_Initial", PCT2, "link"),
    (27, "Initial cost of capital (incl. flex)", f"=C26+{FX['d_init']}", PCT2, "calc"),
    (28, "Terminal cost of capital (live)", f"={LV['term']}+{FX['d_term']}+{FX['d_rf']}", PCT2, "calc"),
    (29, "Riskfree rate (live)", "=WACC!C20", PCT2, "link"),
    (30, "Stable growth rule", '="Terminal growth = riskfree rate (Ginzu default)"', None, "calc"),
    (31, "Terminal growth rate", "=C29", PCT2, "calc"),
    (32, "Terminal return on capital", f"={I('Terminal return on capital (ROIC)')}", PCT, "link"),
    (33, "Terminal value  = FCFF(t+1) / (r - g)", "=N15/(C28-C31)", M, "calc"),
    (34, "PV of terminal value", "=C33*M18", M, "calc"),
    (35, "PV of FCFF, years 1-10", "=SUM(D19:M19)", M, "calc"),
    (36, "Sum of present values", "=C34+C35", M, "calc"),
    (37, "Probability of failure", f"={I('Probability of failure')}", PCT, "link"),
    (38, "Proceeds if the firm fails", f"=C36*{I('Distress proceeds (% of fair value)')}", M, "calc"),
    (39, "Value of operating assets", "=C36*(1-C37)+C38*C37", M, "calc"),
    (40, "Less: debt", f"={I('Debt incl. finance leases')}", M, "link"),
    (41, "Less: minority interests", f"={I('Minority interests')}", M, "link"),
    (42, "Plus: cash & marketable securities", f"={I('Cash & marketable securities')}", M, "link"),
    (43, "Plus: non-operating assets", f"={I('Non-operating assets (cross holdings)')}", M, "link"),
    (44, "Plus: IPO proceeds not yet received", f"={I('IPO proceeds not yet received')}", M, "link"),
    (45, "Value of equity", "=C39-C40-C41+C42+C43+C44", M, "calc"),
    (46, "Less: value of employee options", f"={I('Value of employee options')}", M, "link"),
    (47, "Value of equity in common stock", "=C45-C46", M, "calc"),
    (48, "Shares outstanding (m)", f"={I('Shares outstanding (m)')}", M, "link"),
    (49, "Estimated value per share ($)", "=C47/C48", PS, "calc"),
    (50, "Current share price ($)", f"={I('Share price ($)')}", PS, "link"),
    (51, "Price as % of value", "=C50/C49", PCT, "calc"),
    (52, "Upside / (downside) to value", "=C49/C50-1", PCT, "calc"),
]
for rr, lab, f, fmt, kind in vrows:
    label(dcf, f"B{rr}", lab, bold=rr in (39, 45, 49))
    put(dcf, f"C{rr}", f, kind, fmt, bold=rr in (39, 45, 49))
for rr in (39, 45, 47):
    dcf[f"C{rr}"].border = TOP
dcf["C49"].border = BOX
dcf["C49"].fill = FILL_LIGHT
section(dcf, 54, "R&D capitalization (Ginzu R&D converter)", last="N")
life = f"Inputs!$C${R['rd_life']}"
past = [f"Inputs!$C${R['rd_past0'] + k}" for k in range(5)]
put(dcf, "B55", "Value of research asset", "label")
put(dcf, "C55", f"=Inputs!$C${R['rd_cur']}+" + "+".join(f"{p}*MAX(0,{life}-{k})/{life}" for k, p in enumerate(past, start=1)), "calc", M)
put(dcf, "B56", "Amortization of research asset (this year)", "label")
put(dcf, "C56", "=(" + "+".join(f"{p}*({k}<={life})" for k, p in enumerate(past, start=1)) + f")/{life}", "calc", M)
put(dcf, "B57", "Adjustment to operating income (R&D - amortization)", "label")
put(dcf, "C57", f"=Inputs!$C${R['rd_cur']}-C56", "calc", M)
name(wb, "ValuePerShare", "DCF!$C$49")
name(wb, "SharePrice", f"Inputs!$C${R['Share price ($)']}")
name(wb, "OperatingAssets", "DCF!$C$39")
dcf.freeze_panes = "C6"
printable(dcf)
dcf.page_setup.fitToHeight = 1

# ================================================================= SENSITIVITY tables
section(sens, 11, "Value per share: initial vs terminal cost of capital (change from base)", last="I")
put(sens, "B12", "Rows: change in initial cost of capital | Columns: change in terminal cost of capital", "label").font = NOTE
put(sens, "B13", "=ValuePerShare", "calc", PS)
for c, v in zip("CDEFG", (-0.01, -0.005, 0, 0.005, 0.01)):
    put(sens, f"{c}13", v, "input", PCT2).fill = FILL_LIGHT
for i, v in enumerate((-0.02, -0.015, -0.01, -0.005, 0, 0.005, 0.01, 0.015, 0.02), start=14):
    put(sens, f"B{i}", v, "input", PCT2).fill = FILL_LIGHT
    for c in "CDEFG":
        sens[f"{c}{i}"].number_format = PS
T1 = "C14:G22"
section(sens, 25, "Value per share: 2036 revenue multiplier vs target-margin change (all segments)", last="I")
put(sens, "B26", "Rows: 2036 revenue multiplier (1.89x = breakeven vs price) | Columns: change in target margins", "label").font = NOTE
put(sens, "B27", "=ValuePerShare", "calc", PS)
for c, v in zip("CDEFG", (-0.10, -0.05, 0, 0.05, 0.10)):
    put(sens, f"{c}27", v, "input", PCT).fill = FILL_LIGHT
for i, v in enumerate((0.75, 1.0, 1.25, 1.5, 1.75, round(ENG_BE, 3), 2.0), start=28):
    put(sens, f"B{i}", v, "input", X).fill = FILL_LIGHT
    for c in "CDEFG":
        sens[f"{c}{i}"].number_format = PS
T2 = "C28:G34"
section(sens, 37, "Riskfree rate: value per share and cost of capital (terminal cost of capital and growth move with it)", last="I")
put(sens, "B38", "Change in riskfree rate", "label").font = BOLD
put(sens, "C38", "Value / share", "label").font = BOLD
put(sens, "D38", "Initial WACC", "label").font = BOLD
put(sens, "B39", "(base)", "label").font = NOTE
put(sens, "C39", "=ValuePerShare", "calc", PS)
put(sens, "D39", "=WACC_Initial", "calc", PCT2)
for i, v in enumerate((-0.015, -0.01, -0.005, 0, 0.005, 0.01, 0.015), start=40):
    put(sens, f"B{i}", v, "input", PCT2).fill = FILL_LIGHT
    sens[f"C{i}"].number_format, sens[f"D{i}"].number_format = PS, PCT2
T3 = "C40:D46"
label(sens, "I14", "Current share price", bold=True)
put(sens, "I15", "=SharePrice", "link", PS)
for rng in (T1, T2, "C40:C46"):
    sens.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["$I$15"],
                                                    fill=PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")))
put(sens, "I13", "Shaded green: value at or above the current share price", "label").font = NOTE
printable(sens)
sens.page_setup.fitToHeight = 1

# ================================================================= DIAGNOSTICS
title(diag, "Diagnostics", "The Ginzu 'Diagnostics' sheet: the forecast set against industry averages, with the questions Damodaran "
      "suggests asking")
widths(diag, {"B": 30, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 70})
drows = [
    ("1. Revenue growth", [("Industry (5-yr avg)", f"={I('Revenue growth (5-yr average)')}", PCT), ("Last year", f"=DCF!C7/{I('Revenue, FY2025')}-1", PCT),
                           ("Next year", "=DCF!D8", PCT), ("CAGR yrs 2-5", "=(DCF!H7/DCF!D7)^(1/4)-1", PCT)],
     "Growth above industry: is the company small? Very different from last year's growth: why?"),
    ("2. Dollar revenues", [("Now (TTM)", "=DCF!C7", M), ("Year 1", "=DCF!D7", M), ("Year 5", "=DCF!H7", M), ("Year 10", "=DCF!M7", M)],
     "How big is the total market; what do its biggest firms earn; what share does year 10 imply? (Damodaran rejected the prospectus $28trn TAM as fantasy.)"),
    ("3. Operating margin", [("Industry", f"={I('Pre-tax operating margin (R&D adjusted)')}", PCT), ("Now (R&D adj.)", "=DCF!C9", PCT),
                             ("Year 5", "=DCF!H9", PCT), ("Year 10", "=DCF!M9", PCT)],
     "Industry margins? Unit economics (the cost of the extra unit)? What does competition look like?"),
    ("4. Reinvestment & ROIC", [("Industry sales/capital", f"={I('Sales / capital')}", NUM2), ("Sales/capital now", "=DCF!C7/DCF!C21", NUM2),
                                ("ROIC year 10", "=DCF!M22", PCT), ("Stable ROIC", "=DCF!N22", PCT)],
     "Bounce-back or new growth? Excess capacity? Does efficiency change with scale? H1-2026 capex was $28.5bn vs ~$10bn year-1 reinvestment here."),
    ("5. Cost of capital", [("Industry", f"={I('Cost of capital')}", PCT), ("Years 1-5", "=DCF!D17", PCT),
                            ("Compounded", "=(1/DCF!M18)^(1/10)-1", PCT), ("Stable", "=DCF!N17", PCT)],
     "How does it compare with the industry? What happens to it over time, and why?"),
    ("6. Value vs price", [("Value / price", "=DCF!C49/DCF!C50", PCT), ("Flag", '=IF(DCF!C49/DCF!C50<0.5,"Value low: revisit",IF(DCF!C49/DCF!C50>2,"Value high: revisit","Within 0.5x-2x"))', None)],
     "Ginzu flags: value below 50% or above 200% of price -> revisit growth, target margin, sales/capital and terminal ROIC."),
]
rr = 5
for head, cells, ask in drows:
    section(diag, rr, head, last="H")
    for j, (lab, f, fmt) in enumerate(cells):
        col = "CDEFG"[j]
        c = diag[f"{col}{rr + 1}"]
        c.value, c.font, c.fill = lab, BOLD, FILL_LIGHT
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        put(diag, f"{col}{rr + 2}", f, "link" if f.startswith(("=DCF", "=Inputs")) else "calc", fmt)
    put(diag, f"H{rr + 1}", ask, "label").alignment = Alignment(wrap_text=True, vertical="top")
    diag.merge_cells(f"H{rr + 1}:H{rr + 2}")
    diag.row_dimensions[rr + 1].height = 30
    rr += 4
printable(diag)

# ================================================================= SUMMARY
title(summ, "Valuation Summary: Space Exploration Technologies Corp. (NASDAQ: SPCX)",
      "Intrinsic value, Damodaran FCFF framework | Valuation date 25-Sep-2026 | $ in millions except per share")
widths(summ, {"B": 54, "C": 14, "D": 12, "E": 3, "F": 40, "G": 14, "H": 3, "T": 34, "U": 11, "V": 11, "W": 11, "X": 11})
section(summ, 5, "Key outputs", last="D")
kp = [("Intrinsic value per share ($)", "=ValuePerShare", PS), ("Current share price ($)", "=SharePrice", PS),
      ("Upside / (downside) to value", "=DCF!C52", PCT), ("Price as % of value", "=DCF!C51", PCT),
      ("Market capitalization", f"=SharePrice*{I('Shares outstanding (m)')}", M),
      ("Market enterprise value", f"=C10+{I('Debt incl. finance leases')}-{I('Cash & marketable securities')}", M),
      ("Intrinsic value of operating assets", "=OperatingAssets", M),
      ("Intrinsic value of equity", "=DCF!C47", M), ("Initial / terminal cost of capital", '=TEXT(DCF!C27,"0.00%")&" / "&TEXT(DCF!C28,"0.00%")', None),
      ("Terminal growth (= riskfree)", "=DCF!C31", PCT2), ("2036 revenue (model)", "=DCF!M7", M),
      ("Damodaran June 2026 value per share ($)", f"={I('Damodaran June 2026 value per share ($)')}", PS),
      ("Revenue breakeven vs price (x the 2036 targets)", round(ENG_BE, 3), X)]
for i, (lab, f, fmt) in enumerate(kp, start=6):
    label(summ, f"B{i}", lab, bold=i in (6, 8))
    kind = "input" if isinstance(f, float) else ("link" if "Inputs!" in str(f) else "calc")
    put(summ, f"C{i}", f, kind, fmt, bold=i in (6, 8))
summ["C6"].fill = FILL_LIGHT
summ["C6"].border = BOX
summ["C14"].alignment = Alignment(horizontal="right")
put(summ, "B19", "Breakeven computed by the engine (spacex_model.py); the same ratio sits in Sensitivity, row 33.", "label").font = NOTE
section(summ, 5, "Enterprise value to equity bridge", first="F", last="G")
bridge = [("Value of operating assets", "=DCF!C39"), ("Less: debt", "=-DCF!C40"), ("Less: minority interests", "=-DCF!C41"),
          ("Plus: cash & marketable securities", "=DCF!C42"), ("Plus: non-operating assets", "=DCF!C43"),
          ("Plus: IPO proceeds not yet received", "=DCF!C44"), ("Less: employee options", "=-DCF!C46"),
          ("Value of equity in common stock", "=DCF!C47"), ("Shares outstanding (m)", "=DCF!C48"), ("Value per share ($)", "=DCF!C49")]
for i, (lab, f) in enumerate(bridge, start=6):
    label(summ, f"F{i}", lab, bold=i in (13, 15))
    put(summ, f"G{i}", f, "link", PS if i == 15 else M, bold=i in (13, 15))
summ["G13"].border = TOP
section(summ, 21, "Scenario analysis (live; one change at a time)", first="B", last="D")
hdr(summ, 22, {"B": "Scenario", "C": "Value / share", "D": "vs price"})
for k in range(12):
    label(summ, f"B{23 + k}", f"=Scenarios!C{8 + k}")
    summ[f"B{23 + k}"].font = GREEN
    put(summ, f"C{23 + k}", f"=Scenarios!K{8 + k}", "link", PS)
    put(summ, f"D{23 + k}", f"=Scenarios!L{8 + k}", "link", PCT)
section(summ, 37, "Value walk: Damodaran's June 2026 value to this update ($ per share)", first="B", last="D")
hdr(summ, 38, {"B": "Step (engine attribution; order-dependent)", "C": "Value / share", "D": "Change"})
walk_labels = ["Damodaran June 2026 value (IPO at $135)", "Roll-forward: TTM base year, IPO cash, 13.18bn shares",
               "Riskfree 4.56% -> 5.17% (terminal 8.25% -> 8.86%)", "Beta weights: April mix -> TTM segment revenues",
               "Country ERPs: Jan-2026 -> July-2026 (final)"]
assert len(walk_labels) == len(WALK)
for i, w in enumerate(WALK):
    rrw = 39 + i
    label(summ, f"B{rrw}", walk_labels[i])
    put(summ, f"C{rrw}", w["value"], "input", PS)
    put(summ, f"D{rrw}", "" if i == 0 else f"=C{rrw}-C{rrw - 1}", "calc", '+$#,##0.00;-$#,##0.00')
# chart data (outside the print area)
section(summ, 5, "Chart data: football field ($ per share)", first="T", last="W")
hdr(summ, 6, {"T": "Reference", "U": "Low", "V": "High - low", "W": "High"})
ff = [("52-week trading range", f"={I('52-week low ($)')}", f"={I('52-week high ($)')}"),
      ("Analyst average target (third party)", f"={I('Analyst average price target ($)')}-1", f"={I('Analyst average price target ($)')}+1"),
      ("DCF: cost of capital sensitivity", f"=MIN(Sensitivity!{T1})", f"=MAX(Sensitivity!{T1})"),
      ("DCF: revenue x margin sensitivity", f"=MIN(Sensitivity!{T2})", f"=MAX(Sensitivity!{T2})"),
      ("DCF: scenario range", "=MIN(Scenarios!K8:K19)", "=MAX(Scenarios!K8:K19)"),
      ("Damodaran June 2026 value", f"={I('Damodaran June 2026 value per share ($)')}-1", f"={I('Damodaran June 2026 value per share ($)')}+1"),
      ("DCF base case", "=ValuePerShare-1", "=ValuePerShare+1"),
      ("Current share price", "=SharePrice-1", "=SharePrice+1")]
for i, (lab, lo, hi) in enumerate(ff, start=7):
    label(summ, f"T{i}", lab)
    put(summ, f"U{i}", lo, "calc", PS)
    put(summ, f"W{i}", hi, "calc", PS)
    put(summ, f"V{i}", f"=W{i}-U{i}", "calc", PS)
ch = BarChart()
ch.type, ch.grouping, ch.overlap = "bar", "stacked", 100
ch.title = "Football field: value per share ($)"
ch.add_data(Reference(summ, min_col=21, max_col=22, min_row=6, max_row=14), titles_from_data=True)
ch.set_categories(Reference(summ, min_col=20, min_row=7, max_row=14))
ch.series[0].graphicalProperties.noFill = True
ch.series[0].graphicalProperties.line.noFill = True
ch.series[1].graphicalProperties.solidFill = MID
ch.legend = None
ch.y_axis.majorGridlines = None
ch.y_axis.number_format = '"$"#,##0'
ch.x_axis.delete = False
ch.y_axis.delete = False
ch.height, ch.width = 8.2, 17.5
summ.add_chart(ch, "I5")
# value-walk waterfall (invisible base + down / up / total)
section(summ, 17, "Chart data: value walk waterfall ($ per share)", first="T", last="X")
hdr(summ, 18, {"T": "Step", "U": "Base", "V": "Decrease", "W": "Increase", "X": "Total"})
short = ["June 2026", "Roll-forward", "Riskfree", "Beta weights", "Country ERPs", "Sept 2026"]
n = len(WALK)
for i in range(n + 1):
    rrw = 19 + i
    label(summ, f"T{rrw}", short[i])
    if i == 0:
        put(summ, f"X{rrw}", "=C39", "calc", PS)
    elif i == n:
        put(summ, f"X{rrw}", f"=C{38 + n}", "calc", PS)
    else:
        cur, prev = f"C{39 + i}", f"C{38 + i}"
        put(summ, f"U{rrw}", f"=MIN({cur},{prev})", "calc", PS)
        put(summ, f"V{rrw}", f"=MAX({prev}-{cur},0)", "calc", PS)
        put(summ, f"W{rrw}", f"=MAX({cur}-{prev},0)", "calc", PS)
wf = BarChart()
wf.type, wf.grouping, wf.overlap = "col", "stacked", 100
wf.title = "Value walk: June 2026 to September 2026 ($ per share)"
wf.add_data(Reference(summ, min_col=21, max_col=24, min_row=18, max_row=18 + n + 1), titles_from_data=True)
wf.set_categories(Reference(summ, min_col=20, min_row=19, max_row=19 + n))
wf.series[0].graphicalProperties.noFill = True
wf.series[0].graphicalProperties.line.noFill = True
for s_, color in zip(wf.series[1:], ("C0504D", "4F8A10", NAVY)):
    s_.graphicalProperties.solidFill = color
wf.legend = None
wf.y_axis.majorGridlines = None
wf.y_axis.number_format = '"$"#,##0'
wf.x_axis.delete = False
wf.y_axis.delete = False
wf.height, wf.width = 8.2, 17.5
summ.add_chart(wf, "I24")
printable(summ)
summ.print_area = "A1:S45"
summ.page_setup.fitToHeight = 1

# ================================================================= CHECKS
title(chk, "Model Checks", "Ties the workbook to the aswath-finance engine (calibrated to Damodaran's spreadsheets) and to the Ginzu "
      "rules. Engine reference values are hard-coded outputs of spacex_model.py on 2026-09-25.")
widths(chk, {"B": 6, "C": 62, "D": 16, "E": 16, "F": 12, "G": 60})
hdr(chk, 5, {"B": "#", "C": "Check", "D": "Workbook", "E": "Engine / rule", "F": "Result", "G": "Note"})
BASESTATE = f"AND(Scenarios!$D$5=1,{FX['d_init']}=0,{FX['d_term']}=0,{FX['rev_mult']}=1,{FX['d_margin']}=0,{FX['d_rf']}=0)"
checks = [
    ("Model in base state (scenario 1, flex cells neutral)", f'=IF({BASESTATE},"Base","Flexed")', '="Base"', f"={BASESTATE}",
     "Tie-outs below are only meaningful in base state"),
    ("Value per share ties to engine", "=ValuePerShare", ENG_VPS, f"=IF({BASESTATE},ABS(D7-E7)<0.005,TRUE)", "spacex_model.py sep_2026_inputs.json"),
    ("Initial cost of capital ties to engine", "=WACC_Initial", ENG_WACC, f"=IF({BASESTATE},ABS(D8-E8)<0.000001,TRUE)", "cost_of_capital.py (Damodaran's method)"),
    ("Scenario values tie to engine (max abs difference, $)",
     "=MAX(" + ",".join(f"ABS(Scenarios!K{8 + k}-{v!r})" for k, (_, v) in enumerate(ENG_SCEN)) + ")", 0.0, "=D9<0.01",
     "spacex_model.py --sensitivity, 12 scenarios"),
    ("Year-10 revenue = sum of live 2036 targets", "=DCF!M7", "=SUM(Segments!D10:D13)", "=ABS(D10-E10)<0.001", "Segment path lands on the 2036 targets"),
    ("Terminal growth <= riskfree (Ginzu rule)", "=DCF!C31", "=DCF!C29", "=D11<=E11+0.0000001", "Stable growth capped at riskfree"),
    ("Terminal ROIC >= terminal cost of capital (Ginzu rule)", "=DCF!C32", "=DCF!C28", "=D12>=E12", "Otherwise growth destroys value"),
    ("Sensitivity table 1 centre = base value", "=Sensitivity!E18", "=ValuePerShare", f"=IF({BASESTATE},ABS(D13-E13)<0.005,TRUE)", "Data table recalculated by Excel"),
    ("Sensitivity table 2 (1.00x, 0%) = base value", "=Sensitivity!E29", "=ValuePerShare", f"=IF({BASESTATE},ABS(D14-E14)<0.005,TRUE)", ""),
    ("Equity bridge adds up", "=DCF!C45", "=DCF!C39-DCF!C40-DCF!C41+DCF!C42+DCF!C43+DCF!C44", "=ABS(D15-E15)<0.001", ""),
    ("Value walk ends at the engine value", f"=Summary!C{38 + n}", ENG_VPS, "=ABS(D16-E16)<0.005", "Engine attribution"),
    ("Breakeven multiplier reproduces the price", "=Sensitivity!E33", "=SharePrice", f"=IF({BASESTATE},ABS(D17-E17)<0.05,TRUE)", "Row 1.89x, 0% margin change"),
]
for i, (lab, wv, ev, res, note) in enumerate(checks, start=6):
    put(chk, f"B{i}", i - 5, "label", "0")
    label(chk, f"C{i}", lab)
    put(chk, f"D{i}", wv, "calc", '0.0000%' if i in (8,) else (PCT2 if i in (11, 12) else ('$#,##0.0000' if i in (7, 9, 13, 14, 16, 17) else M)))
    put(chk, f"E{i}", ev, "input" if isinstance(ev, float) else "calc",
        '0.0000%' if i in (8,) else (PCT2 if i in (11, 12) else ('$#,##0.0000' if i in (7, 9, 13, 14, 16, 17) else M)))
    put(chk, f"F{i}", res, "calc")
    put(chk, f"G{i}", note, "label").font = NOTE
last = 5 + len(checks)
chk.conditional_formatting.add(f"F6:F{last}", CellIsRule(operator="equal", formula=["TRUE"], fill=PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")))
chk.conditional_formatting.add(f"F6:F{last}", CellIsRule(operator="equal", formula=["FALSE"], fill=PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")))
label(chk, f"C{last + 2}", "ALL CHECKS PASS", bold=True)
put(chk, f"F{last + 2}", f"=COUNTIF(F6:F{last},FALSE)=0", "calc", bold=True).border = BOX
name(wb, "AllChecksPass", f"Checks!$F${last + 2}")
printable(chk)

# ================================================================= SOURCES & NOTES
title(src, "Sources, Methodology & Notes", "Every input traces to a filing, market data, or Damodaran's own model/data")
widths(src, {"B": 4, "C": 34, "D": 110})
srows = [
    ("Filings", "SpaceX Q2-2026 Form 10-Q: https://www.sec.gov/Archives/edgar/data/0001181412/000162828026052535/spcx-20260630.htm"),
    ("", "SpaceX Q2-2026 earnings release: https://www.sec.gov/Archives/edgar/data/1181412/000162828026052515/earningsreleaseq22608042.htm"),
    ("Market data", "StockAnalysis SPCX (close, 52-week range, analyst target), 25-Sep-2026: https://stockanalysis.com/stocks/spcx/"),
    ("", "US 10-year Treasury 5.17%, 25-Sep-2026: https://www.cnbc.com/2026/09/25/treasury-yields-bonds-debt.html"),
    ("", "Q2-2026 results context (AI capex): https://www.cnbc.com/2026/08/04/spacex-spcx-earnings-live-updates-q2-2026.html"),
    ("Damodaran model", "Damodaran, SpaceX2026IPOUpdated.xlsx (June 2026): https://pages.stern.nyu.edu/~adamodar/pc/blog/SpaceX2026IPOUpdated.xlsx"),
    ("Damodaran posts", "Revisiting the SpaceX Valuation (2026-06): https://aswathdamodaran.blogspot.com/2026/06/a-weeks-ago-i-assessed-value-of-spacex.html"),
    ("", "To a Trillion(s) Dollars and beyond: A SpaceX IPO Odyssey (2026-04): https://aswathdamodaran.blogspot.com/2026/04/to-trillion-dollars-and-beyond-spacex.html"),
    ("Damodaran data", "fcffsimpleginzu.xlsx industry averages & rating spreads (Jan-2026); ctrypremJuly26.xlsx country ERPs (Jul-2026): pages.stern.nyu.edu/~adamodar"),
    ("Methodology", "FCFF DCF transcribed from Damodaran's fcffsimpleginzu spreadsheet; the aswath-finance engine reproduces 31 published Damodaran valuations exactly."),
    ("", "Segment mechanics transcribed from Damodaran's June 2026 SpaceX workbook (rebuilt from raw inputs to its $97.8277/share)."),
    ("", "Cost of capital: Damodaran's worksheet method (multi-business beta, operating-country ERP, rating spread), reproduced exactly at 8.3745% for June."),
    ("Judgment calls", "1) Beta weights updated from the stale April mix to TTM segment revenues (the literal June weights: see Scenario 2)."),
    ("", "2) Terminal cost of capital = the June model's 8.25% moved with the riskfree rate (+0.61pt), following Damodaran's riskfree-adjustment convention."),
    ("", "3) Base segment margins, operating-country weights, Ba2/BB rating and excluded options kept from the June model (not disclosed in the 10-Q)."),
    ("", "4) Share count from the 10-Q cover (13,182m); StockAnalysis shows 13.57bn, which would lower value/share by about 2.9%."),
    ("", "5) The story (2036 revenue targets, target margins, sales-to-capital) is Damodaran's June 2026 story, unchanged."),
    ("Damodaran's stance (June 2026)", 'At the $135 offer price: "too richly priced for my tastes"; "as an investor, I would not sell short, notwithstanding my value assessment".'),
    ("Disclaimer", "Educational application of Damodaran's public framework and data. Not investment advice, not a recommendation, and not affiliated with or endorsed by Damodaran or any bank."),
]
section(src, 5, "Sources and notes", first="B", last="D")
for i, (k, v) in enumerate(srows, start=6):
    label(src, f"C{i}", k, bold=bool(k)).alignment = Alignment(vertical="top")
    c = put(src, f"D{i}", v, "label")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if "http" in v:
        c.hyperlink = v[v.index("http"):].split()[0]
        c.font = Font(name=F, size=10, color="1F4E79", underline="single")
printable(src)

# ================================================================= COVER
cover.sheet_view.showGridLines = False
widths(cover, {"A": 3, "B": 4, "C": 38, "D": 70})
for rr in range(2, 9):
    for col in "BCD":
        cover[f"{col}{rr}"].fill = FILL_NAVY
cover["C3"] = "Space Exploration Technologies Corp. (NASDAQ: SPCX)"
cover["C3"].font = Font(name=F, size=20, bold=True, color="FFFFFF")
cover["C5"] = "Intrinsic Valuation: Damodaran FCFF (Ginzu) Framework"
cover["C5"].font = Font(name=F, size=13, color="FFFFFF")
cover["C7"] = "Valuation date: 25 September 2026  |  $ in millions except per share  |  For internal discussion"
cover["C7"].font = Font(name=F, size=10, italic=True, color="DCE6F1")
label(cover, "C10", "Headline", bold=True)
put(cover, "C11", "Intrinsic value per share", "label")
put(cover, "D11", "=ValuePerShare", "calc", PS, bold=True)
put(cover, "C12", "Share price (25-Sep-2026)", "label")
put(cover, "D12", "=SharePrice", "calc", PS)
put(cover, "C13", "Price as % of value", "label")
put(cover, "D13", "=DCF!C51", "calc", PCT)
put(cover, "C14", "Model checks", "label")
put(cover, "D14", '=IF(AllChecksPass,"All checks pass","CHECK FAILURES - see Checks tab")', "calc", bold=True)
for c in ("D11", "D12", "D13", "D14"):
    cover[c].alignment = Alignment(horizontal="left")
label(cover, "C16", "Contents", bold=True)
toc = [("Summary", "Key outputs, EV-to-equity bridge, scenarios, football field, value walk"),
       ("Inputs", "All hard-coded inputs with sources (blue)"), ("Scenarios", "Live scenario selector and scenario table"),
       ("Sensitivity", "Live data tables: cost of capital, revenue x margin, riskfree"),
       ("WACC", "Cost of capital build (Damodaran's method)"), ("Segments", "Segment revenue, margin and reinvestment build"),
       ("DCF", "FCFF projection, terminal value and equity bridge"), ("Diagnostics", "Ginzu diagnostic checks vs industry"),
       ("Checks", "Tie-out to the calibrated engine and to the Ginzu rules"), ("Sources", "Sources, methodology, judgment calls, disclaimer")]
for i, (sh, desc) in enumerate(toc, start=17):
    c = cover[f"C{i}"]
    c.value, c.hyperlink = sh, f"#'{sh}'!A1"
    c.font = Font(name=F, size=10, color="1F4E79", underline="single")
    put(cover, f"D{i}", desc, "label")
label(cover, "C28", "Colour convention", bold=True)
put(cover, "C29", "Blue: hard-coded input", "input")
put(cover, "C30", "Black: formula", "calc")
put(cover, "C31", "Green: link to another sheet", "link")
flex_c = put(cover, "C32", "Yellow fill: flex / selector cell (keep at base)", "label")
flex_c.fill = FILL_FLEX
put(cover, "C34", "Educational application of Damodaran's public framework and data; not investment advice. See Sources.",
    "label").font = NOTE
printable(cover, landscape=False)

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                size = 10 if (c.font is None or c.font.name == "Calibri") else (c.font.size or 10)
                c.font = Font(name=F, size=size, bold=c.font.bold if c.font else False,
                              italic=c.font.italic if c.font else False, color=c.font.color if c.font else None)
    ws.sheet_view.zoomScale = 85
wb.active = 0
wb.save(OUT)

# ---------------------------------------------------------------- data-table + recalc instructions for Excel
ps1 = HERE / "finalize_workbook.ps1"
ps1.write_text(f"""$ErrorActionPreference = 'Stop'
$path = Join-Path $PSScriptRoot '{OUT.name}'
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false; $xl.DisplayAlerts = $false
$wb = $xl.Workbooks.Open($path)
$m = [System.Reflection.Missing]::Value
$s = $wb.Worksheets.Item('Scenarios');   $null = $s.Range('N7:O19').Table($m, $s.Range('D5'))
$t = $wb.Worksheets.Item('Sensitivity')
$null = $t.Range('B13:G22').Table($t.Range('C6'), $t.Range('C5'))
$null = $t.Range('B27:G34').Table($t.Range('C8'), $t.Range('C7'))
$null = $t.Range('B39:D46').Table($m, $t.Range('C9'))
foreach ($r in @($t.Range('C14:G22'), $t.Range('C28:G34'), $t.Range('C40:C46'), $s.Range('O8:O19'))) {{
  $r.Font.Name = 'Arial'; $r.Font.Size = 10; $r.NumberFormat = '$#,##0.00_);($#,##0.00)' }}
$t.Range('D40:D46').Font.Name = 'Arial'; $t.Range('D40:D46').Font.Size = 10; $t.Range('D40:D46').NumberFormat = '0.00%'
$xl.CalculateFull()
$wb.Worksheets.Item('Cover').Activate()
$wb.Save(); $wb.Close($false); $xl.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($xl) | Out-Null
Write-Output 'finalized'
""", encoding="utf-8")
print(f"built {OUT.name}; engine value/share {ENG_VPS:.4f}, WACC {ENG_WACC:.6f}, breakeven x{ENG_BE:.3f}")
